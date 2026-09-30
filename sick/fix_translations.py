
import os
import sys
import glob
import time
from openpyxl import load_workbook

OUTPUT_DIR = "SICK_Products"

# Default replacements (key: old text, value: new text)
DEFAULT_REPLACEMENTS = {
    "Sự miêu tả": "Mô tả",
}


def parse_replacements():
    """Parse --replace 'old=>new' args."""
    replacements = dict(DEFAULT_REPLACEMENTS)
    args = sys.argv
    i = 0
    while i < len(args):
        if args[i] == "--replace" and i + 1 < len(args):
            spec = args[i + 1]
            if "=>" in spec:
                old, new = spec.split("=>", 1)
                replacements[old.strip()] = new.strip()
            i += 2
        else:
            i += 1
    return replacements


def fix_file(path, replacements, dry_run=False, vi_only=False, en_only=False):
    """Tra ve so cell duoc thay doi.
    vi_only: neu True, chi replace cot A(1) B(2) C(3) - khong dung cot D(4) E(5) (EN).
    en_only: neu True, chi replace cot D(4) E(5) - khong dung cot A(1) B(2) C(3).
    """
    wb = load_workbook(path)
    ws = wb.active

    changed = 0
    for row in ws.iter_rows():
        for cell in row:
            # Filter cot theo flag
            if vi_only and cell.column >= 4:
                continue
            if en_only and cell.column < 4:
                continue
            v = cell.value
            if not isinstance(v, str):
                continue
            new_v = v
            for old, new in replacements.items():
                if old in new_v:
                    new_v = new_v.replace(old, new)
            if new_v != v:
                if not dry_run:
                    cell.value = new_v
                changed += 1

    if changed > 0 and not dry_run:
        wb.save(path)
    wb.close()
    return changed


def main():
    dry_run = "--dry-run" in sys.argv
    vi_only = "--vi-only" in sys.argv
    en_only = "--en-only" in sys.argv
    family_filter = None
    if "--family" in sys.argv:
        idx = sys.argv.index("--family")
        if idx + 1 < len(sys.argv):
            family_filter = sys.argv[idx + 1]

    replacements = parse_replacements()
    print(f"Replacements ({len(replacements)}):")
    for old, new in replacements.items():
        print(f"  {old!r} -> {new!r}")
    print()

    if family_filter:
        folders = [d for d in os.listdir(OUTPUT_DIR) if family_filter.lower() in d.lower()]
    else:
        folders = sorted(d for d in os.listdir(OUTPUT_DIR) if os.path.isdir(os.path.join(OUTPUT_DIR, d)))

    print(f"Scanning {len(folders)} folders...")
    folder_files = []
    for folder in folders:
        folder_path = os.path.join(OUTPUT_DIR, folder)
        if not os.path.isdir(folder_path):
            continue
        files = [f for f in glob.glob(os.path.join(folder_path, "*.xlsx"))
                 if not os.path.basename(f).startswith("~$")]
        if files:
            folder_files.append((folder, files))
    grand_total = sum(len(fs) for _, fs in folder_files)
    print(f"Total {grand_total} files.\n")

    total_files = 0
    files_changed = 0
    total_cells_changed = 0
    errors = 0
    last_print = time.time()

    for folder, files in folder_files:
        for f in files:
            total_files += 1
            try:
                n = fix_file(f, replacements, dry_run=dry_run, vi_only=vi_only, en_only=en_only)
                if n > 0:
                    files_changed += 1
                    total_cells_changed += n
            except Exception as e:
                errors += 1
                print(f"  ERROR {f}: {e}")

            now = time.time()
            if now - last_print > 1 or total_files % 100 == 0:
                pct = total_files * 100 // max(grand_total, 1)
                print(f"  [{pct:3d}%] {total_files:>5}/{grand_total} | changed: {files_changed:>5} | now: {folder}", flush=True)
                last_print = now

    print()
    print("=" * 70)
    print(f"Total scanned: {total_files}")
    if dry_run:
        print(f"Files would change: {files_changed}")
        print(f"Cells would change: {total_cells_changed}")
    else:
        print(f"Files changed:   {files_changed}")
        print(f"Cells changed:   {total_cells_changed}")
    print(f"Errors:          {errors}")


if __name__ == "__main__":
    main()
