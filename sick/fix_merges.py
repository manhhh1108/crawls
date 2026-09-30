"""
Fix bug merge B:C va D:E cho cac row Title mo ta ngan, Mo ta thong so trong cac
file da qua add_columns.py.

Nguyen nhan: openpyxl insert_rows() khong shift merge cho 2 row cuoi cua content_rows
(Title mo ta ngan, Mo ta thong so) -> 2 row nay bi mat merge sau khi insert.

Script nay dam bao moi row content (3-11) co merge B:C va D:E. Idempotent.

Usage:
  python fix_merges.py                    # fix tat ca
  python fix_merges.py --dry-run          # preview
  python fix_merges.py --family WTT12-S   # 1 dong (test)
"""
import os
import sys
import glob
import time
from openpyxl import load_workbook
from openpyxl.styles import Alignment

OUTPUT_DIR = "SICK_Products"

# Cac row content can co merge B:C va D:E (theo format moi sau add_columns)
CONTENT_ROWS = list(range(3, 12))  # rows 3-11: Ma, Danh muc sp, Thuong hieu, URL, Keyword, Meta title, Meta desc, Title mo ta ngan, Mo ta thong so

WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def fix_file(path, dry_run=False):
    """Tra ve so merge da fix (0 neu khong can)."""
    wb = load_workbook(path)
    ws = wb.active

    existing_merges = {mr.coord for mr in ws.merged_cells.ranges}
    fixed = 0

    for r in CONTENT_ROWS:
        # Chi fix neu row co content (vd label o col A)
        label = ws.cell(row=r, column=1).value
        if not label:
            continue

        bc = f"B{r}:C{r}"
        de = f"D{r}:E{r}"

        if bc not in existing_merges and not dry_run:
            ws.merge_cells(bc)
            ws.cell(row=r, column=2).alignment = WRAP_TOP
            fixed += 1
        elif bc not in existing_merges:
            fixed += 1

        if de not in existing_merges and not dry_run:
            ws.merge_cells(de)
            ws.cell(row=r, column=4).alignment = WRAP_TOP
            fixed += 1
        elif de not in existing_merges:
            fixed += 1

    if fixed > 0 and not dry_run:
        wb.save(path)
    wb.close()
    return fixed


def main():
    dry_run = "--dry-run" in sys.argv
    family_filter = None
    if "--family" in sys.argv:
        idx = sys.argv.index("--family")
        if idx + 1 < len(sys.argv):
            family_filter = sys.argv[idx + 1]

    if family_filter:
        folders = [d for d in os.listdir(OUTPUT_DIR) if family_filter.lower() in d.lower()]
    else:
        folders = sorted(d for d in os.listdir(OUTPUT_DIR) if os.path.isdir(os.path.join(OUTPUT_DIR, d)))

    print(f"Scanning {len(folders)} folders...")
    folder_files = []
    for folder in folders:
        folder_path = os.path.join(OUTPUT_DIR, folder)
        if not os.path.isdir(folder_path):
            continue
        files = [f for f in glob.glob(os.path.join(folder_path, "*.xlsx"))
                 if not os.path.basename(f).startswith("~$")]
        if files:
            folder_files.append((folder, files))
    grand_total = sum(len(fs) for _, fs in folder_files)
    print(f"Total {grand_total} files to check.\n")

    total_files = 0
    files_fixed = 0
    total_merges_fixed = 0
    errors = 0
    last_print = time.time()

    for folder, files in folder_files:
        for f in files:
            total_files += 1
            try:
                n = fix_file(f, dry_run=dry_run)
                if n > 0:
                    files_fixed += 1
                    total_merges_fixed += n
            except Exception as e:
                errors += 1
                print(f"  ERROR {f}: {e}")

            now = time.time()
            if now - last_print > 1 or total_files % 100 == 0:
                pct = total_files * 100 // max(grand_total, 1)
                print(f"  [{pct:3d}%] {total_files:>5}/{grand_total} | fixed: {files_fixed:>5} | now: {folder}", flush=True)
                last_print = now

    print()
    print("=" * 70)
    print(f"Total scanned:    {total_files}")
    if dry_run:
        print(f"Files would fix:  {files_fixed}")
        print(f"Merges would fix: {total_merges_fixed}")
    else:
        print(f"Files fixed:      {files_fixed}")
        print(f"Total merges fixed: {total_merges_fixed}")
    print(f"Errors:           {errors}")


if __name__ == "__main__":
    main()
