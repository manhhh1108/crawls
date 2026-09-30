import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Thong ke SP SICK tren Web"

title_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
cat_font = Font(name='Arial', size=11, bold=True)
normal_font = Font(name='Arial', size=10)
sub_font = Font(name='Arial', size=10, bold=True, color='2E75B6')

title_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
cat_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
total_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')

thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 55
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 22

ws.merge_cells('A1:D1')
c = ws['A1']
c.value = 'THONG KE SAN PHAM TREN WEB SICK.COM/SG/EN'
c.font = title_font
c.fill = title_fill
c.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 35

headers = ['STT', 'Dong san pham', 'Ngay thong ke', 'So luong dong SP']
for col, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border

data = [
    ("Detection Sensors (Cam bien phat hien)", [
        ("Photoelectric Sensors", [
            "G2","G6","G10","G20","GR18","H18 Sure Sense","HISIC",
            "RAY10","RAY26","Roller Sensor Bar","V12","V18 Laser","V180",
            "W10","W100-2","W100 Laser","W11-2","W11G-2","W12","W16",
            "W18-3 EX","W2","W24","W250-2","W26","W27","W280-2","W34",
            "W4","W45","W8","W8 Laser","W8G","W9","Z18","Zonecontrol"
        ]),
        ("Inductive Proximity Sensors", [
            "IDF","IMA","IMB","IMC","IME","IMF","IMG","IMI","IMM","IMN",
            "IMP","IMR","IMS","IMW","IMX","IQB","IQE","IQG","IQL","IQM",
            "IQV","IQY","SAM"
        ]),
        ("Capacitive Proximity Sensors", ["CM","CMB","CQ","CQF"]),
        ("Magnetic Sensors", ["MIS","MLS","MME","MMN","MQB"]),
        ("Color Sensors", ["CSM","CSS/CSX High Speed","CSS High Resolution"]),
        ("Contrast Sensors", ["KT8","KTC","KTL180","KTM","KTS","KTX","OLS","RS10"]),
        ("Luminescence Sensors", ["LUT1","LUT3","LUT8","LUT9","LUTM"]),
        ("Fork Sensors", ["CFM/CFX","MF","UD18","UF","UFS","UFW","WF","WFE","WFL","WFM","WFS"]),
        ("Cylinder Sensors", [
            "MPA","MPS-C","MPS-G","MPS-M","MPS-T","MZ2Q-C","MZ2Q-T",
            "MZC1","MZC1 TWIN","MZC1 VIA","MZC2","MZCG","MZCG VIA","MZN",
            "MZT7","MZT7 TWIN","MZT8","MZT8 ATEX","MZT8 TWIN","MZT8 VIA",
            "RZC1","RZN","RZT7"
        ]),
        ("Automation Light Grids", [
            "ELG","FlexChain","FLG","HLG","MLG-2","MLG-2 WebChecker",
            "PLG","SLG","SLG-2","WLG"
        ]),
        ("Fiber Optic Sensors", ["Fiber Optic Cables","GLL170","WLL24","WLL80"]),
        ("Glare Sensors", ["Glare"]),
        ("Condition Monitoring Sensors", ["Multi Physics Box"]),
        ("Array Sensors", ["AS30"]),
    ]),
    ("Distance Sensors (Cam bien khoang cach)", [
        ("Laser Distance Sensors", [
            "DT20 Hi","Dx100","Dx1000","Dx35","Dx50","Dx50-2","Dx500","Dx80",
            "OD Mini","OD Precision","OD Value","OD1000","OD200","OD2000",
            "OD5000","OD7000","OL1","Profiler","WTT12 PowerProx",
            "WTT190 PowerProx","WTT2 PowerProx","WTT280 PowerProx","WTT4 PowerProx"
        ]),
        ("Ultrasonic Distance Sensors", ["UC12","UC30","UC4","UC40","UM12","UM18","UM30"]),
    ]),
    ("LiDAR and Radar Sensors", [
        ("LiDAR Sensors", [
            "CL1","LD-LRS","LD-MRS","LD-OEM","LMC1xx","LMS1000","LMS1xx",
            "LMS4000","LMS5xx","LRS4000","MRS1000","MRS6000","multiScan100",
            "NAV2xx","NAV3xx","picoScan100","TiM"
        ]),
        ("Radar Sensors", ["RMS1000","RMS2000"]),
    ]),
    ("Safety (An toan)", [
        ("Safety Light Curtains", [
            "C4-RD","C4000 Advanced","C4000 Advanced ATEX II 3G/3D",
            "C4000 Advanced EX","C4000 Entry/Exit","C4000 Entry/Exit ATEX II 3G/3D",
            "C4000 Fusion","C4000 Fusion ATEX II 3G/3D","C4000 Fusion Ex",
            "C4000 Palletizer","C4000 Palletizer ATEX II 3G/3D","deTec",
            "M4000 Advanced Curtain","M4000 Area","miniTwin","Sense2",
            "TWINOX4","Upgrade Kit FGS to deTec4 Core"
        ]),
        ("Safety Light-Beam Sensors", [
            "deTem","L21","L25","L26","L29","L4000 Systems","L41",
            "M2000 Standard","M4000 Advanced","M4000 Advanced AP",
            "M4000 Standard","M4000 Standard AP","Upgrade Kit MSL to M4000",
            "WSU/WEU26-3"
        ]),
        ("Safety Switches", [
            "E100","ER12","ES11","ES21","flexLock","i10 Lock","i10H","i10P",
            "i10R","i110 Lock","i110P","i110R","i110RP","i110S","i12S",
            "i14 Lock","i150RP","i15 Lock","i16S","i17S","i200 Lock",
            "IME2S","IN4000 Direct","IN4000 Standard","IQB2S","MLP1",
            "RE1","RE2","safeIDS","STR1","TR10 Lock","TR110 Lock","TR4 Direct"
        ]),
        ("Safety Controllers", [
            "Flexi Classic","Flexi Compact","Flexi Gateway","Flexi IO",
            "Flexi Mobile","Flexi Net","Flexi Soft","Flexi Soft Designer",
            "Safety Designer"
        ]),
        ("Safety Laser Scanners", [
            "microScan3","multiScan100-S","nanoScan3","outdoorScan3",
            "S300","S300 Mini","TiM-S"
        ]),
        ("Safe 3D Cameras", ["safeVisionary2"]),
        ("Safe Motion Monitoring and Control", [
            "Flexi Soft Drive Monitor","Speed Monitor","Standstill Monitor"
        ]),
        ("Safe Series Connection", ["Flexi Loop"]),
        ("Safety Multibeam Scanners", ["scanGrid2"]),
        ("Safety Radar Sensors", ["safeRS3"]),
        ("Safety Relays", ["Rely"]),
        ("Safe Motor Feedback Systems", [
            "EDS/EDM35-S","EES/EEM37-S","EKS/EKM36-S","SFS/SFM60-S",
            "SKS/SKM36-S","SRS/SRM50-S","STL70-S","TTK50-S","TTK70-S"
        ]),
        ("Safety Encoders", ["AFS/AFM60S Pro","DFS60S Pro"]),
        ("Safety Distance Sensors", ["DT35-S","WTT12-S"]),
        ("Safety Systems", [
            "AGV Dynamic Weather Assist","End of Arm Safeguard",
            "Safe AGV Easy","Safe Brake Assist","Safe EFI Pro System",
            "Safe Entry Exit","Safe Portal","Safe Robotics Area Protection",
            "Safeguard Detector"
        ]),
    ]),
    ("Machine Vision and Identification", [
        ("Fixed Mount Barcode Scanners", [
            "CLV50x","CLV60x","CLV61x","CLV62x","CLV63x","CLV64x","CLV65x","CLV69x"
        ]),
        ("Image-Based Code Readers", [
            "GLS100","GLS6","ICD8xx","Lector61x","Lector62x","Lector63x",
            "Lector64x/Lector65x","Lector83x","Lector85x"
        ]),
        ("Machine Vision", [
            "Inspector","Inspector83x","Inspector85x",
            "InspectorP Rack Fine Positioning","InspectorP61x","InspectorP62x",
            "midiCam2","picoCam2","Ranger3","Ruler3000","sensingCam SEC100",
            "Trispector1000","TrispectorP1000","Visionary-B Two",
            "Visionary-S","Visionary-T Mini","Visionary-T300"
        ]),
        ("Mobile Handheld Scanners", ["HW19","IDM","Zx36"]),
        ("RFID", ["RFH5xx","RFH6xx","RFU61x","RFU62x","RFU63x","RFU65x"]),
    ]),
    ("Motion Control Sensors", [
        ("Absolute Encoders", [
            "A3M60","ACM60","ACS/ACM36","AFS/AFM60 Ethernet",
            "AFS/AFM60 SSI","AHS/AHM36","ANS/ANM58","ARS60","MAS"
        ]),
        ("Incremental Encoders", ["DBS3650","DBS60","DFS2x","DFS60","DGS80","DLS40","DUS60"]),
        ("Inertial Sensors", ["TMS/TMM22","TMS/TMM61","TMS/TMM88","TMS/TMM88 Dynamic"]),
        ("Linear Encoders", ["DAX","KH53","MAX","OLM","STL/ETL70","TTK50","TTK70"]),
        ("Measuring Wheel Encoders", ["DBV50","DFV60","DKV60","DUV60","MWS075","MWS120"]),
        ("Motor Feedback Systems", [
            "EDS/EDM35","EES/EEM37","EKS/EKM36","ELS/ELM35","SCON",
            "SEK/SEL","SES/SEM","SFS/SFM60","SHUB","SKS/SKM36",
            "SRS/SRM50","STS","VFS60"
        ]),
        ("Non-Contact Motion Sensors", ["SPEETEC 1D"]),
        ("Wire Draw Encoders", ["EcoLine","HighLine","VarioLine"]),
    ]),
    ("Process Sensors (Cam bien qua trinh)", [
        ("Flow Sensors", ["BulkScan","FFU","FTMG","T-EASIC FTS"]),
        ("Level Sensors", [
            "GRF18S","LBV300","LBV301","LFC","LFP Cubic","LFP Inox",
            "LFV200","LXRC","LXRH","UP56"
        ]),
        ("Pressure Sensors", [
            "LFH","PAC50","PBS Hygienic","PBS Plus","PBST","PBT",
            "PET","PFT","PFT-2","PHT","PTA"
        ]),
        ("Temperature Sensors", ["TBS","TBT","TCT","THTE","THTL","THTS","TSP"]),
    ]),
    ("Network and Connection Technology", [
        ("Connectors and Cables", [
            "Field Wireable Connectors","Industrial Ethernet/Fieldbus Cables",
            "Other Connectors","Passive Distribution Boxes",
            "Sensor/Actuator Cable","Y and T Splitters"
        ]),
        ("Displays", ["SID"]),
        ("Edge Computing Devices", [
            "SIM10xx","SIM1200","SIM200","SIM2x00","SIM800","SM1000",
            "Telematic Data Collector"
        ]),
        ("Junction Boxes", ["CDB","CDE50","CDM"]),
        ("Network Devices", ["CDE100","CDF","SIG100","SIG200","SIG300","SIG350"]),
        ("Power Supply", ["Power Supply Cables","Power Supply Units"]),
        ("System Plugs and Extension Modules", [
            "DMM4","System Plugs deTec/deTem","System Plugs Flexi Compact",
            "System Plugs Flexi Soft","System Plugs LMS5xx",
            "System Plugs microScan3","System Plugs miniTwin",
            "System Plugs multiScan/picoScan","System Plugs nanoScan3",
            "System Plugs outdoorScan3","System Plugs S300",
            "System Plugs S3000","UE401","UE402","UE403"
        ]),
    ]),
    ("Accessories (Phu kien)", [
        ("Actuators and Bolts", [
            "Actuators for Electromechanical Safety Switches",
            "Actuators for Magnetic Safety Switches",
            "Actuators for RFID Safety Switches","MB1"
        ]),
        ("Antennas", ["GPS Antennas","GSM/UMTS/LTE Antennas","RFID Antennas","WLAN Antennas"]),
        ("Codes", ["Magnetic Coded Tags","Positioning Codes"]),
        ("Commissioning Aids", ["Alignment Aids","Configuration Devices","Test Equipment"]),
        ("Device Protection and Care", [
            "Cleaning Agents","Cooling Devices","Drying Agent","Front Screens",
            "Heating Devices","Protection Filter","Protective Caps",
            "Protective Housing","Purge Air Equipment"
        ]),
        ("Integration Modules", [
            "Bus Adapters","Cloning Modules","Display Modules","Fieldbus Modules",
            "HIPERFACE Adapters","Interface Modules","Power Supply Modules",
            "SSI Parallel Adapter"
        ]),
        ("Magnets", ["Magnetic Tapes","Position Magnets"]),
        ("Manual Unlocking", ["Locks"]),
        ("Measuring Wheels", ["Measuring Wheel Mechanics","Measuring Wheels"]),
        ("Mounting Systems", [
            "Assembly Accessories","Brackets","Device Columns",
            "Flanges and Nozzles","Muting Arms","Stator Couplings"
        ]),
        ("Optical Data Transmission", ["ISD300","ISD400"]),
        ("Programming Devices", [
            "CPA","PGT-01-S","PGT-08-S","PGT-10-Pro","PGT-11-S",
            "PGT-12-Pro","PGT-13-S","PGT-14","PGT-15"
        ]),
        ("Rechargeable Batteries", ["Base Stations","Rechargeable Batteries"]),
        ("Reflectors and Optics", [
            "Illuminations","Lenses","Mirror","Optical Apertures",
            "Optical Filters","Reflectors"
        ]),
        ("RFID Transponders", ["HF Transponders","UHF Transponders"]),
        ("Shaft Adaptation", ["Collets","Shaft Couplings"]),
        ("Signal Transmitters", ["Acoustic Signal Transmitters","Optical Signal Transmitters"]),
        ("Storage Media", ["Memory Cards","USB Sticks"]),
        ("Wire Draw Mechanism", ["Wire Draw for Rope Pull","Wire Draw for Encoders"]),
    ]),
    ("Digital Services and Software", [
        ("Application Software", [
            "3D Object Detection","Baggage Analytics","Code Loc","Field Analytics",
            "LiDAR Loc","Logistics Diagnostic Analytics","Monitoring Box",
            "Package Analytics","Safety Laser Scanner Visualization",
            "SICK Analytics Assurance","SICK AssetHub","SICK Maritime Suite",
            "SICK Safety Assistant App","Smart Motor Sensors Software",
            "Tire Analytics","Track and Trace Vision Software"
        ]),
        ("Engineering Tools", [
            "Function Block Factory","SENTIO Creator","SICK AppEngine",
            "SICK AppManager","SICK AppStudio","SICK dStudio",
            "SICK STREAM Software","SOPAS ET"
        ]),
        ("Insight Tools", [
            "Asset Analytics","Incoming Goods Suite","Safety Machine Analytics",
            "SICK AR Assistant","Track and Trace Analytics Lite"
        ]),
        ("Integration Software", ["SICK connectX"]),
        ("SICK SensorApps", [
            "3D Belt Pick","Color Inspection and Sorting",
            "HERMES Standard Solution","SICK Nova","Static Package Dimensioning"
        ]),
    ]),
    ("Systems (He thong)", [
        ("Localization Systems", [
            "Automated Load Detect Ident","OccuPID System",
            "Tag-LOC System","TRITON Floor LOC"
        ]),
        ("Object Detection Systems", [
            "AOS LiDAR","AOS Radar","Area Hotspot Detection",
            "Backup Assistance System","Conveyor Hotspot Detection",
            "Drive Assist System","Overheight Detection","Vehicle Hotspot Detection"
        ]),
        ("Profiling Systems", [
            "Axle Classification","Free Flow Profiler","Load Volume Measurement",
            "Multi Lane Profiling","TIC501","TICX02","VPS Pro"
        ]),
        ("Quality Control Systems", [
            "Foreign Object Detection","Label Checker","Pallet Classification",
            "Pallet Integrity Inspection","Pallet Serialization"
        ]),
        ("Robot Guidance Systems", ["Body Position System","PALLOC","PLB","PLOC2D","PLR"]),
        ("Track and Trace Systems", [
            "CLV Identification System","DWS Dynamic","DWS Pallet",
            "ICR Identification System","Ident Gate System",
            "Inspector Logistics System","Lector Identification System",
            "Lector65x System","Master Data Analyzer",
            "Master Data Analyzer Vision","RF Identification System",
            "VML","VMS4x00/5x00","VMS6x00/7x00","VMV Dimensioning System"
        ]),
    ]),
    ("Service (Dich vu)", [
        ("Consulting", ["Accident Investigation","Concept and Feasibility","Safety and Risk Assessment"]),
        ("Engineering and Integration", [
            "Acceptance Services","Extended Warranty","Initial Verification",
            "Installation and Commissioning","Solution Engineering"
        ]),
        ("Maintenance and Inspection", [
            "Electrical Equipment Check","Inspection of Protective Devices",
            "Machine Safety Inspection","Maintenance","Monitoring Services",
            "Performance Check","Service Agreements","Smart Parts",
            "Software/Firmware Maintenance","Stop Time Measurement"
        ]),
        ("Modernization", ["Upgrade and Retrofit"]),
        ("Technical Support", ["Preparation for Reverification","Repairs","Troubleshooting"]),
    ]),
    ("Training (Dao tao)", [
        ("Training and Education", [
            "Product/System/Software","Safety Competence",
            "Specialized and Methodology","Technology Trainings"
        ]),
    ]),
    ("Analyzers (Phan tich)", [
        ("Gas Analyzers", ["SICK Gas Analyzers"]),
    ]),
]

row = 3
main_stt = 0
grand_total = 0

for cat_name, subcats in data:
    main_stt += 1
    total_in_cat = sum(len(fams) for _, fams in subcats)
    grand_total += total_in_cat

    # Category header
    for col in range(1, 5):
        c = ws.cell(row=row, column=col)
        c.fill = cat_fill
        c.font = cat_font
        c.border = thin_border
    ws.cell(row=row, column=1, value=main_stt).alignment = Alignment(horizontal='center')
    ws.cell(row=row, column=2, value=cat_name)
    ws.cell(row=row, column=4, value=total_in_cat).alignment = Alignment(horizontal='center')
    row += 1

    sub_stt = 0
    for subcat_name, families in subcats:
        sub_stt += 1
        stt_str = f"{main_stt}.{sub_stt}"
        for col in range(1, 5):
            c = ws.cell(row=row, column=col)
            c.font = sub_font
            c.border = thin_border
        ws.cell(row=row, column=1, value=stt_str).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=2, value=f"  {subcat_name}")
        ws.cell(row=row, column=4, value=len(families)).alignment = Alignment(horizontal='center')
        row += 1

        for idx, fam in enumerate(families, 1):
            for col in range(1, 5):
                c = ws.cell(row=row, column=col)
                c.font = normal_font
                c.border = thin_border
            ws.cell(row=row, column=1, value=f"{stt_str}.{idx}").alignment = Alignment(horizontal='center')
            ws.cell(row=row, column=2, value=f"      {fam}")
            row += 1

# Total row
row += 1
for col in range(1, 5):
    c = ws.cell(row=row, column=col)
    c.font = Font(name='Arial', size=12, bold=True)
    c.fill = total_fill
    c.border = thin_border
ws.cell(row=row, column=2, value="TONG SO DONG SAN PHAM")
ws.cell(row=row, column=4, value=grand_total).alignment = Alignment(horizontal='center')

# === Sheet 2: Summary ===
ws2 = wb.create_sheet("Tong hop")
ws2.column_dimensions['A'].width = 6
ws2.column_dimensions['B'].width = 50
ws2.column_dimensions['C'].width = 15
ws2.column_dimensions['D'].width = 15

ws2.merge_cells('A1:D1')
c = ws2['A1']
c.value = "TONG HOP THEO DANH MUC CHINH"
c.font = title_font
c.fill = title_fill
c.alignment = Alignment(horizontal='center', vertical='center')

for col, h in enumerate(["STT","Danh muc","So nhom con","So dong SP"], 1):
    c = ws2.cell(row=2, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center')
    c.border = thin_border

r = 3
ts = 0
tf = 0
for i, (cat_name, subcats) in enumerate(data, 1):
    ns = len(subcats)
    nf = sum(len(f) for _, f in subcats)
    ts += ns
    tf += nf
    for col in range(1, 5):
        ws2.cell(row=r, column=col).border = thin_border
    ws2.cell(row=r, column=1, value=i).alignment = Alignment(horizontal='center')
    ws2.cell(row=r, column=2, value=cat_name)
    ws2.cell(row=r, column=3, value=ns).alignment = Alignment(horizontal='center')
    ws2.cell(row=r, column=4, value=nf).alignment = Alignment(horizontal='center')
    r += 1

for col, val in [(1, ""), (2, "TONG CONG"), (3, ts), (4, tf)]:
    c = ws2.cell(row=r, column=col, value=val)
    c.font = Font(name='Arial', size=11, bold=True)
    c.fill = total_fill
    c.border = thin_border
    c.alignment = Alignment(horizontal='center')

filepath = "C:/Users/Admin/Desktop/New folder/SICK_Web_Products_Statistics.xlsx"
wb.save(filepath)
print(f"Saved: {filepath}")
print(f"Total: {grand_total} product families, {ts} subcategories, {len(data)} main categories")
