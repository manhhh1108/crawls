"""
Script sửa lỗi dịch sản phẩm SICK.
Đọc file lỗi, tìm và thay thế giá trị sai trong từng file sản phẩm xlsx.
"""
import openpyxl
import os
import unicodedata
from collections import defaultdict

ERROR_FILE = "loi dich san pham.xlsx"
PATH_PREFIX = "SICK_Products/"  # Error file paths need this prefix added

def load_errors():
    """Load all errors grouped by file path for efficiency."""
    wb = openpyxl.load_workbook(ERROR_FILE)
    ws = wb.active

    errors_by_file = defaultdict(list)
    for row in range(2, ws.max_row + 1):
        error = {
            "family": ws.cell(row, 1).value,
            "sku": ws.cell(row, 2).value,
            "xlsx_path": ws.cell(row, 3).value,
            "field": ws.cell(row, 4).value,
            "error_type": ws.cell(row, 5).value,
            "vi_value": ws.cell(row, 6).value,
            "en_value": ws.cell(row, 7).value,
            "suggested_fix": ws.cell(row, 8).value,
        }
        errors_by_file[error["xlsx_path"]].append(error)

    wb.close()
    return errors_by_file

def normalize(s):
    """Normalize string: replace non-breaking spaces and normalize unicode."""
    if s is None:
        return ""
    return unicodedata.normalize("NFKC", str(s)).strip()

def fix_product_file(xlsx_path, errors):
    """Fix translation errors in a single product xlsx file."""
    full_path = PATH_PREFIX + xlsx_path
    if not os.path.exists(full_path):
        return False, f"File not found: {full_path}"

    wb = openpyxl.load_workbook(full_path)
    ws = wb.active
    fixed_count = 0

    for error in errors:
        field = normalize(error["field"])
        vi_value = normalize(error["vi_value"])
        suggested_fix = error["suggested_fix"]

        if not suggested_fix:
            continue

        # Strategy 1: Find row where English spec name (col 4) matches field,
        # then fix Vietnamese value in col 3
        for r in range(1, ws.max_row + 1):
            en_spec_name = ws.cell(r, 4).value
            vi_val = ws.cell(r, 3).value

            if en_spec_name and normalize(en_spec_name) == field:
                if vi_val and normalize(vi_val) == vi_value:
                    ws.cell(r, 3).value = suggested_fix
                    fixed_count += 1
                    break
        else:
            # Strategy 2: Field found but value didn't match exactly - try partial
            for r in range(1, ws.max_row + 1):
                en_spec_name = ws.cell(r, 4).value
                if en_spec_name and normalize(en_spec_name) == field:
                    vi_val = ws.cell(r, 3).value
                    if vi_val and vi_value in normalize(vi_val):
                        ws.cell(r, 3).value = normalize(vi_val).replace(vi_value, str(suggested_fix))
                        fixed_count += 1
                        break

        # Also fix in short description (row 10, col 2) if it contains the wrong value
        short_desc = ws.cell(10, 2).value
        if short_desc and vi_value and vi_value in normalize(short_desc):
            ws.cell(10, 2).value = normalize(short_desc).replace(vi_value, str(suggested_fix))

    wb.save(full_path)
    wb.close()
    return True, fixed_count

def main():
    print("Loading errors...")
    errors_by_file = load_errors()
    total_files = len(errors_by_file)
    total_errors = sum(len(v) for v in errors_by_file.values())
    print(f"Found {total_errors} errors across {total_files} files")

    fixed_total = 0
    skipped_files = []
    failed_files = []

    for i, (xlsx_path, errors) in enumerate(errors_by_file.items(), 1):
        if i % 100 == 0 or i == total_files:
            print(f"Processing {i}/{total_files}...")

        try:
            success, result = fix_product_file(xlsx_path, errors)
            if success:
                fixed_total += result
            else:
                skipped_files.append(result)
        except Exception as e:
            failed_files.append(f"{xlsx_path}: {e}")

    print(f"\n=== Results ===")
    print(f"Total errors: {total_errors}")
    print(f"Fixed: {fixed_total}")
    print(f"Skipped (file not found): {len(skipped_files)}")
    print(f"Failed: {len(failed_files)}")

    if skipped_files:
        print(f"\nSkipped files (first 10):")
        for s in skipped_files[:10]:
            print(f"  {s}")

    if failed_files:
        print(f"\nFailed files (first 10):")
        for f in failed_files[:10]:
            print(f"  {f}")

if __name__ == "__main__":
    main()
