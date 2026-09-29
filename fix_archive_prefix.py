"""
Fix prefix "Lưu trữ" / "Archive" trong file xlsx CHI cho cac product co
category = "Archive" tren SICK (vd Sales Kit, EOL products).

CACH HOAT DONG (an toan, khong dung global replace):
  1. Detect file co la "Archive" khong: row 8 (Meta title) col B start with "Lưu trữ "
  2. Neu CO, sua DUNG 6 cell:
     - Row 8 col B (Meta title VI): bo prefix "Lưu trữ "
     - Row 9 col B (Meta desc VI):  "Mua Lưu trữ "  -> "Mua "
     - Row 11 col B (Mo ta thong so VI): "Sản phẩm Lưu trữ " -> "Sản phẩm "
     - Row 8 col D (Meta title EN): bo prefix "Archive "
     - Row 9 col D (Meta desc EN):  "Buy genuine Archive " -> "Buy genuine "
     - Row 11 col D (Mo ta thong so EN): "The product Archive " -> "The product "

KHONG dung global replace -> KHONG anh huong cell khac, file khac.

Usage:
  python fix_archive_prefix.py            # apply tat ca
  python fix_archive_prefix.py --dry-run  # preview
  python fix_archive_prefix.py --family AFS_AFM60S_Pro  # 1 dong
"""
import os
import sys
import glob
import time
from openpyxl import load_workbook

OUTPUT_DIR = "SICK_Products"


def fix_file(path, dry_run=False):
    """
    Tra ve so cell sua. 0 neu file khong phai Archive.
    """
    wb = load_workbook(path)
    ws = wb.active

    # Detect Archive: row 8 col B (Meta title VI) start with "Lưu trữ "
    meta_title_vi = ws.cell(row=8, column=2).value or ""
    meta_title_en = ws.cell(row=8, column=4).value or ""
    is_archive_vi = isinstance(meta_title_vi, str) and meta_title_vi.startswith("Lưu trữ ")
    is_archive_en = isinstance(meta_title_en, str) and meta_title_en.startswith("Archive ")

    if not (is_archive_vi or is_archive_en):
        wb.close()
        return 0

    changed = 0

    # === VI side (col B) ===
    if is_archive_vi:
        # Row 8: bo "Lưu trữ " prefix
        v = ws.cell(row=8, column=2).value
        if isinstance(v, str) and v.startswith("Lưu trữ "):
            new_v = v[len("Lưu trữ "):]
            if not dry_run:
                ws.cell(row=8, column=2, value=new_v)
            changed += 1

        # Row 9: "Mua Lưu trữ " -> "Mua "
        v = ws.cell(row=9, column=2).value
        if isinstance(v, str) and "Mua Lưu trữ " in v:
            new_v = v.replace("Mua Lưu trữ ", "Mua ")
            if not dry_run:
                ws.cell(row=9, column=2, value=new_v)
            changed += 1

        # Row 11: "Sản phẩm Lưu trữ " -> "Sản phẩm "
        v = ws.cell(row=11, column=2).value
        if isinstance(v, str) and "Sản phẩm Lưu trữ " in v:
            new_v = v.replace("Sản phẩm Lưu trữ ", "Sản phẩm ")
            if not dry_run:
                ws.cell(row=11, column=2, value=new_v)
            changed += 1

    # === EN side (col D) ===
    if is_archive_en:
        # Row 8
        v = ws.cell(row=8, column=4).value
        if isinstance(v, str) and v.startswith("Archive "):
            new_v = v[len("Archive "):]
            if not dry_run:
                ws.cell(row=8, column=4, value=new_v)
            changed += 1

        # Row 9
        v = ws.cell(row=9, column=4).value
        if isinstance(v, str) and "Buy genuine Archive " in v:
            new_v = v.replace("Buy genuine Archive ", "Buy genuine ")
            if not dry_run:
                ws.cell(row=9, column=4, value=new_v)
            changed += 1

        # Row 11
        v = ws.cell(row=11, column=4).value
        if isinstance(v, str) and "The product Archive " in v:
            new_v = v.replace("The product Archive ", "The product ")
            if not dry_run:
                ws.cell(row=11, column=4, value=new_v)
            changed += 1

    if changed > 0 and not dry_run:
        wb.save(path)
    wb.close()
    return changed


def main():
    dry_run = "--dry-run" in sys.argv
    family_filter = None
    if "--family" in sys.argv:
        idx = sys.argv.index("--family")
        if idx + 1 < len(sys.argv):
            family_filter = sys.argv[idx + 1]

    if family_filter:
        folders = [d for d in os.listdir(OUTPUT_DIR) if family_filter.lower() in d.lower()]
    else:
        folders = sorted(d for d in os.listdir(OUTPUT_DIR) if os.path.isdir(os.path.join(OUTPUT_DIR, d)))

    print(f"Scanning {len(folders)} folders...")
    folder_files = []
    for folder in folders:
        folder_path = os.path.join(OUTPUT_DIR, folder)
        if not os.path.isdir(folder_path):
            continue
        files = [f for f in glob.glob(os.path.join(folder_path, "*.xlsx"))
                 if not os.path.basename(f).startswith("~$")]
        if files:
            folder_files.append((folder, files))
    grand_total = sum(len(fs) for _, fs in folder_files)
    print(f"Total {grand_total} files.\n")

    total_files = 0
    archive_files = 0
    total_cells = 0
    errors = 0
    last_print = time.time()

    for folder, files in folder_files:
        for f in files:
            total_files += 1
            try:
                n = fix_file(f, dry_run=dry_run)
                if n > 0:
                    archive_files += 1
                    total_cells += n
            except Exception as e:
                errors += 1
                print(f"  ERROR {f}: {e}")

            now = time.time()
            if now - last_print > 1 or total_files % 100 == 0:
                pct = total_files * 100 // max(grand_total, 1)
                print(f"  [{pct:3d}%] {total_files:>5}/{grand_total} | archive files fixed: {archive_files}", flush=True)
                last_print = now

    print()
    print("=" * 70)
    print(f"Total scanned:        {total_files}")
    if dry_run:
        print(f"Archive files (would fix): {archive_files}")
        print(f"Cells (would change):      {total_cells}")
    else:
        print(f"Archive files fixed:  {archive_files}")
        print(f"Cells changed:        {total_cells}")
    print(f"Errors:               {errors}")


if __name__ == "__main__":
    main()
