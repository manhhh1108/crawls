# -*- coding: utf-8 -*-
"""Build the consolidated workbook from out/*.json."""
import os, io, re, json, glob, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OLDBOOK = os.path.join(os.path.dirname(HERE), "data", "Danh muc san pham 16 hang.xlsx")
OUT = os.path.join(HERE, "out")
DEST = os.path.join(os.path.dirname(HERE), "data", "Danh muc san pham 16 hang - FULL.xlsx")

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")
LEAF_FILL = PatternFill("solid", fgColor="FFF2CC")
LVL_FILL = {1: PatternFill("solid", fgColor="DDEBF7"), 2: PatternFill("solid", fgColor="EDF4FB")}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def old_tree(sheet):
    """Rebuild the tree of a sheet in the previous workbook (STT dot depth)."""
    import openpyxl as _ox
    rows = [r for r in sheet.iter_rows(values_only=True)][3:]
    cur, out = {}, []
    for r in rows:
        if not r or not r[0]:
            continue
        stt = str(r[0]).strip()
        name = (r[1] or "").strip()
        if not name or not re.match(r"^[\d.]+$", stt):
            continue
        d = stt.count(".") + 1
        cur[d] = name
        out.append(([cur[i] for i in range(1, d + 1) if i in cur], r[2]))
    return [(p, c) for p, c in out if len(p) == (str(p) and len(p))]


def norm(p):
    return tuple(re.sub(r"\s+", " ", str(x)).strip().casefold() for x in p)


# Only where the fresh crawl is demonstrably thinner than the previous file AND
# both use the same language for the category names. Everywhere else the two
# trees are simply different vintages/translations and merging them would
# duplicate rows.
MERGE_FROM_OLD = {"02 Pepperl+Fuchs"}


def merge_old(brand, nodes, oldwb):
    if oldwb is None or brand not in MERGE_FROM_OLD or brand not in oldwb.sheetnames:
        return nodes, 0
    have = {norm(n["path"]) for n in nodes}
    added = 0
    for p, code in old_tree(oldwb[brand]):
        k = norm(p)
        if k in have:
            continue
        if len(p) > 1 and norm(p[:-1]) not in have:
            continue
        nodes.append({"path": list(p), "type": "range", "url": None,
                      "code": code, "src": "old"})
        have.add(k)
        added += 1
    return nodes, added


def tree(nodes):
    """Return rows: (stt, level, path, node, is_leaf) in hierarchical order."""
    by_path = {}
    for n in nodes:
        by_path[tuple(n["path"])] = n
    # make sure every ancestor exists
    for p in list(by_path):
        for i in range(1, len(p)):
            if p[:i] not in by_path:
                by_path[p[:i]] = {"path": list(p[:i]), "type": "category", "url": None, "code": None}
    children = {}
    for p in by_path:
        children.setdefault(p[:-1], []).append(p)
    for k in children:
        children[k].sort(key=lambda x: str(x[-1]).lower())

    rows = []

    def walk(parent, prefix):
        for i, p in enumerate(children.get(parent, []), 1):
            stt = "%s%d" % (prefix, i)
            leaf = not children.get(p)
            rows.append((stt, len(p), p, by_path[p], leaf))
            walk(p, stt + ".")

    walk((), "")
    return rows


def sheet_for(wb, brand, data):
    rows = tree(data["nodes"])
    maxlv = max([r[1] for r in rows] or [1])
    ws = wb.create_sheet(brand[:31])
    ws["A1"] = "DANH MUC SAN PHAM - %s" % brand
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Website: %s" % data.get("site", "")
    ws["A3"] = data.get("note", "")
    ws["A3"].font = Font(italic=True, size=9, color="808080")
    hdr = ["STT", "Ten danh muc (co thut dong)", "Cap"] + \
          ["Cap %d" % i for i in range(1, maxlv + 1)] + \
          ["La cap ngay tren ma SP", "Link", "Ma DM"]
    ws.append([])
    ws.append(hdr)
    hr = ws.max_row
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=hr, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for stt, lv, p, n, leaf in rows:
        name = p[-1]
        line = [stt, ("    " * (lv - 1)) + str(name), lv]
        line += [p[i] if i < len(p) else None for i in range(maxlv)]
        line += ["x" if leaf else "", n.get("url") or "", n.get("code") or "",
                 "file cu" if n.get("src") == "old" else "crawl"]
        ws.append(line)
        r = ws.max_row
        if lv in LVL_FILL:
            for c in range(1, 4):
                ws.cell(row=r, column=c).fill = LVL_FILL[lv]
        if lv <= 2:
            ws.cell(row=r, column=2).font = Font(bold=True)
        if leaf:
            ws.cell(row=r, column=3 + maxlv + 1).fill = LEAF_FILL
    ws.freeze_panes = ws.cell(row=hr + 1, column=4)
    ws.auto_filter.ref = "A%d:%s%d" % (hr, get_column_letter(len(hdr)), ws.max_row)
    widths = [10, 62, 6] + [26] * maxlv + [12, 60, 14, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    counts = {}
    for _s, lv, _p, _n, leaf in rows:
        counts[lv] = counts.get(lv, 0) + 1
    return counts, len(rows), maxlv


def main():
    files = sorted(glob.glob(os.path.join(OUT, "*.json")))
    try:
        import openpyxl
        oldwb = openpyxl.load_workbook(OLDBOOK, read_only=True)
    except Exception as e:
        print("  [warn] khong doc duoc file cu:", e)
        oldwb = None
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("00 TONG HOP")
    stats = []
    for f in files:
        with io.open(f, encoding="utf-8") as fh:
            data = json.load(fh)
        brand = data["brand"]
        data["nodes"], nadd = merge_old(brand, data["nodes"], oldwb)
        if nadd:
            print("  (+%d muc bo sung tu file cu cho %s)" % (nadd, brand))
        counts, total, maxlv = sheet_for(wb, brand, data)
        stats.append((brand, data.get("site", ""), counts, total, maxlv, data.get("note", "")))
        print("  %-28s %5d rows  levels=%s" % (brand, total, dict(sorted(counts.items()))))

    maxlv = max([s[4] for s in stats] or [1])
    summary["A1"] = "TONG HOP DANH MUC SAN PHAM - 16 HANG (crawl den cap ngay tren ma SP)"
    summary["A1"].font = TITLE_FONT
    summary["A2"] = "Ngay crawl: %s" % datetime.date.today().isoformat()
    summary["A2"].font = Font(italic=True, size=9, color="808080")
    hdr = ["STT", "Ten hang", "Website"] + ["DM cap %d" % i for i in range(1, maxlv + 1)] + \
          ["Tong DM", "So cap sau", "Ghi chu"]
    summary.append([])
    summary.append(hdr)
    hr = summary.max_row
    for c in range(1, len(hdr) + 1):
        cell = summary.cell(row=hr, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    tot = [0] * maxlv
    for i, (brand, site, counts, total, mx, note) in enumerate(stats, 1):
        row = [i, brand, site] + [counts.get(l, 0) for l in range(1, maxlv + 1)] + [total, mx, note]
        summary.append(row)
        for l in range(maxlv):
            tot[l] += counts.get(l + 1, 0)
    summary.append([None, "TONG CONG", None] + tot + [sum(tot), None, None])
    for c in range(1, len(hdr) + 1):
        summary.cell(row=summary.max_row, column=c).font = Font(bold=True)
    widths = [6, 30, 40] + [10] * maxlv + [10, 10, 70]
    for i, w in enumerate(widths, 1):
        summary.column_dimensions[get_column_letter(i)].width = w
    summary.freeze_panes = summary.cell(row=hr + 1, column=1)
    wb.save(DEST)
    print("\nSaved:", DEST)


if __name__ == "__main__":
    main()
