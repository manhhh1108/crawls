import requests
import json
import time
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE_URL = "https://www.sick.com/api/fact-finder"
CHANNEL = "enSG-sick"
HEADERS = {
    "Accept": "application/json",
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
}

def search(filters, page=1, hits_per_page=100):
    """Search FactFinder API"""
    params = {
        "query": "*",
        "page": page,
        "hitsPerPage": hits_per_page,
        "showPermutedSearchParams": "false"
    }
    for f in filters:
        params.setdefault("filter", [])

    url = f"{BASE_URL}/search/{CHANNEL}"
    filter_params = "&".join([f"filter={f}" for f in filters])
    full_url = f"{url}?query=*&{filter_params}&page={page}&hitsPerPage={hits_per_page}&showPermutedSearchParams=false"

    resp = requests.get(full_url, headers=HEADERS, timeout=30)
    return resp.json()

def get_children(parent_code):
    """Get children of a category/master product"""
    data = search([f"Parent:{parent_code}"], hits_per_page=200)
    total = data.get("totalHits", 0)
    hits = data.get("hits", [])
    results = []
    for hit in hits:
        mv = hit.get("masterValues", {})
        results.append({
            "code": mv.get("DefArticleNo", ""),
            "name": mv.get("DefName", ""),
            "type": mv.get("DefType", ""),
            "url": mv.get("DefUrl", ""),
        })
    return total, results

def count_products(master_code):
    """Count actual product variants under a master product"""
    data = search([f"Parent:{master_code}"], hits_per_page=1)
    return data.get("totalHits", 0)

def get_record_detail(code):
    """Get detail info for a record"""
    url = f"{BASE_URL}/records/detail/{CHANNEL}/{code}?usePersonalization=false&withCampaigns=false&withRecommendations=false&withSimilarProducts=false"
    resp = requests.get(url, headers=HEADERS, timeout=30)
    return resp.json()

print("=" * 70)
print("  CRAWLING SICK.COM PRODUCT DATA")
print("=" * 70)

# Step 1: Get top-level product categories
# Root: g568268 (Products)
print("\nStep 1: Getting top-level categories...")
_, top_categories = get_children("g568268")
print(f"Found {len(top_categories)} top-level categories")
time.sleep(0.5)

# Step 2: For each top category, get subcategories, then product families
all_data = []

for cat in top_categories:
    cat_code = cat["code"]
    cat_name = cat["name"]
    cat_type = cat["type"]

    print(f"\n--- {cat_name} ({cat_code}, type={cat_type}) ---")

    # Get subcategories
    _, subcats = get_children(cat_code)
    time.sleep(0.3)

    cat_entry = {
        "name": cat_name,
        "code": cat_code,
        "subcategories": []
    }

    for sub in subcats:
        sub_code = sub["code"]
        sub_name = sub["name"]
        sub_type = sub["type"]

        print(f"  {sub_name} ({sub_code}, type={sub_type})")

        if sub_type == "Category":
            # This is a subcategory, get its children (product families/master products)
            _, families = get_children(sub_code)
            time.sleep(0.3)

            sub_entry = {
                "name": sub_name,
                "code": sub_code,
                "families": []
            }

            for fam in families:
                fam_code = fam["code"]
                fam_name = fam["name"]
                fam_type = fam["type"]

                if fam_type == "MasterProduct":
                    # Count product variants
                    prod_count = count_products(fam_code)
                    time.sleep(0.2)
                    print(f"    {fam_name}: {prod_count} products")
                    sub_entry["families"].append({
                        "name": fam_name,
                        "code": fam_code,
                        "product_count": prod_count
                    })
                elif fam_type == "Category":
                    # Another level of subcategory - get its children
                    _, sub_families = get_children(fam_code)
                    time.sleep(0.3)

                    fam_sub_entry = {
                        "name": fam_name,
                        "code": fam_code,
                        "families": []
                    }

                    for sf in sub_families:
                        if sf["type"] == "MasterProduct":
                            pc = count_products(sf["code"])
                            time.sleep(0.2)
                            print(f"    {fam_name} > {sf['name']}: {pc} products")
                            fam_sub_entry["families"].append({
                                "name": sf["name"],
                                "code": sf["code"],
                                "product_count": pc
                            })
                        else:
                            # Even deeper - just count
                            pc2 = count_products(sf["code"])
                            time.sleep(0.2)
                            print(f"    {fam_name} > {sf['name']} ({sf['type']}): {pc2} products")
                            fam_sub_entry["families"].append({
                                "name": sf["name"],
                                "code": sf["code"],
                                "product_count": pc2
                            })

                    if fam_sub_entry["families"]:
                        sub_entry["families"].extend(fam_sub_entry["families"])
                    else:
                        pc = count_products(fam_code)
                        time.sleep(0.2)
                        print(f"    {fam_name}: {pc} products (leaf category)")
                        sub_entry["families"].append({
                            "name": fam_name,
                            "code": fam_code,
                            "product_count": pc
                        })
                else:
                    # Unknown type, just count
                    pc = count_products(fam_code)
                    time.sleep(0.2)
                    print(f"    {fam_name} ({fam_type}): {pc} products")
                    sub_entry["families"].append({
                        "name": fam_name,
                        "code": fam_code,
                        "product_count": pc
                    })

            cat_entry["subcategories"].append(sub_entry)

        elif sub_type == "MasterProduct":
            # Direct master product under category (no subcategory level)
            prod_count = count_products(sub_code)
            time.sleep(0.2)
            print(f"    -> {prod_count} products")

            # Add as a single-family subcategory
            if not cat_entry["subcategories"] or cat_entry["subcategories"][-1]["name"] != cat_name:
                cat_entry["subcategories"].append({
                    "name": cat_name,
                    "code": cat_code,
                    "families": []
                })
            cat_entry["subcategories"][-1]["families"].append({
                "name": sub_name,
                "code": sub_code,
                "product_count": prod_count
            })

    all_data.append(cat_entry)

# Save raw data as JSON
with open("C:/Users/Admin/Desktop/New folder/sick_products_raw.json", "w", encoding="utf-8") as f:
    json.dump(all_data, f, ensure_ascii=False, indent=2)

print("\n\n" + "=" * 70)
print("  CREATING EXCEL FILE")
print("=" * 70)

# Create Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Thong ke SP SICK tren Web"

title_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
cat_font = Font(name='Arial', size=11, bold=True)
sub_font = Font(name='Arial', size=10, bold=True, color='2E75B6')
normal_font = Font(name='Arial', size=10)

title_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
cat_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
total_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')

thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 55
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 22

# Title
ws.merge_cells('A1:D1')
c = ws['A1']
c.value = 'THONG KE SAN PHAM TREN WEB SICK.COM/SG/EN'
c.font = title_font
c.fill = title_fill
c.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 35

# Headers
for col, h in enumerate(['STT', 'Dong san pham', 'Category Code', 'So luong san pham'], 1):
    c = ws.cell(row=2, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border

row = 3
main_stt = 0
grand_total = 0

for cat_data in all_data:
    main_stt += 1
    cat_name = cat_data["name"]
    cat_total = sum(
        sum(f["product_count"] for f in sub["families"])
        for sub in cat_data["subcategories"]
    )
    grand_total += cat_total

    # Category header row
    for col in range(1, 5):
        c = ws.cell(row=row, column=col)
        c.fill = cat_fill
        c.font = cat_font
        c.border = thin_border
    ws.cell(row=row, column=1, value=main_stt).alignment = Alignment(horizontal='center')
    ws.cell(row=row, column=2, value=cat_name)
    ws.cell(row=row, column=3, value=cat_data["code"])
    ws.cell(row=row, column=4, value=cat_total).alignment = Alignment(horizontal='center')
    row += 1

    sub_stt = 0
    for sub_data in cat_data["subcategories"]:
        sub_stt += 1
        sub_total = sum(f["product_count"] for f in sub_data["families"])
        stt_str = f"{main_stt}.{sub_stt}"

        for col in range(1, 5):
            c = ws.cell(row=row, column=col)
            c.font = sub_font
            c.border = thin_border
        ws.cell(row=row, column=1, value=stt_str).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=2, value=f"  {sub_data['name']}")
        ws.cell(row=row, column=3, value=sub_data["code"])
        ws.cell(row=row, column=4, value=sub_total).alignment = Alignment(horizontal='center')
        row += 1

        for idx, fam in enumerate(sub_data["families"], 1):
            for col in range(1, 5):
                c = ws.cell(row=row, column=col)
                c.font = normal_font
                c.border = thin_border
            ws.cell(row=row, column=1, value=f"{stt_str}.{idx}").alignment = Alignment(horizontal='center')
            ws.cell(row=row, column=2, value=f"      {fam['name']}")
            ws.cell(row=row, column=3, value=fam["code"])
            ws.cell(row=row, column=4, value=fam["product_count"]).alignment = Alignment(horizontal='center')
            row += 1

# Total row
row += 1
for col in range(1, 5):
    c = ws.cell(row=row, column=col)
    c.font = Font(name='Arial', size=12, bold=True)
    c.fill = total_fill
    c.border = thin_border
ws.cell(row=row, column=2, value="TONG SO SAN PHAM")
ws.cell(row=row, column=4, value=grand_total).alignment = Alignment(horizontal='center')

# Summary sheet
ws2 = wb.create_sheet("Tong hop")
ws2.column_dimensions['A'].width = 6
ws2.column_dimensions['B'].width = 50
ws2.column_dimensions['C'].width = 15
ws2.column_dimensions['D'].width = 18

ws2.merge_cells('A1:D1')
c = ws2['A1']
c.value = "TONG HOP THEO DANH MUC CHINH"
c.font = title_font
c.fill = title_fill
c.alignment = Alignment(horizontal='center', vertical='center')

for col, h in enumerate(["STT", "Danh muc", "So dong SP", "Tong san pham"], 1):
    c = ws2.cell(row=2, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center')
    c.border = thin_border

r = 3
total_lines = 0
total_products = 0
for i, cat_data in enumerate(all_data, 1):
    n_lines = sum(len(sub["families"]) for sub in cat_data["subcategories"])
    n_products = sum(sum(f["product_count"] for f in sub["families"]) for sub in cat_data["subcategories"])
    total_lines += n_lines
    total_products += n_products
    for col in range(1, 5):
        ws2.cell(row=r, column=col).border = thin_border
    ws2.cell(row=r, column=1, value=i).alignment = Alignment(horizontal='center')
    ws2.cell(row=r, column=2, value=cat_data["name"])
    ws2.cell(row=r, column=3, value=n_lines).alignment = Alignment(horizontal='center')
    ws2.cell(row=r, column=4, value=n_products).alignment = Alignment(horizontal='center')
    r += 1

for col, val in [(1, ""), (2, "TONG CONG"), (3, total_lines), (4, total_products)]:
    c = ws2.cell(row=r, column=col, value=val)
    c.font = Font(name='Arial', size=11, bold=True)
    c.fill = total_fill
    c.border = thin_border
    c.alignment = Alignment(horizontal='center')

filepath = "C:/Users/Admin/Desktop/New folder/SICK_Web_Products_Statistics.xlsx"
wb.save(filepath)
print(f"\nSaved: {filepath}")
print(f"Total: {grand_total} products across {total_lines} product lines in {len(all_data)} categories")
