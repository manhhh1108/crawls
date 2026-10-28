# -*- coding: utf-8 -*-
"""Mitsubishi Electric - lap ban do 'nhanh danh muc' -> URL trang san pham.

    python mitsu_map.py

mitsubishielectric.com/fa khong co API danh muc: cap 4 trong file danh muc la
nhan cua the <a class="c-linkWithImage"> tren trang 'Product List' cua tung
dong (vd 'CPU module', 'Base unit'). -> Bo het trang trong /fa/products/, gom
(nhan the, H1, URL) roi doi chieu theo ten.

Ket qua: out/07 Mitsubishi Electric/urlmap.json
"""
import io
import json
import os
import re
import sys
import html as H
import collections

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, REPO)

from shared.fetcher import try_fetch                  # noqa: E402

SITE = "https://www.mitsubishielectric.com"
ROOT = "/fa/products/"
OUT = os.path.join(HERE, "out", "07 Mitsubishi Electric")
NTREND_XLSX = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
SHEET = "07 Mitsubishi Electric"
MAX_PAGES = 1500

LINK = re.compile(r'<a class="c-linkWithImage"[^>]*href="([^"]+)"(.*?)</a>', re.S)
LABEL = re.compile(r'c-linkWithImage__link[^>]*>(.*?)</span>', re.S)
HREF = re.compile(r'href="([^"?#]+)"')
H1 = re.compile(r"<h1[^>]*>(.*?)</h1>", re.S)


def txt(s):
    return H.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", s or ""))).strip()


def norm(s):
    """Ten danh muc -> khoa so sanh (bo het ky tu khong phai chu/so)."""
    return re.sub(r"[^0-9a-z]+", "", (s or "").lower())


def crawl():
    """BFS trong /fa/products/, tra ve {khoa ten: [url, ...]}."""
    seen, queue = set(), collections.deque([ROOT + "index.html"])
    names = collections.defaultdict(list)
    n = 0
    while queue and n < MAX_PAGES:
        path = queue.popleft()
        if path in seen:
            continue
        seen.add(path)
        html = try_fetch(SITE + path, retries=2, timeout=45)
        n += 1
        if not html:
            continue
        m = H1.search(html)
        if m:
            names[norm(txt(m.group(1)))].append(path)
        for href, blk in LINK.findall(html):
            lb = LABEL.search(blk)
            if lb:
                names[norm(txt(lb.group(1)))].append(href.split("#")[0])
        for href in HREF.findall(html):
            if href.startswith(ROOT) and href.endswith((".html", "/")):
                if href not in seen:
                    queue.append(href)
        if n % 50 == 0:
            print("   da doc %d trang, hang doi %d" % (n, len(queue)))
    return names, n


BRAND_VI = "Mitsubishi Electric"


def cat_of(v):
    """Cot danh muc Ntrend chi ghi moi ten hang = chua phan loai."""
    c = (v or "").strip()
    if c == BRAND_VI:
        return ("Không có danh mục sản phẩm,  chỉ nằm trong danh mục "
                "Thương Hiệu " + BRAND_VI)
    return c


def wanted():
    """345 dong SP (cot 'La cap ngay tren ma SP' = x)."""
    from openpyxl import load_workbook
    ws = load_workbook(NTREND_XLSX, read_only=True, data_only=True)[SHEET]
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or len(r) < 9 or str(r[7] or "").strip() != "x":
            continue
        path = [str(x).strip() for x in r[3:7] if x and str(x).strip()]
        if not path:
            continue
        out.append({"stt": str(r[0]).strip(), "name": path[-1], "path": path,
                    "cat_vi": cat_of(r[8])})
    return out


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    os.makedirs(OUT, exist_ok=True)
    names, npage = crawl()
    print("da doc %d trang | %d ten khac nhau" % (npage, len(names)))

    lines = wanted()
    keys = list(names)
    hit = 0
    for l in lines:
        # ten trong danh muc doi khi co tien to "--> " hoac bi cat ngan so voi
        # nhan tren web -> doi chieu them theo tien to / chua trong nhau.
        k = norm(l["name"])
        cand = names.get(k) or []
        if not cand and len(k) >= 6:
            near = [kk for kk in keys if kk.startswith(k) or k.startswith(kk)]
            if len(near) == 1:
                cand = names[near[0]]
            elif not near:
                near = [kk for kk in keys if len(kk) >= 6 and (k in kk or kk in k)]
                if len(near) == 1:
                    cand = names[near[0]]
        # uu tien URL nam trong nhanh cha (ten cap 2/3 xuat hien trong duong dan)
        if len(cand) > 1:
            par = [norm(p) for p in l["path"][:-1]]
            best = [u for u in cand
                    if any(k[:6] in u.lower() for k in par if len(k) > 5)]
            cand = best or cand
        if cand:
            l["url"] = SITE + collections.Counter(cand).most_common(1)[0][0]
            hit += 1
    with io.open(os.path.join(OUT, "urlmap.json"), "w", encoding="utf-8") as f:
        json.dump({"lines": lines}, f, ensure_ascii=False)
    print("dong SP: %d | khop: %d | chua khop: %d"
          % (len(lines), hit, len(lines) - hit))
    for l in lines:
        if not l.get("url"):
            print("   [?]", " > ".join(l["path"]))


if __name__ == "__main__":
    main()
