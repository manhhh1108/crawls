# -*- coding: utf-8 -*-
"""Pepperl+Fuchs - crawl noi dung bai dang cho tung dong san pham.

    python c_pepperl.py            # chay tiep nhung dong chua co
    python c_pepperl.py 5          # chi 5 dong dau (xem thu)
    python c_pepperl.py --force    # lam lai tu dau

Can chay truoc: python pf_classmap.py   (map ten dong SP -> trang ?class=)

Khac SICK: pepperl-fuchs.com khong co API va chan curl -> moi trang deu phai
render bang Chrome that (Playwright). Trang chi hien 50 ma dau, phai dong banner
cookie roi bam 'Load more' den het.

Ket qua: out/02 Pepperl+Fuchs/content.json  -> build_pepperl_excel.py dung tiep.
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, REPO)

from c_pepperl_try import (BRAND, UA, OUT_DIR as XLSX_DIR, CARD, TOPLINE,  # noqa: E402
                           HEADLINE, DESC_FULL, HREF, IMG, SITE, txt,
                           accept_cookies, load_all, parse_items, long_desc,
                           short_desc, slugify, ntrend_row)

OUT = os.path.join(HERE, "out", "02 Pepperl+Fuchs")
CONTENT = os.path.join(OUT, "content.json")
# map boc tu payload bo loc trong cache HTML - co ma class theo TUNG NHANH
# (xem pf_map4.py). Khoa dang "ten cha||ten dong SP".
MAP_JSON = os.path.join(OUT, "classmap2.json")
SRC_IMG = os.path.join(OUT, "src_img")
IMG_DIR = os.path.join(XLSX_DIR, "images")
for d in (OUT, SRC_IMG, IMG_DIR):
    os.makedirs(d, exist_ok=True)


def clean_class_url(u):
    """Bo filter tim kiem dinh kem, chi giu ?class=<id>.

    Mot so the danh muc tro den URL da loc san (...?class=30734&q=nsb10*-s2d2),
    vao do chi thay vai ma SP chu khong phai ca dong.
    """
    u = (u or "").replace("&amp;", "&")
    m = re.search(r"[?&]class=(\d+)", u)
    if not m:
        return u
    return "%s/en-sg/products-gp25581?class=%s" % (SITE, m.group(1))


def wanted():
    """309 dong SP + ten tieng Viet + danh muc Ntrend, theo thu tu trong file."""
    from openpyxl import load_workbook
    from c_pepperl_try import NTREND_XLSX
    ws = load_workbook(NTREND_XLSX, read_only=True, data_only=True)["02 Pepperl+Fuchs"]
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or len(r) < 10 or str(r[8] or "").strip() != "x":
            continue
        name = (r[1] or "").strip()
        if not name:
            continue
        cat = (r[9] or "").strip()
        if cat == BRAND:
            cat = ("Không có danh mục sản phẩm,  chỉ nằm trong danh mục "
                   "Thương Hiệu Pepperl+Fuchs")
        out.append({"stt": str(r[0]).strip(), "name": name,
                    "name_vi": (r[2] or "").strip(), "cat_vi": cat,
                    "path": [str(x).strip() for x in (r[4], r[5], r[6], r[7])
                             if x and str(x).strip()]})
    return out


CARDHREF = re.compile(r'<a class="action-card[^"]*"[^>]*href="([^"]+)"(.*?)</a>', re.S)
CARDHEAD = re.compile(r'action-card__headline[^>]*>(.*?)<', re.S)


def resolve_list_url(pg, url, line_en):
    """Trang /products/<...> la trang DANH MUC, khong co ma SP nao.

    86/203 dong map ra trang danh muc -> mo ra, tim the con co ?class=:
      - uu tien the trung ten dong SP
      - khong co thi lay the dau tien (trang danh muc cua chinh dong do)
    """
    if "class=" in url:
        return url
    pg.goto(url, wait_until="networkidle", timeout=90000)
    pg.wait_for_timeout(1500)
    html = pg.content()
    cards = []
    for href, blk in CARDHREF.findall(html):
        if "class=" not in href:
            continue
        m = CARDHEAD.search(blk)
        t = txt(m.group(1)) if m else ""
        u = href if href.startswith("http") else SITE + href
        cards.append((t, clean_class_url(u)))
    if not cards:
        return url
    for t, u in cards:
        if t.lower() == line_en.lower():
            return u
    return cards[0][1]


def fetch_line(pg, url, line_en=""):
    """Mot trang ?class= -> danh sach ma SP (da bam het 'Load more')."""
    url = resolve_list_url(pg, url, line_en)
    pg.goto(url, wait_until="networkidle", timeout=90000)
    pg.wait_for_timeout(2000)
    load_all(pg)
    return parse_items(pg.content()), url


def grab_image(items, line_en):
    """Anh dai dien = anh ma SP dau tien, ghep nen + logo NTREND."""
    import urllib.request
    url = next((it.get("img") for it in items if it.get("img")), "")
    if not url:
        return ""
    base = slugify(line_en)
    ext = os.path.splitext(url.split("?")[0])[1][:5] or ".png"
    src = os.path.join(SRC_IMG, base + ext)
    try:
        if not (os.path.exists(src) and os.path.getsize(src) > 500):
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            data = urllib.request.urlopen(req, timeout=60).read()
            if len(data) < 500:
                return ""
            with open(src, "wb") as f:
                f.write(data)
        import brandimg
        return os.path.basename(
            brandimg.compose(src, os.path.join(IMG_DIR, base + ".jpg")))
    except Exception as e:
        print("     [anh loi]", repr(e)[:60])
        return ""


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    force = "--force" in sys.argv
    limit = next((int(a) for a in sys.argv[1:] if a.isdigit()), 0)

    lines = wanted()
    if limit:
        lines = lines[:limit]
    cmap = {}
    if os.path.exists(MAP_JSON):
        cmap = json.load(io.open(MAP_JSON, encoding="utf-8"))
    done = {}
    if os.path.exists(CONTENT) and not force:
        # khoa theo URL chu KHONG theo ten: 16 nhom trung ten nhung khac trang
        # (vd 'Accessories' o 16 nhanh) - khoa theo ten se de len nhau, mat bai.
        done = {a.get("url_src"): a for a in
                json.load(io.open(CONTENT, encoding="utf-8"))["articles"]
                if a.get("url_src")}

    # Gan trang cho tung dong (khop CA NHANH CHA) va bo trung THEO TRANG:
    # nhieu dong cung tro ve mot trang thi noi dung y het nhau -> chi lam 1 bai.
    # Cac dong trung ten nhung KHAC trang van giu rieng (vd 18 dong 'Accessories'
    # thuoc 16 nhanh -> 16 trang khac nhau -> 16 bai).
    import pf_map2
    seen_url, nomap, uniq = {}, [], []
    for l in lines:
        hit, _how = pf_map2.lookup(cmap, l)
        if not hit:
            nomap.append(l["name"])
            continue
        u = clean_class_url(hit["url"])
        if u in seen_url:
            seen_url[u].append(l["name"])
            continue
        seen_url[u] = []
        uniq.append(dict(l, url=u))
    ndup = sum(len(v) for v in seen_url.values())
    todo = [l for l in uniq if l["url"] not in done]
    print("dong SP: %d | bo trung cung trang: %d | bai se co: %d"
          % (len(lines), ndup, len(uniq)))
    print("da xong: %d | se lam: %d | chua co trang: %d"
          % (len(done), len(todo), len(nomap)))
    if nomap:
        print("  (chua co trang -> chay lai pf_classmap.py):", ", ".join(nomap[:6]),
              "..." if len(nomap) > 6 else "")
    if not todo:
        return

    from playwright.sync_api import sync_playwright
    from vi_engine import translate_many
    from translate import _flush

    with sync_playwright() as pw:
        b = pw.chromium.launch(headless=False, channel="chrome")
        pg = b.new_page()
        first = True
        for i, ln in enumerate(todo, 1):
            url = ln["url"]
            try:
                items, url = fetch_line(pg, url, ln["name"])
            except Exception as e:
                print("  %3d/%d %-34s LOI %s" % (i, len(todo), ln["name"][:34],
                                                 str(e)[:40]))
                continue
            if first:
                accept_cookies(pg)
                first = False
                try:
                    items, url = fetch_line(pg, url, ln["name"])
                except Exception:
                    pass
            long_en = long_desc(items, ln["name"])
            short_en = short_desc(items, ln["name"])
            img = grab_image(items, ln["name"])

            # Review 22/10: ten SP trong bang ma cung phai co ban tieng Viet
            n_en = sorted({it.get("name", "") for it in items if it.get("name")})
            n_vi = dict(zip(n_en, translate_many(n_en, "key"))) if n_en else {}
            for it in items:
                it["name_vi"] = n_vi.get(it.get("name", ""), "")
            d_en = [it["desc"] for it in items]
            long_vi = translate_many(long_en, "value") if long_en else []
            d_vi = translate_many(d_en, "value") if d_en else []
            for it, v in zip(items, d_vi):
                it["desc_vi"] = v

            done[url] = dict(ln, url_src=url, items=items, img_file=img,
                                    long_en=long_en, long_vi=long_vi,
                                    short_en=short_en,
                                    short_vi=long_vi[:len(short_en)])
            print("  %3d/%d %-34s ma=%-4d ts=%-3d %s"
                  % (i, len(todo), ln["name"][:34], len(items), len(long_en),
                     "img" if img else "---"))
            if i % 5 == 0:
                _flush(True)
                save(uniq, done)
        b.close()
    _flush(True)
    save(uniq, done)
    print("\nxong %d bai -> %s" % (len(done), CONTENT))


def save(lines, done):
    """Ghi content.json. CHI nhan 'uniq' (moi muc co khoa 'url') - truyen nham
    'lines' thi khong muc nao co 'url', file bi ghi de thanh rong (da gap 20/10).
    """
    arts, seen = [], set()
    if lines and not any(l.get("url") for l in lines):
        raise ValueError("save(): danh sach khong co 'url' - sai tham so")
    for l in lines:
        u = l.get("url")
        if u in done and u not in seen:
            seen.add(u)
            arts.append(done[u])
    with io.open(CONTENT, "w", encoding="utf-8") as f:
        json.dump({"brand": "02 Pepperl+Fuchs", "articles": arts},
                  f, ensure_ascii=False)


if __name__ == "__main__":
    main()
