# -*- coding: utf-8 -*-
"""SICK - PHA 0: dung danh sach dong SP can crawl tu danh muc moi nhat.

Dau vao:
  data/Danh muc san pham 16 hang - FULL.xlsx, sheet "01 SICK"
      cot H  'La cap ngay tren ma SP' = x  -> dong san pham
      o BOI VANG (FFFF00)                  -> nhom KHONG ban, bo qua
  data/SICK_Web_Products_Statistics.xlsx, sheet "Thong ke SP SICK Web SP"
      lay lai 'Danh muc SP Tieng Viet/Anh' theo Category Code (gcode)

Dau ra:
  data/SICK - dong SP can crawl.xlsx  (dung format ma c_sick_fetch.py doc)
"""
import io
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)

from openpyxl import Workbook, load_workbook              # noqa: E402
from openpyxl.styles import Font, PatternFill            # noqa: E402
from openpyxl.utils import get_column_letter             # noqa: E402

CATALOG = os.path.join(REPO, "data", "Danh muc san pham 16 hang - FULL.xlsx")
CAT_SHEET = "01 SICK"
STATS = os.path.join(REPO, "data", "SICK_Web_Products_Statistics.xlsx")
STATS_SHEET = "Thong ke SP SICK Web SP"
OUT = os.path.join(REPO, "data", "SICK - dong SP can crawl.xlsx")

YELLOW = "FFFFFF00"

# Nhom 4 'Digital services and software' - software/tools, khong ban
SKIP_PATH = ("Digital services and software",)

# Software nam lan trong nhom khac
SKIP_NAME = ("Flexi Soft Designer", "Safety Designer", "Triton Floor-LOC")

PLACEHOLDER_FULL = ("Không có danh mục sản phẩm,  chỉ nằm trong danh mục "
                    "Thương Hiệu SICK Sensor")

ANALYZER_VI = "Thiết bị phân tích khí và bụi"
ANALYZER_EN = "Gas and Dust Analyzers"

# Dong SP moi (co trong danh muc 12/10 nhung file thong ke cu chua co)
# -> map tay vao danh muc Ntrend dang dung. gcode: (VI, EN)
CAT_FIX = {
    "g555781": ("Cảm biến công nghiệp, Cảm biến tiệm cận",     # IMX
                "Industrial Sensors, Proximity Sensors"),
    "g184875": ("Cảm biến công nghiệp, Cảm biến quang điện",   # W100
                "Industrial Sensors, Photoelectric Sensors"),
    "g203757": ("Cảm biến công nghiệp, Cảm biến quang điện",   # W250
                "Industrial Sensors, Photoelectric Sensors"),
    "g204129": ("Cảm biến công nghiệp, Cảm biến quang điện",   # W280
                "Industrial Sensors, Photoelectric Sensors"),
    "g543722": ("Cảm biến công nghiệp, Cảm biến lưu lượng",    # Digital flow solutions
                "Industrial Sensors, Flow Sensors"),
    "g543714": ("Cảm biến công nghiệp, Cảm biến lưu lượng",    # Flow measurement instruments
                "Industrial Sensors, Flow Sensors"),
    "g472151": ("Cảm biến công nghiệp, Cảm biến lưu lượng",    # Flow metering systems
                "Industrial Sensors, Flow Sensors"),
    "g543715": ("Cảm biến công nghiệp, Cảm biến lưu lượng",    # Flow meters
                "Industrial Sensors, Flow Sensors"),
    "g625244": ("Vision và Nhận dạng, Máy đọc mã vạch",        # sensingCam SEC200
                "Vision & Identification, Barcode Scanners"),
    "g623175": ("Encoder bộ mã hóa vòng quay",                 # ENS/ENM35
                "Rotary Encoders"),
    "g628391": ("Bộ kết nối mạng công nghiệp, Đầu nối và cáp",  # Power cables
                "Industrial Network Connectivity, Connectors & Cables"),
    # 2 danh muc MOI tren web Ntrend (chot 13/10)
    "g54724": (ANALYZER_VI, ANALYZER_EN),    # Gas Analyzers
    "g54727": (ANALYZER_VI, ANALYZER_EN),    # Dust Measuring Devices
    "g284260": (ANALYZER_VI, ANALYZER_EN),   # Smoke detectors
    "g407184": (ANALYZER_VI, ANALYZER_EN),   # Air quality measuring devices
    "g284262": (ANALYZER_VI, ANALYZER_EN),   # Visual range measuring devices
    "g284256": (ANALYZER_VI, ANALYZER_EN),   # Digital analyzer solutions
    "g628337": ("Cảm biến công nghiệp, Cảm biến radar",          # RMS
                "Industrial Sensors, Radar Sensors"),
    # khong thuoc danh muc ban hang nao -> dung chung cau nhu 98 dong Accessories/Systems
    "g575324": (PLACEHOLDER_FULL, PLACEHOLDER_FULL),            # Other products
    "g624914": (PLACEHOLDER_FULL, PLACEHOLDER_FULL),            # Pallet Serialization System
}


def is_yellow(cell):
    f = cell.fill
    return bool(f and f.patternType == "solid" and f.fgColor
                and f.fgColor.rgb == YELLOW)


PLACEHOLDER = "Không có danh mục sản phẩm"


def cat_map():
    """gcode -> (danh muc VI, danh muc EN) lay tu file thong ke cu.

    File thong ke co 26 ma bi lap 2 ban: 1 ban co danh muc Ntrend that, 1 ban
    ghi 'Khong co danh muc san pham...'. Luon uu tien ban co danh muc that.
    """
    ws = load_workbook(STATS, data_only=True, read_only=True)[STATS_SHEET]
    out = {}
    for r in ws.iter_rows(values_only=True):
        if len(r) < 7 or str(r[6]).strip().lower() != "x" or not r[2]:
            continue
        code = str(r[2]).strip()
        vi, en = (r[4] or "").strip(), (r[5] or "").strip()
        real = bool(vi) and not vi.startswith(PLACEHOLDER)
        old = out.get(code)
        if old is None or (real and not old[2]):
            out[code] = (vi, en, real)
    return {k: (v[0], v[1]) for k, v in out.items()}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    cats = cat_map()
    ws = load_workbook(CATALOG)[CAT_SHEET]

    keep, skipped = [], []
    for row in ws.iter_rows(min_row=6):
        v = [c.value for c in row]
        if len(v) < 10 or v[7] != "x":
            continue
        name = (v[1] or "").strip()
        path = [p for p in (v[3], v[4], v[5]) if p]
        if any(is_yellow(c) for c in row):
            skipped.append((name, "boi vang"))
        elif path and path[0] in SKIP_PATH:
            skipped.append((name, "nhom software/tools"))
        elif name in SKIP_NAME:
            skipped.append((name, "software le"))
        else:
            code = (v[9] or "").strip()
            cvi, cen = CAT_FIX.get(code) or cats.get(code, ("", ""))
            keep.append({"stt": str(v[0]).strip(), "name": name, "code": code,
                         "n": 0, "cat_vi": cvi, "cat_en": cen,
                         "path": " > ".join(path)})

    wb = Workbook()
    w = wb.active
    w.title = STATS_SHEET
    w["A1"] = "DONG SAN PHAM SICK CAN CRAWL NOI DUNG"
    w["A1"].font = Font(bold=True, size=13, color="1F4E79")
    w.append(["STT", "Dòng sản phẩm", "Category Code", "So luong san pham",
              "Danh mục SP Tiếng Việt", "Danh mục SP tiếng Anh", "x", "Nhánh danh mục"])
    hdr = PatternFill("solid", fgColor="DDEBF7")
    for c in range(1, 9):
        w.cell(row=2, column=c).fill = hdr
        w.cell(row=2, column=c).font = Font(bold=True)
    for i, wd in enumerate([10, 42, 14, 10, 50, 50, 4, 60], 1):
        w.column_dimensions[get_column_letter(i)].width = wd
    w.freeze_panes = "A3"

    for k in keep:
        w.append([k["stt"], k["name"], k["code"], k["n"],
                  k["cat_vi"], k["cat_en"], "x", k["path"]])

    wb.save(OUT)
    no_cat = [k["name"] for k in keep if not k["cat_vi"]]
    print("giu   : %d dong SP" % len(keep))
    print("bo    : %d dong" % len(skipped))
    for lab in ("boi vang", "nhom software/tools", "software le"):
        print("        %-22s %d" % (lab, sum(1 for _, r in skipped if r == lab)))
    print("thieu danh muc Ntrend VI: %d dong" % len(no_cat))
    for n in no_cat:
        print("        -", n)
    print("->", OUT)


if __name__ == "__main__":
    main()
