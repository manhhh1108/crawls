# -*- coding: utf-8 -*-
"""Pepperl+Fuchs - map 'dong SP' -> trang danh sach ma SP (?class=<id>).

Vi sao can: danh muc trong 'Danh mục SP web NTREND.xlsx' co 309 dong SP nhung
161 dong KHONG co link, 148 dong con lai ghi link cua CAP CHA (nhieu dong dung
chung mot ?class=). Muon crawl noi dung thi moi dong phai co trang rieng.

Cach lam: duyet cay danh muc bang Chrome that. Moi trang danh muc co cac
"action card" tro xuong cap duoi, kem san ?class=<id> dung cua tung cap.
    trang /products/<...>-gp#####   -> the con (category hoac ?class=)
    trang ?class=<id>               -> the con (neu con) + danh sach ma SP

Ket qua: out/02 Pepperl+Fuchs/classmap.json
    {"<ten dong SP>": {"url": "...?class=123", "path": ["L1","L2","L3"]}}

Chay:
    python pf_classmap.py            # bo sung cho dong con thieu
    python pf_classmap.py --force    # duyet lai tu dau
"""
import io
import json
import os
import re
import sys
import html as H

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

SITE = "https://www.pepperl-fuchs.com"
HUB = SITE + "/en-sg/products-gp25581"
OUT_DIR = os.path.join(HERE, "out", "02 Pepperl+Fuchs")
MAP_JSON = os.path.join(OUT_DIR, "classmap.json")
NTREND_XLSX = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
os.makedirs(OUT_DIR, exist_ok=True)

CARD = re.compile(r'<a class="action-card[^"]*"[^>]*href="([^"]+)"(.*?)</a>', re.S)
CARD_TITLE = re.compile(r'action-card__(?:headline|title)[^>]*>(.*?)<', re.S)
BREAD = re.compile(r'breadcrumb__button button"[^>]*>(.*?)</a>', re.S)


def txt(s):
    return H.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", s or ""))).strip()


def lines_wanted():
    """309 dong SP can co trang, lay tu sheet 02 Pepperl+Fuchs."""
    from openpyxl import load_workbook
    ws = load_workbook(NTREND_XLSX, read_only=True, data_only=True)["02 Pepperl+Fuchs"]
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or len(r) < 9:
            continue
        if str(r[8] or "").strip() != "x":
            continue
        name = (r[1] or "").strip()
        path = [p for p in (r[4], r[5], r[6], r[7]) if p and str(p).strip()]
        if name:
            out.append({"name": name, "path": [str(p).strip() for p in path]})
    return out


def card_links(html):
    """[(url, tieu de)] cua cac the dan xuong cap duoi."""
    out = []
    for href, blk in CARD.findall(html):
        u = href if href.startswith("http") else SITE + href
        t = CARD_TITLE.search(blk)
        out.append((u, txt(t.group(1)) if t else txt(blk)[:60]))
    return out


def walk(seeds, want_names, log=print):
    """Duyet rong tu cac trang goc, dung lai khi da tim du ten can."""
    from playwright.sync_api import sync_playwright
    found, seen, queue = {}, set(), list(seeds)
    want = {n.lower() for n in want_names}
    with sync_playwright() as pw:
        b = pw.chromium.launch(headless=False, channel="chrome")
        pg = b.new_page()
        first = True
        while queue:
            url = queue.pop(0)
            if url in seen or len(seen) > 600:
                continue
            seen.add(url)
            try:
                pg.goto(url, wait_until="networkidle", timeout=90000)
                pg.wait_for_timeout(1500)
            except Exception as e:
                log("  [loi] %s %s" % (url[:70], str(e)[:40]))
                continue
            if first:
                try:
                    from c_pepperl_try import accept_cookies
                    accept_cookies(pg)
                except Exception:
                    pass
                first = False
            html = pg.content()
            crumbs = [txt(x) for x in BREAD.findall(html)]
            crumbs = [c for c in crumbs if c.lower() != "country"]
            for u, title in card_links(html):
                if "/support/" in u or "/news/" in u:
                    continue
                key = title.lower()
                if key in want and key not in found:
                    found[key] = {"name": title, "url": u, "path": crumbs[1:]}
                    log("  + %-42s %s" % (title[:42], u[-46:]))
                if u not in seen:
                    queue.append(u)
            log("    (da duyet %d trang, tim duoc %d/%d dong)"
                % (len(seen), len(found), len(want)))
            if len(found) >= len(want):
                break
        b.close()
    return found


def seeds():
    """Trang danh muc that de bat dau duyet.

    HUB (/products-gp25581) la trang TIM KIEM, khong co the danh muc nao -> phai
    lay cac trang /products/<...>-gp##### trong cay danh muc da crawl truoc do.
    """
    out = []
    try:
        with io.open(os.path.join(REPO, "catalog", "out",
                                  "02 Pepperl+Fuchs.json"), encoding="utf-8") as f:
            nodes = json.load(f)["nodes"]
        for n in sorted(nodes, key=lambda x: int(x.get("level") or 9)):
            u = n.get("url") or ""
            if "/products/" in u and u not in out:
                out.append(u)
    except Exception:
        pass
    return out or [HUB]


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    force = "--force" in sys.argv
    old = {}
    if os.path.exists(MAP_JSON) and not force:
        old = json.load(io.open(MAP_JSON, encoding="utf-8"))
    lines = lines_wanted()
    todo = [l["name"] for l in lines if l["name"].lower() not in old]
    print("dong SP can map: %d | da co: %d | con thieu: %d"
          % (len(lines), len(old), len(todo)))
    if not todo:
        return
    sd = seeds()
    print("trang goc de duyet: %d" % len(sd))
    got = walk(sd, todo)
    old.update(got)
    json.dump(old, io.open(MAP_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    miss = [l["name"] for l in lines if l["name"].lower() not in old]
    print("\nxong: map duoc %d/%d dong | con thieu %d" % (len(old), len(lines), len(miss)))
    for m in miss[:20]:
        print("   thieu:", m)
    print("->", MAP_JSON)


if __name__ == "__main__":
    main()
