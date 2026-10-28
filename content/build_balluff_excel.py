# -*- coding: utf-8 -*-
"""Balluff - PHA 3: content.json -> moi bai 1 file Excel.

    python build_balluff_excel.py

Layout theo 'Template nội dung Balluff.docx'. Bang ma SP co cot DONG: moi dong
SP mot bo cot rieng (Use / Dimension / Material / Version / Focal length ...)
lay tu `columns` cua API. Khoi cot tieng Viet dat truoc, khoi tieng Anh (de
soat) dat sau.
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from openpyxl import Workbook                              # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill   # noqa: E402
from openpyxl.utils import get_column_letter               # noqa: E402

import out_dir                                         # noqa: E402
from bai_prune import prune                             # noqa: E402
from c_pepperl_try import (slugify, SPECDESC, META_TPL,    # noqa: E402
                           SPECDESC_EN, META_TPL_EN)

BRAND = "Balluff"
SRC = os.path.join(HERE, "out", "17 Balluff", "content.json")
OUT_DIR = out_dir.brand_dir("Noi dung Balluff")

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
FLD_FILL = PatternFill("solid", fgColor="DDEBF7")
WRAP = Alignment(wrap_text=True, vertical="top")
MAX_COLS = 8
MAX_TAGS = 12


def safe(s, n=58):
    return re.sub(r'[\\/:*?"<>|]+', "-", s or "").strip()[:n].rstrip(" .")


def ctx(a):
    en = a.get("title_en") or a["name"]
    nv = a.get("name_vi") or en
    first = nv.split(" ")[0]
    cv = {"type_c": nv[:1].upper() + nv[1:],
          "type_l": nv if (len(first) > 1 and first.isupper())
                    else nv[:1].lower() + nv[1:],
          "name": nv, "name_brand": BRAND, "brand": BRAND}
    ce = {"type_c": en, "type_l": en.lower(),
          "name": en, "name_brand": BRAND, "brand": BRAND}
    return cv, ce


def tag_text(a, name):
    """The tag = [Ma san pham], [Thuong hieu]."""
    codes = [it.get("art") for it in (a.get("items") or []) if it.get("art")]
    head = ", ".join(codes[:MAX_TAGS]) if codes else name
    return "%s, %s" % (head, BRAND)


def build_one(a, idx, out_dir):
    cv, ce = ctx(a)
    nv, en = cv["name"], ce["name"]
    wb = Workbook()
    ws = wb.active
    ws.title = "NOI DUNG"
    ws["A1"] = "NỘI DUNG SẢN PHẨM NTREND - %s" % en
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["Trường", "Tiếng Việt (bản đăng web)", "Tiếng Anh (soát chính tả)"])
    for c in range(1, 4):
        ws.cell(row=3, column=c).fill = HDR_FILL
        ws.cell(row=3, column=c).font = Font(bold=True, color="FFFFFF")

    def row(k, vi, en_=""):
        ws.append([k, vi, en_])
        r = ws.max_row
        ws.cell(row=r, column=1).fill = FLD_FILL
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).alignment = WRAP

    kw_vi = "%s %s" % (nv, BRAND)
    kw_en = "%s %s" % (en, BRAND)
    row("Ảnh sản phẩm", a.get("img_file") or "(chưa có ảnh)")
    row("Danh mục sản phẩm", a.get("cat_vi") or "(chưa phân loại)")
    row("Thẻ tag", tag_text(a, nv), tag_text(a, en))
    row("Thương hiệu", BRAND, BRAND)
    row("URL", slugify(kw_vi), slugify(kw_en))
    row("Keyword chính", kw_vi, kw_en)
    row("Meta title", "%s %s" % (cv["type_c"], BRAND), kw_en)
    row("Meta description (155 - 165 characters)",
        META_TPL.format(**cv), META_TPL_EN.format(**ce))
    # Mo ta dai: doan gioi thieu phan Feature + thong so gom tu bang ma SP
    row("Mô tả dài của sản phẩm",
        "\n".join((a.get("prose_vi") or [])
                  + ["- " + x for x in a.get("long_vi") or []]),
        "\n".join((a.get("prose_en") or [])
                  + ["- " + x for x in a.get("long_en") or []]))
    row("Mô tả ngắn của sản phẩm",
        "\n".join("- " + x for x in a.get("short_vi") or []),
        "\n".join("- " + x for x in a.get("short_en") or []))
    row("Mô tả thông số", SPECDESC.format(**cv), SPECDESC_EN.format(**ce))

    items = a.get("items") or []
    cols_en = (a.get("cols_en") or [])[:MAX_COLS]
    cols_vi = (a.get("cols_vi") or [])[:MAX_COLS]
    while len(cols_vi) < len(cols_en):
        cols_vi.append(cols_en[len(cols_vi)])
    if items:
        hdr = ["Bảng mã sản phẩm + Thông tin sản phẩm",
               "Mã đặt hàng (Part No.)", "Model"] + cols_vi \
            + ["%s (EN)" % c for c in cols_en]
        ws.append(hdr)
        r = ws.max_row
        ws.cell(row=r, column=1).fill = FLD_FILL
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in range(2, len(hdr) + 1):
            ws.cell(row=r, column=c).font = Font(bold=True, italic=True)
            ws.cell(row=r, column=c).alignment = WRAP
        n = len(cols_en)
        for it in items:
            vv = (it.get("vals_vi") or it.get("vals") or [])[:n]
            ve = (it.get("vals") or [])[:n]
            line = ["", it.get("art", ""), it.get("name", "")] \
                + list(vv) + [""] * (n - len(vv)) \
                + list(ve) + [""] * (n - len(ve))
            ws.append(line)
            for c in range(2, len(line) + 1):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP
    else:
        ws.append(["Bảng mã sản phẩm + Thông tin sản phẩm",
                   "(trang hãng không có dữ liệu)"])
        ws.cell(row=ws.max_row, column=1).fill = FLD_FILL
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    ws.append([])
    ws.append(["Nhánh danh mục", " > ".join(a.get("path") or []), ""])
    ws.append(["Nguồn", a.get("url_src", ""), ""])
    widths = [34, 22, 30] + [26] * (max(len(cols_vi), 2) * 2)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    fname = "%03d - %s.xlsx" % (idx, safe(en))
    wb.save(os.path.join(out_dir, fname))
    return fname


def build_index(arts, names, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "MUC LUC"
    ws["A1"] = "MỤC LỤC BÀI ĐĂNG - 17 Balluff"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["STT", "Dòng sản phẩm", "Tên tiếng Việt", "File", "Ảnh",
               "Số mã SP", "Số cột thông số", "Danh mục Ntrend",
               "Nhánh danh mục", "Nguồn"])
    for c in range(1, 11):
        ws.cell(row=3, column=c).fill = FLD_FILL
        ws.cell(row=3, column=c).font = Font(bold=True)
    for i, (a, f) in enumerate(zip(arts, names), 1):
        ws.append([i, a.get("title_en") or a["name"], a.get("name_vi", ""), f,
                   "x" if a.get("img_file") else "",
                   len(a.get("items") or []), len(a.get("cols_en") or []),
                   a.get("cat_vi", ""), " > ".join(a.get("path") or []),
                   a.get("url_src", "")])
    for i, w in enumerate([6, 42, 38, 46, 6, 10, 14, 50, 50, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    wb.save(os.path.join(out_dir, "00 MUC LUC.xlsx"))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    arts = json.load(io.open(SRC, encoding="utf-8"))["articles"]
    os.makedirs(OUT_DIR, exist_ok=True)
    names = [build_one(a, i, OUT_DIR) for i, a in enumerate(arts, 1)]
    build_index(arts, names, OUT_DIR)
    prune(OUT_DIR, names)
    nit = sum(len(a.get("items") or []) for a in arts)
    print("Saved %d bai (%d ma SP) -> %s" % (len(arts), nit, OUT_DIR))
    for lbl, bad in (("khong co ma SP", [a for a in arts if not a.get("items")]),
                     ("khong co anh", [a for a in arts if not a.get("img_file")])):
        if bad:
            print("  %s (%d): %s" % (lbl, len(bad),
                  ", ".join((x.get("title_en") or x["name"])[:26] for x in bad[:8])))


if __name__ == "__main__":
    main()
