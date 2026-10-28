# -*- coding: utf-8 -*-
"""Reer - PHA 3: content.json -> moi bai 1 file Excel.

    python build_reer_excel.py

Layout theo 'data/Reer_template.docx'. Khac P+F o BANG MA SAN PHAM: cot thong
so la DONG - moi dong SP mot bo khac nhau, lay tu <table class="shop_attributes">
cua chinh trang hang, ten cot dich sang tieng Viet:

    Model | Ma dat hang | Cap an toan | Chieu cao bao ve (mm) | Do phan giai (mm) | ...
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from openpyxl import Workbook                              # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill   # noqa: E402
from openpyxl.utils import get_column_letter               # noqa: E402

from c_pepperl_try import slugify, SPECDESC, META_TPL, SPECDESC_EN, META_TPL_EN      # noqa: E402

BRAND = "Reer"
SRC = os.path.join(HERE, "out", "03 Reer", "content.json")
OUT_DIR = os.path.join(REPO, "data", "Noi dung Reer")

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
FLD_FILL = PatternFill("solid", fgColor="DDEBF7")
WRAP = Alignment(wrap_text=True, vertical="top")
MAX_COLS = 8          # so cot thong so toi da trong bang ma SP


def safe(s, n=58):
    return re.sub(r'[\\/:*?"<>|]+', "-", s or "").strip()[:n].rstrip(" .")


def ctx(a):
    # ten hien thi = <h1> tren web (day du), khong phai ten rut gon trong danh muc
    en = a.get("title_en") or a["name"]
    nv = a.get("name_vi") or en
    first = nv.split(" ")[0]
    cv = {"type_c": nv[:1].upper() + nv[1:],
          "type_l": nv if (len(first) > 1 and first.isupper())
                    else nv[:1].lower() + nv[1:],
          "name": nv, "name_brand": BRAND, "brand": BRAND}
    ce = {"type_c": en, "type_l": en.lower(),
          "name": en, "name_brand": BRAND, "brand": BRAND}
    return cv, ce


def build_one(a, idx, out_dir):
    cv, ce = ctx(a)
    nv = cv["name"]
    wb = Workbook()
    ws = wb.active
    ws.title = "NOI DUNG"
    ws["A1"] = "NỘI DUNG SẢN PHẨM NTREND - %s" % (a.get("title_en") or a["name"])
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["Trường", "Tiếng Việt (bản đăng web)", "Tiếng Anh (soát chính tả)"])
    for c in range(1, 4):
        ws.cell(row=3, column=c).fill = HDR_FILL
        ws.cell(row=3, column=c).font = Font(bold=True, color="FFFFFF")

    def row(k, vi, en=""):
        ws.append([k, vi, en])
        r = ws.max_row
        ws.cell(row=r, column=1).fill = FLD_FILL
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).alignment = WRAP

    kw_vi = "%s %s" % (nv, BRAND)
    kw_en = "%s %s" % (ce["name"], BRAND)
    row("Ảnh sản phẩm", a.get("img_file") or "(chưa có ảnh)")
    row("Danh mục sản phẩm", a.get("cat_vi") or "(chưa phân loại)")
    row("Thẻ tag", "%s, %s" % (nv, BRAND), "%s, %s" % (ce["name"], BRAND))
    row("Thương hiệu", BRAND, BRAND)
    row("URL", slugify(kw_vi), slugify(kw_en))
    row("Keyword chính", kw_vi, kw_en)
    row("Meta title", "%s %s" % (cv["type_c"], BRAND), kw_en)
    row("Meta description (155 - 165 characters)",
        META_TPL.format(**cv), META_TPL_EN.format(**ce))
    row("Mô tả dài của sản phẩm",
        "\n".join("- " + x for x in a.get("long_vi") or []),
        "\n".join("- " + x for x in a.get("long_en") or []))
    row("Mô tả ngắn của sản phẩm",
        "\n".join("- " + x for x in a.get("short_vi") or []),
        "\n".join("- " + x for x in a.get("short_en") or []))
    row("Mô tả thông số", SPECDESC.format(**cv), SPECDESC_EN.format(**ce))

    items = a.get("items") or []
    cols_en = (a.get("cols_en") or [])[:MAX_COLS]
    cols_vi = (a.get("cols_vi") or [])[:MAX_COLS]
    while len(cols_vi) < len(cols_en):
        cols_vi.append(cols_en[len(cols_vi)])
    if items:
        # review 23/10: khoi cot tieng Viet truoc, khoi tieng Anh (de soat) sau
        hdr = ["Bảng mã sản phẩm + Thông tin sản phẩm", "Model",
               "Mã đặt hàng (Part No.)"] + cols_vi             + ["%s (EN)" % c for c in cols_en]
        ws.append(hdr)
        r = ws.max_row
        ws.cell(row=r, column=1).fill = FLD_FILL
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in range(2, len(hdr) + 1):
            ws.cell(row=r, column=c).font = Font(bold=True, italic=True)
            ws.cell(row=r, column=c).alignment = WRAP
        for it in items:
            d = {k: v for k, v in (it.get("attrs_vi") or [])}
            de = {k: v for k, v in (it.get("attrs") or [])}
            line = ["", it.get("name", ""), it.get("art", "")]
            line += [d.get(k, "") for k in cols_en]
            line += [de.get(k, "") for k in cols_en]
            ws.append(line)
            for c in range(2, len(line) + 1):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP
    else:
        ws.append(["Bảng mã sản phẩm + Thông tin sản phẩm",
                   "(trang hãng không có dữ liệu)"])
        ws.cell(row=ws.max_row, column=1).fill = FLD_FILL
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    ws.append([])
    ws.append(["Nhánh danh mục", " > ".join(a.get("path") or []), ""])
    ws.append(["Nguồn", a.get("url_src", ""), ""])
    widths = [34, 40, 22] + [26] * (max(len(cols_vi), 2) * 2)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    fname = "%03d - %s.xlsx" % (idx, safe(a["name"]))
    wb.save(os.path.join(out_dir, fname))
    return fname


def build_index(arts, names, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "MUC LUC"
    ws["A1"] = "MỤC LỤC BÀI ĐĂNG - 03 Reer"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["STT", "Dòng sản phẩm", "Tên tiếng Việt", "File", "Ảnh",
               "Số mã SP", "Số cột thông số", "Danh mục Ntrend",
               "Nhánh danh mục", "Nguồn"])
    for c in range(1, 11):
        ws.cell(row=3, column=c).fill = FLD_FILL
        ws.cell(row=3, column=c).font = Font(bold=True)
    for i, (a, f) in enumerate(zip(arts, names), 1):
        ws.append([i, a.get("title_en") or a["name"], a.get("name_vi", ""), f,
                   "x" if a.get("img_file") else "",
                   len(a.get("items") or []), len(a.get("cols_en") or []),
                   a.get("cat_vi", ""), " > ".join(a.get("path") or []),
                   a.get("url_src", "")])
    for i, w in enumerate([6, 38, 34, 46, 6, 10, 14, 50, 52, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    wb.save(os.path.join(out_dir, "00 MUC LUC.xlsx"))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    arts = json.load(io.open(SRC, encoding="utf-8"))["articles"]
    os.makedirs(OUT_DIR, exist_ok=True)
    names = [build_one(a, i, OUT_DIR) for i, a in enumerate(arts, 1)]
    build_index(arts, names, OUT_DIR)
    nit = sum(len(a.get("items") or []) for a in arts)
    print("Saved %d bai (%d ma SP) -> %s" % (len(arts), nit, OUT_DIR))
    empty = [a["name"] for a in arts if not a.get("items")]
    if empty:
        print("  bai khong co ma SP (%d): %s" % (len(empty), ", ".join(empty[:8])))


if __name__ == "__main__":
    main()
