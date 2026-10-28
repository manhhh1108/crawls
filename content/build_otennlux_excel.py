# -*- coding: utf-8 -*-
"""OTENNLUX Lighting - PHA 3: content.json -> moi bai 1 file Excel.

    python build_otennlux_excel.py

Layout theo 'Template nội dung OTENNLUX Lighting.docx'. Moi bai la MOT san pham
cu the (khong phai dong SP) nen co hai khoi bang:

    * thuoc tinh san pham (Item no. / Color / Product weight ...)
    * bang model trong phan bai viet (Item No. | power | length ...)

Khoi cot tieng Viet dat truoc, khoi tieng Anh (de soat) dat sau.
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from openpyxl import Workbook                              # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill   # noqa: E402
from openpyxl.utils import get_column_letter               # noqa: E402

import out_dir                                         # noqa: E402
from bai_prune import prune                             # noqa: E402
from c_pepperl_try import (slugify, SPECDESC, META_TPL,    # noqa: E402
                           SPECDESC_EN, META_TPL_EN)

BRAND = "OTENNLUX"
BRAND_FULL = "OTENNLUX Lighting"
SRC = os.path.join(HERE, "out", "16 OTENNLUX Lighting", "content.json")
OUT_DIR = out_dir.brand_dir("Noi dung OTENNLUX Lighting")

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
FLD_FILL = PatternFill("solid", fgColor="DDEBF7")
WRAP = Alignment(wrap_text=True, vertical="top")
MAX_TAGS = 12


def safe(s, n=58):
    return re.sub(r'[\\/:*?"<>|]+', "-", s or "").strip()[:n].rstrip(" .")


def ctx(a):
    en = a.get("title_en") or a["name"]
    nv = a.get("name_vi") or en
    first = nv.split(" ")[0]
    cv = {"type_c": nv[:1].upper() + nv[1:],
          "type_l": nv if (len(first) > 1 and first.isupper())
                    else nv[:1].lower() + nv[1:],
          "name": nv, "name_brand": BRAND, "brand": BRAND}
    ce = {"type_c": en, "type_l": en.lower(),
          "name": en, "name_brand": BRAND, "brand": BRAND}
    return cv, ce


def tag_text(a, name):
    codes = a.get("codes") or []
    head = ", ".join(codes[:MAX_TAGS]) if codes else name
    return "%s, %s" % (head, BRAND)


def build_one(a, idx, out_dir):
    cv, ce = ctx(a)
    nv, en = cv["name"], ce["name"]
    wb = Workbook()
    ws = wb.active
    ws.title = "NOI DUNG"
    ws["A1"] = "NỘI DUNG SẢN PHẨM NTREND - %s" % en
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["Trường", "Tiếng Việt (bản đăng web)", "Tiếng Anh (soát chính tả)"])
    for c in range(1, 4):
        ws.cell(row=3, column=c).fill = HDR_FILL
        ws.cell(row=3, column=c).font = Font(bold=True, color="FFFFFF")

    def row(k, vi, en_=""):
        ws.append([k, vi, en_])
        r = ws.max_row
        ws.cell(row=r, column=1).fill = FLD_FILL
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).alignment = WRAP

    def block(hdr):
        ws.append(hdr)
        r = ws.max_row
        ws.cell(row=r, column=1).fill = FLD_FILL
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in range(2, len(hdr) + 1):
            ws.cell(row=r, column=c).font = Font(bold=True, italic=True)
            ws.cell(row=r, column=c).alignment = WRAP

    kw_vi = "%s %s" % (nv, BRAND)
    kw_en = "%s %s" % (en, BRAND)
    row("Ảnh sản phẩm", a.get("img_file") or "(chưa có ảnh)")
    row("Danh mục sản phẩm", a.get("cat_vi") or "(chưa phân loại)")
    row("Thẻ tag", tag_text(a, nv), tag_text(a, en))
    row("Thương hiệu", BRAND_FULL, BRAND_FULL)
    row("URL", slugify(kw_vi), slugify(kw_en))
    row("Keyword chính", kw_vi, kw_en)
    row("Meta title", "%s %s" % (cv["type_c"], BRAND), kw_en)
    row("Meta description (155 - 165 characters)",
        META_TPL.format(**cv), META_TPL_EN.format(**ce))
    # Mo ta dai = full phan thong tin phia tren bang thong so ky thuat
    row("Mô tả dài của sản phẩm",
        "\n".join(a.get("prose_vi") or a.get("long_vi") or []),
        "\n".join(a.get("prose_en") or a.get("long_en") or []))
    row("Mô tả ngắn của sản phẩm",
        "\n".join(a.get("short_vi") or []),
        "\n".join(a.get("short_en") or []))
    row("Mô tả thông số", SPECDESC.format(**cv), SPECDESC_EN.format(**ce))

    attrs = a.get("attrs_en") or []
    av = a.get("attrs_vi") or attrs
    models = a.get("models_en") or []
    mv = a.get("models_vi") or models

    if attrs:
        block(["Thông số kỹ thuật", "Tên thông số", "Giá trị",
               "Parameter (EN)", "Value (EN)"])
        for (ke, ve), (kv, vv) in zip(attrs, av):
            ws.append(["", kv, vv, ke, ve])
            for c in range(2, 6):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP

    for n, (t_en, t_vi) in enumerate(zip(models, mv), 1):
        ncol = len(t_en["head"])
        block(["Bảng mã sản phẩm %d" % n if len(models) > 1
               else "Bảng mã sản phẩm"]
              + list(t_vi["head"]) + ["%s (EN)" % c for c in t_en["head"]])
        for r_en, r_vi in zip(t_en["rows"], t_vi["rows"]):
            line = [""] + (list(r_vi) + [""] * ncol)[:ncol] \
                        + (list(r_en) + [""] * ncol)[:ncol]
            ws.append(line)
            for c in range(2, len(line) + 1):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP

    if not attrs and not models:
        block(["Thông số kỹ thuật", "(trang hãng không có dữ liệu)"])

    ws.append([])
    ws.append(["Nhánh danh mục", " > ".join(a.get("path") or []), ""])
    ws.append(["Nguồn", a.get("url_src", ""), ""])
    ncol = max([len(t["head"]) for t in models] + [2])
    for i, w in enumerate([34, 40, 46] + [24] * (ncol * 2), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    fname = "%03d - %s.xlsx" % (idx, safe(en))
    wb.save(os.path.join(out_dir, fname))
    return fname


def build_index(arts, names, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "MUC LUC"
    ws["A1"] = "MỤC LỤC BÀI ĐĂNG - 16 OTENNLUX Lighting"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["STT", "Sản phẩm", "Tên tiếng Việt", "File", "Ảnh", "Số mã SP",
               "Số thông số", "Danh mục Ntrend", "Nhánh danh mục", "Nguồn"])
    for c in range(1, 11):
        ws.cell(row=3, column=c).fill = FLD_FILL
        ws.cell(row=3, column=c).font = Font(bold=True)
    for i, (a, f) in enumerate(zip(arts, names), 1):
        ws.append([i, a.get("title_en") or a["name"], a.get("name_vi", ""), f,
                   "x" if a.get("img_file") else "",
                   len(a.get("codes") or []), len(a.get("attrs_en") or []),
                   a.get("cat_vi", ""), " > ".join(a.get("path") or []),
                   a.get("url_src", "")])
    for i, w in enumerate([6, 50, 44, 48, 6, 10, 12, 46, 44, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    wb.save(os.path.join(out_dir, "00 MUC LUC.xlsx"))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    arts = json.load(io.open(SRC, encoding="utf-8"))["articles"]
    os.makedirs(OUT_DIR, exist_ok=True)
    names = [build_one(a, i, OUT_DIR) for i, a in enumerate(arts, 1)]
    build_index(arts, names, OUT_DIR)
    prune(OUT_DIR, names)
    print("Saved %d bai (%d ma SP, %d bai co bang model) -> %s"
          % (len(arts), sum(len(a.get("codes") or []) for a in arts),
             sum(1 for a in arts if a.get("models_en")), OUT_DIR))
    for lbl, bad in (("khong co bang nao",
                      [a for a in arts
                       if not a.get("models_en") and not a.get("attrs_en")]),
                     ("khong co anh",
                      [a for a in arts if not a.get("img_file")])):
        if bad:
            print("  %s (%d): %s" % (lbl, len(bad),
                  ", ".join((x.get("title_en") or x["name"])[:26] for x in bad[:8])))


if __name__ == "__main__":
    main()
