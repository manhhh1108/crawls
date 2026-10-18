# -*- coding: utf-8 -*-
"""Dich TEN DONG SAN PHAM SICK sang tieng Viet (review 13/10, muc 4).

Yeu cau: "Co nhung san pham khong co ten ma ma no la ten san pham, thi chi de
ban dich" -> moi ten dong SP la CUM TIENG ANH MO TA deu phai hien thi bang
tieng Viet o cot B (The tag, Keyword, Meta title, Meta description, mo ta).
Ten la MA / TEN RIENG cua SICK (W16, microScan3, Flexi Soft, Lector, OD Mini...)
thi giu nguyen.

Cach lam:
  1. Danh sach da chot tay (sick_names_vi.NAME_VI + SUGGESTED) -> uu tien cao nhat.
  2. Cac ten con lai: hoi Gemini phan loai GIU NGUYEN / DICH, kem ban dich.
  3. Ket qua ghi ra:
       content/out/01 SICK/names_vi.json       (pipeline doc)
       data/SICK - ten dong SP can dich.xlsx   (de Ntrend soat & sua tay)
     Muon sua 1 ten: sua cot C trong file xlsx roi chay lai script nay
     (--from-xlsx) hoac chinh thang names_vi.json.

Chay:
    python sick_name_translate.py              # dich phan con thieu
    python sick_name_translate.py --force      # hoi lai tu dau
    python sick_name_translate.py --from-xlsx  # chi nap lai ban sua tay tu xlsx
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

OUT_DIR = os.path.join(HERE, "out", "01 SICK")
NAMES_JSON = os.path.join(OUT_DIR, "names_vi.json")
XLSX = os.path.join(REPO, "data", "SICK - ten dong SP can dich.xlsx")
SHEET = "TEN DONG SP"

KEEP = "GIU NGUYEN"

SYSTEM = """Bạn là biên tập viên nội dung thương mại điện tử ngành tự động hóa công nghiệp, am hiểu danh mục sản phẩm của hãng SICK (Đức).

Nhiệm vụ: với mỗi TÊN DÒNG SẢN PHẨM tiếng Anh, quyết định 1 trong 2:

A. Trả về đúng chuỗi "GIU NGUYEN" — nếu tên đó là MÃ sản phẩm hoặc TÊN RIÊNG một dòng/họ sản phẩm của SICK. Dấu hiệu:
   - có chữ số hoặc hậu tố mã: W16, DT35-S, IME2S, MLG-2, i10 Lock, W100 Laser, S300 Mini, AFS/AFM60 Ethernet, PGT-10 Pro, MZC1 Twin, DFS60S Pro, Visionary-T Mini
   - tên thương hiệu dòng SP của SICK: Flexi Soft, Flexi Loop, Flexi Classic, Lector, Inspector, Bulkscan, Visionary, Profiler, Glare, deTec, microScan3, nanoScan3, outdoorScan3, picoScan, multiScan, miniTwin, T-Easic, Coriva, OccupID, PALLOC, ReLy, Safe Portal
   - tên riêng + một từ tiếng Anh làm biến thể (Lock, Pro, Mini, Twin, Direct, Standard, Laser, Ethernet, Dynamic, Hi, Value, Precision, Cubic, Inox, Hygienic, plus)

B. Trả về BẢN DỊCH TIẾNG VIỆT — nếu tên đó chỉ là CỤM TỪ MÔ TẢ chủng loại hàng hóa, khách Việt sẽ tra cứu bằng tiếng Việt. Ví dụ:
   Protective housing = Vỏ bảo vệ | Reflectors = Gương phản xạ | Positioning codes = Mã định vị
   Power supply units = Bộ nguồn | Smoke detectors = Đầu báo khói | Memory cards = Thẻ nhớ
   Area Hotspot Detection System = Hệ thống phát hiện điểm nóng theo vùng

QUY TẮC KHI DỊCH:
1. Tiếng Việt CÓ DẤU đầy đủ, viết hoa chữ đầu, không viết hoa toàn bộ.
2. GIỮ NGUYÊN trong bản dịch: từ viết tắt kỹ thuật (GPS, RFID, WLAN, GSM, UMTS, LTE, HF, UHF, SSI, USB, IO-Link, HIPERFACE, Ethernet, CLV, ICR, RF, VMV, AGV, DWS) và tên riêng của SICK (Flexi, Lector, EFI-pro, deTec, deTem, microScan3...).
   Ví dụ: "GPS antennas" = "Anten GPS"; "Safe EFI-pro System" = "Hệ thống an toàn EFI-pro"; "System plugs Flexi Soft" = "Phích cắm hệ thống Flexi Soft".
3. Dịch theo nghĩa ngành: housing = vỏ (không phải "nhà ở"); bus = bus truyền thông; actuator = cơ cấu tác động; collet = ống kẹp; bracket = giá đỡ.
4. Ngắn gọn như tên danh mục hàng hóa, không thêm giải thích, không thêm chữ "SICK".
5. Không chắc chắn thì trả "GIU NGUYEN"."""


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").replace("\xa0", " ")).strip()


def curated():
    """Danh sach da chot tay trong sick_names_vi.py (ca 2 nhom - review da duyet
    dich HET cac ten dang cum tieng Anh)."""
    import sick_names_vi as S
    out = {}
    for src in (S.SUGGESTED, S.NAME_VI):       # NAME_VI de sau = uu tien hon
        for k, v in src.items():
            out[_norm(k)] = _norm(v)
    return out


def all_names():
    with io.open(os.path.join(OUT_DIR, "raw.json"), encoding="utf-8") as f:
        return [_norm(l["name"]) for l in json.load(f)["lines"]]


def load():
    if os.path.exists(NAMES_JSON):
        with io.open(NAMES_JSON, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save(d):
    os.makedirs(OUT_DIR, exist_ok=True)
    with io.open(NAMES_JSON, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=1, sort_keys=True)


def ask_llm(names, chunk=40):
    """Hoi Gemini: moi ten -> 'GIU NGUYEN' hoac ban dich tieng Viet."""
    import llm_translate as L

    out = {}
    for i in range(0, len(names), chunk):
        part = names[i:i + chunk]
        prompt = ("Xu ly %d ten duoi day. Tra ve DUNG mot mang JSON gom %d chuoi, "
                  "cung thu tu, moi phan tu la \"GIU NGUYEN\" hoac ban dich tieng Viet.\n%s"
                  % (len(part), len(part), json.dumps(part, ensure_ascii=False, indent=1)))
        L._resolve_endpoints()
        last = None
        for _try in range(len(L.ENDPOINTS) * 2):
            prov, model = L.ENDPOINTS[L._cur]
            try:
                got = L._parse(L.CALLERS[prov](model, SYSTEM, prompt, len(part)),
                               len(part))
                break
            except Exception as e:
                last = e
                if ("429" in str(e) or "503" in str(e) or "rate_limit" in str(e)) \
                        and L._cur + 1 < len(L.ENDPOINTS):
                    L._cur += 1
                    print("  [chuyen -> %s/%s]" % L.ENDPOINTS[L._cur])
                    continue
        else:
            raise last
        for en, vi in zip(part, got):
            vi = _norm(vi)
            out[en] = "" if (not vi or vi.upper().replace("Ữ", "U") == KEEP
                             or vi.lower() == en.lower()) else vi
        print("  da hoi %d/%d" % (min(i + chunk, len(names)), len(names)))
    return out


def write_xlsx(names, mapping, src):
    """Xuat file de Ntrend soat tay. Cot C sua duoc, chay --from-xlsx de ap dung."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = "TÊN DÒNG SP SICK – BẢN DỊCH TIẾNG VIỆT DÙNG Ở CỘT B (bản đăng web)"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws["A2"] = ("Cột C trống = giữ nguyên tên tiếng Anh (là mã/tên riêng của SICK). "
                "Sửa cột C rồi chạy: python sick_name_translate.py --from-xlsx")
    ws.append([])
    ws.append(["STT", "Tên dòng SP (EN)", "Tiếng Việt dùng ở cột B", "Nguồn"])
    hdr = ws.max_row
    for c in range(1, 5):
        ws.cell(row=hdr, column=c).fill = PatternFill("solid", fgColor="1F4E79")
        ws.cell(row=hdr, column=c).font = Font(bold=True, color="FFFFFF")
    n = 0
    for nm in names:
        vi = mapping.get(nm, "")
        if not vi:
            continue
        n += 1
        ws.append([n, nm, vi, src.get(nm, "")])
        for c in range(1, 5):
            ws.cell(row=ws.max_row, column=c).alignment = Alignment(
                wrap_text=True, vertical="top")
    for i, w in enumerate([6, 48, 48, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    wb.save(XLSX)
    return n


def from_xlsx():
    """Nap lai ban sua tay o cot C cua file xlsx -> names_vi.json."""
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, data_only=True)
    ws = wb[SHEET]
    d = load()
    changed = 0
    for r in ws.iter_rows(values_only=True):
        if not r or len(r) < 3 or r[0] is None:
            continue
        if not str(r[0]).strip().isdigit():
            continue
        en, vi = _norm(r[1]), _norm(r[2])
        if d.get(en, "") != vi:
            d[en] = vi
            changed += 1
    wb.close()
    save(d)
    print("nap tu %s: %d ten thay doi, tong %d ten co ban dich"
          % (os.path.basename(XLSX), changed, sum(1 for v in d.values() if v)))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if "--from-xlsx" in sys.argv:
        return from_xlsx()

    force = "--force" in sys.argv
    names = all_names()
    cur = curated()
    done = {} if force else load()

    src = {}
    mapping = {}
    for nm in names:
        if nm in cur:
            mapping[nm] = cur[nm]
            src[nm] = "chốt tay"
    todo = [n for n in names if n not in mapping and n not in done]
    print("tong %d ten | da chot tay %d | da co san %d | can hoi LLM %d"
          % (len(names), len(mapping), len(done), len(todo)))

    if todo:
        done.update(ask_llm(todo))
    for nm in names:
        if nm in mapping:
            continue
        mapping[nm] = done.get(nm, "")
        if mapping[nm]:
            src[nm] = "Gemini"
    save(mapping)
    n = write_xlsx(names, mapping, src)
    print("\n%d/%d ten dung ban dich tieng Viet (con lai giu nguyen ma/ten rieng)"
          % (n, len(names)))
    print("  -> %s" % NAMES_JSON)
    print("  -> %s" % XLSX)


if __name__ == "__main__":
    main()
