# -*- coding: utf-8 -*-
"""So sanh Google Translate (dang dung) vs Gemini (de xuat) tren cung mot bo cau.

    python compare_translate.py            # bo 24 cau kho mac dinh
    python compare_translate.py --from-data 30   # lay 30 cau that tu content.json

In ra bang 3 cot va ghi them data/So sanh dich Google vs Gemini.xlsx de doi chieu.
"""
import io
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import llm_translate                              # noqa: E402
import vi_terms                                   # noqa: E402

# cac cau may dich hay sai nhat - lay tu chinh loi da gap trong du an
HARD = [
    ("key", "Housing material"),
    ("key", "Housing design"),
    ("key", "Bus termination"),
    ("key", "Drive concept"),
    ("key", "Shaft"),
    ("key", "Measuring wheel mechanics"),
    ("key", "Wire draw mechanism"),
    ("key", "Front screen heating"),
    ("key", "Degree of protection"),
    ("key", "Response time"),
    ("value", "Male connector, 4-pin"),
    ("value", "Female connector, M12, 5-pin"),
    ("value", "Teach-in by wire"),
    ("value", "1,200 mm"),
    ("prose", "The track and trace system reads codes and assigns them to the "
              "associated object."),
    ("prose", "The housing is made of die-cast zinc and the front screen of PMMA."),
    ("prose", "Bus-capable sensors reduce wiring effort in the control cabinet."),
    ("prose", "The drive monitors speed and standstill via two independent "
              "channels."),
    ("prose", "Rugged design for harsh ambient conditions, protection class IP69K."),
    ("prose", "Quick and easy mounting thanks to the universal bracket."),
    ("prose", "The light curtain switches off the machine when a hand enters the "
              "protective field."),
    ("prose", "Order no. 1055909, delivery from stock."),
    ("prose", "Suitable for conveyor belts with speeds of up to 2.5 m/s."),
    ("prose", "The sensor detects objects regardless of colour, surface or "
              "material."),
]


def from_data(n):
    p = os.path.join(HERE, "out", "01 SICK", "content.json")
    arts = json.load(io.open(p, encoding="utf-8"))["articles"]
    out, seen = [], set()
    for a in arts:
        for s in (a.get("overview_en") or [])[:1]:
            if s and s not in seen and 60 < len(s) < 320:
                seen.add(s)
                out.append(("prose", s))
        for k, _v in (a.get("specs_en") or [])[:1]:
            if k and k not in seen:
                seen.add(k)
                out.append(("key", k))
        if len(out) >= n:
            break
    return out[:n]


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    rows = HARD
    if len(sys.argv) > 2 and sys.argv[1] == "--from-data":
        rows = from_data(int(sys.argv[2]))

    import vi_engine
    terms = vi_terms.terms()

    # dung dung pipeline hien tai: vi_engine (tu dien + mask + Google + postfix)
    print("Dang dich %d cau bang PIPELINE HIEN TAI (tu dien + Google)..." % len(rows))
    g = [None] * len(rows)
    for kind in ("key", "value", "prose"):
        idx = [i for i, (k, _s) in enumerate(rows) if k == kind]
        if not idx:
            continue
        got = vi_engine.translate_many([rows[i][1] for i in idx], kind)
        for i, v in zip(idx, got):
            g[i] = v

    print("Dang dich bang Gemini (%s)..." % llm_translate.MODEL)
    m = []
    for kind in ("key", "value", "prose"):
        idx = [i for i, (k, _s) in enumerate(rows) if k == kind]
        if not idx:
            continue
        got = llm_translate.translate_batch([rows[i][1] for i in idx], kind, terms)
        m += list(zip(idx, got))
    gem = [None] * len(rows)
    for i, v in m:
        gem[i] = v

    print()
    for (kind, en), a, b in zip(rows, g, gem):
        mark = "  " if a.strip() == (b or "").strip() else "**"
        print("%s [%s] %s" % (mark, kind, en[:96]))
        print("     Hiện tại: %s" % (a or "")[:110])
        print("     Gemini: %s" % (b or "")[:110])
        print()
    diff = sum(1 for a, b in zip(g, gem) if a.strip() != (b or "").strip())
    print("Khac nhau %d/%d cau (** = khac)" % (diff, len(rows)))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook(); w = wb.active; w.title = "SO SANH"
        w["A1"] = "SO SÁNH BẢN DỊCH: pipeline hiện tại (từ điển + Google) vs Gemini"
        w["A1"].font = Font(bold=True, size=13, color="1F4E79")
        w.append([])
        w.append(["Loại", "Tiếng Anh", "Pipeline hiện tại (từ điển + Google)", "Gemini (đề xuất)", "Chọn bản nào?"])
        for c in range(1, 6):
            w.cell(row=3, column=c).fill = PatternFill("solid", fgColor="DDEBF7")
            w.cell(row=3, column=c).font = Font(bold=True)
        for (kind, en), a, b in zip(rows, g, gem):
            w.append([kind, en, a, b, ""])
            for c in range(1, 6):
                w.cell(row=w.max_row, column=c).alignment = Alignment(wrap_text=True,
                                                                     vertical="top")
        for i, wd in enumerate([8, 52, 52, 52, 16], 1):
            w.column_dimensions[get_column_letter(i)].width = wd
        w.freeze_panes = "A4"
        p = os.path.join(REPO, "data", "So sanh dich Google vs Gemini.xlsx")
        wb.save(p)
        print("->", p)
    except Exception as e:
        print("(khong xuat duoc xlsx:", repr(e), ")")


if __name__ == "__main__":
    main()
