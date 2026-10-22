# -*- coding: utf-8 -*-
"""Pepperl+Fuchs - map 'dong SP' -> trang ?class=, KHOP THEO CA NHANH CHA.

Van de cua ban dau (classmap_all.json): map khoa theo TEN dong SP. Nhung 48 ten
bi lap - rieng 'Accessories' lap 18 lan duoi 18 nhanh khac nhau, 'Peripherals'
4 lan... Khop theo ten thi ca 18 dong deu tro ve mot trang => 17 bai sai.

Cach lam: moi the danh muc duoc ghi kem TEN TRANG CHA tim thay no, nen khoa la
    (ten cha, ten the)
Khi tra, uu tien khop ca cap; khong co moi lui ve khop ten.

Nguon:
  1. catalog/out/02 Pepperl+Fuchs.json  - cay danh muc da crawl (co path day du)
  2. .cache/pepperl/*.gz                - HTML cac trang da tai lan truoc
  3. out/02 Pepperl+Fuchs/classmap.json - walk pf_classmap.py (co path)

Ket qua: out/02 Pepperl+Fuchs/classmap2.json
    {"<cha>||<ten>": {"name":..., "url":..., "parent":...}}
"""
import ast
import gzip
import io
import json
import os
import re
import sys
import html as H

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
OUT = os.path.join(HERE, "out", "02 Pepperl+Fuchs")
CACHE = os.path.join(REPO, ".cache", "pepperl")
SITE = "https://www.pepperl-fuchs.com"
MAP2 = os.path.join(OUT, "classmap2.json")

CARD = re.compile(r'<a class="action-card[^"]*"[^>]*href="([^"]+)"(.*?)</a>', re.S)
HEAD = re.compile(r"action-card__headline[^>]*>(.*?)<", re.S)
BREAD = re.compile(r'breadcrumb__button button"[^>]*>(.*?)</a>', re.S)
H1 = re.compile(r"<h1[^>]*>(.*?)</h1>", re.S)


def txt(s):
    return H.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", s or ""))).strip()


def full(u):
    return u if u.startswith("http") else SITE + u


def key(parent, name):
    return "%s||%s" % ((parent or "").strip().lower(), (name or "").strip().lower())


def page_parent(html):
    """Ten cap cuoi trong breadcrumb = 'cha' cua cac the tren trang do."""
    cr = [txt(x) for x in BREAD.findall(html)]
    cr = [c for c in cr if c.lower() not in ("country", "products")]
    if cr:
        return cr[-1]
    h = H1.search(html)
    return txt(h.group(1)) if h else ""


def from_cache():
    out = {}
    if not os.path.isdir(CACHE):
        return out
    for f in os.listdir(CACHE):
        if not f.endswith(".gz"):
            continue
        try:
            h = gzip.open(os.path.join(CACHE, f), "rt",
                          encoding="utf-8", errors="replace").read()
        except Exception:
            continue
        par = page_parent(h)
        for href, blk in CARD.findall(h):
            m = HEAD.search(blk)
            if not m:
                continue
            nm = txt(m.group(1))
            if not nm:
                continue
            out.setdefault(key(par, nm), {"name": nm, "parent": par,
                                          "url": full(href.replace("&amp;", "&"))})
    return out


def from_nodes():
    out = {}
    p = os.path.join(REPO, "catalog", "out", "02 Pepperl+Fuchs.json")
    if not os.path.exists(p):
        return out
    for n in json.load(io.open(p, encoding="utf-8"))["nodes"]:
        u = n.get("url")
        if not u:
            continue
        path = n.get("path")
        if isinstance(path, str):
            try:
                path = ast.literal_eval(path)
            except Exception:
                path = [path]
        path = [str(x).strip() for x in (path or []) if str(x).strip()]
        if not path:
            continue
        nm = path[-1]
        par = path[-2] if len(path) > 1 else ""
        out.setdefault(key(par, nm), {"name": nm, "parent": par, "url": full(u)})
    return out


def from_walk():
    p = os.path.join(OUT, "classmap.json")
    if not os.path.exists(p):
        return {}
    out = {}
    for _k, v in json.load(io.open(p, encoding="utf-8")).items():
        path = v.get("path") or []
        par = path[-1] if path else ""
        out.setdefault(key(par, v["name"]), {"name": v["name"], "parent": par,
                                             "url": v["url"]})
    return out


def lines():
    from openpyxl import load_workbook
    p = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
    ws = load_workbook(p, read_only=True, data_only=True)["02 Pepperl+Fuchs"]
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or len(r) < 9 or str(r[8] or "").strip() != "x":
            continue
        nm = (r[1] or "").strip()
        if not nm:
            continue
        path = [str(x).strip() for x in (r[4], r[5], r[6], r[7])
                if x and str(x).strip()]
        out.append({"stt": str(r[0]).strip(), "name": nm, "path": path})
    return out


def lookup(m, ln):
    """Tra trang cho 1 dong SP: khop (cha, ten) truoc, roi moi khop ten."""
    path = ln["path"]
    nm = ln["name"]
    # cha trong danh muc = phan tu ngay truoc ten dong SP
    idx = len(path) - 1
    while idx >= 0 and path[idx].lower() != nm.lower():
        idx -= 1
    par = path[idx - 1] if idx > 0 else (path[-2] if len(path) > 1 else "")
    for cand in (par,) + tuple(reversed(path[:-1])):
        hit = m.get(key(cand, nm))
        if hit:
            return hit, "cha=%s" % cand
    for k, v in m.items():
        if k.endswith("||" + nm.lower()):
            return v, "chi khop ten"
    return None, ""


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    m = {}
    for src, lab in ((from_nodes(), "cay danh muc"), (from_cache(), "cache HTML"),
                     (from_walk(), "walk")):
        before = len(m)
        for k, v in src.items():
            m.setdefault(k, v)
        print("  %-14s +%d muc (tong %d)" % (lab, len(m) - before, len(m)))
    json.dump(m, io.open(MAP2, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    ls = lines()
    ok = exact = loose = 0
    miss = []
    for ln in ls:
        hit, how = lookup(m, ln)
        if not hit:
            miss.append(ln["name"])
            continue
        ok += 1
        if how == "chi khop ten":
            loose += 1
        else:
            exact += 1
    print()
    print("dong SP: %d | tim duoc trang: %d (khop ca nhanh: %d, chi khop ten: %d)"
          % (len(ls), ok, exact, loose))
    print("con thieu: %d" % len(miss))
    print("->", MAP2)


if __name__ == "__main__":
    main()
