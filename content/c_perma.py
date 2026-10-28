# -*- coding: utf-8 -*-
"""Crawl NOI DUNG 78 bai (dong san pham) Perma cho web NTREND.
Dau vao: catalog/out/05 Perma.json (cac node la = dong SP).
Dau ra:  content/out/05 Perma/content.json + anh goc (src_img/) + anh da ghep
         logo NTREND (images/)."""
import io, json, os, re, html as H, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, REPO)
sys.path.insert(0, HERE)

from shared.fetcher import try_fetch  # noqa: E402
from translate import vi as _vi, _flush  # noqa: E402
import brandimg                        # noqa: E402

# sua loi dich may cho thuat ngu nganh (ren, ma dat hang, don vi bar...)
_FIX_VI = [
    (re.compile(r"\bĐiều\s+(\d{4,7})\b"), r"Art. No. \1"),
    (re.compile(r"\bnam\b"), "ren ngoài"),
    (re.compile(r"\bnữ\b"), "ren trong"),
    (re.compile(r"^Phạm vi\b"), "Dòng"),
    (re.compile(r"(\d)\s*thanh\b"), r"\1 bar"),
    (re.compile(r"\bLái xe\b"), "Truyền động"),
    (re.compile(r"\bNhà ở\b"), "Vỏ"),
]

# ten thong so dich thang bang tu dien (khong qua Google de khoi sai nganh)
SPEC_KEY_VI = {
    "Drive": "Truyền động",
    "Power supply": "Nguồn cấp",
    "Housing": "Vỏ",
    "Discharge period": "Chu kỳ xả",
    "Lubricant volume": "Dung tích chất bôi trơn",
    "Operating temperature": "Nhiệt độ làm việc",
    "Pressure build-up": "Áp suất tạo được",
    "Protection class": "Cấp bảo vệ",
    "Ambient temperature": "Nhiệt độ môi trường",
    "Dimensions": "Kích thước",
    "Weight": "Khối lượng",
    "Certification": "Chứng nhận",
    "Capacity": "Dung tích",
    "Operating time": "Thời gian hoạt động",
    "Battery pack": "Bộ pin",
    "Oil level": "Mức dầu",
    "Connection": "Đầu nối",
}


def vi(text):
    s = _vi(text)
    for pat, rep in _FIX_VI:
        s = pat.sub(rep, s)
    return s


def vi_spec_key(k):
    return SPEC_KEY_VI.get(k.strip()) or vi(k)

CAT_JSON = os.path.join(REPO, "catalog", "out", "05 Perma.json")
OUT_DIR = os.path.join(HERE, "out", "05 Perma")
SRC_IMG = os.path.join(OUT_DIR, "src_img")
BRAND_IMG = os.path.join(OUT_DIR, "images")
for p in (OUT_DIR, SRC_IMG, BRAND_IMG):
    os.makedirs(p, exist_ok=True)

# "Loai san pham" cho meta/mo ta, theo nhom cha trong cay danh muc
TYPE_VI = {
    "Single-Point Lubrication Systems": "Bộ bôi trơn tự động đơn điểm",
    "Multi-Point Lubrication Systems": "Bộ bôi trơn tự động đa điểm",
    "Lubricants": "Chất bôi trơn",
    "Angles": "Cút nối góc",
    "Brackets": "Giá đỡ",
    "Cartridges / Pails / Bottles / Canisters for pre-filling":
        "Bình chứa chất bôi trơn nạp sẵn",
    "Extensions": "Ống nối dài",
    "Hose / tube connectors": "Đầu nối ống",
    "INSTALLATION KITs": "Bộ phụ kiện lắp đặt",
    "Lubrication gears": "Bánh răng bôi trơn",
    "Oil brushes": "Chổi tra dầu",
    "Product accessories": "Phụ kiện sản phẩm",
    "Reducers": "Đầu chuyển ren",
    "Special brushes / Special mounting brackets": "Chổi & giá đỡ chuyên dụng",
    "Accessories": "Phụ kiện",
}
TYPE_EN = {
    "Single-Point Lubrication Systems": "Single-point lubrication system",
    "Multi-Point Lubrication Systems": "Multi-point lubrication system",
    "Lubricants": "Lubricant",
    "Angles": "Angle connector",
    "Brackets": "Mounting bracket",
    "Cartridges / Pails / Bottles / Canisters for pre-filling":
        "Pre-filled lubricant container",
    "Extensions": "Extension",
    "Hose / tube connectors": "Hose / tube connector",
    "INSTALLATION KITs": "Installation kit",
    "Lubrication gears": "Lubrication gear",
    "Oil brushes": "Oil brush",
    "Product accessories": "Product accessory",
    "Reducers": "Reducer",
    "Special brushes / Special mounting brackets": "Special brush / bracket",
    "Accessories": "Accessory",
}

BAD_PARA = ("Watch video", "Download", "Newsletter", "Sign up", "moreperma",
            "Cookie", "Imprint", "Privacy", "all lubrication systems",
            "Art. No", "More info")  # text card san pham, khong phai overview
BAD_LI = ("instruction", "flyer", "data sheet", "datasheet", "drawing", "video",
          "login", "contact", "download", "certificate", "newsletter", "more")


def txt(s):
    return re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", "", s))).strip()


def slugify(name):
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return re.sub(r"-{2,}", "-", s)


# Trang perma nhet ca breadcrumb va nhan tab vao CUNG mot the <p> voi noi dung
# that: "Lubrication Systems ... perma NOVA 360° perma NOVA  perma NOVA can be
# used..." hoac "ApplicationsFeaturesTechnical Data...perma NOVA was specifically
# developed...". Phai cat bo phan dau (review 22/10).
TAB_WORDS = ("Applications", "Features", "Technical Data", "Downloads",
             "Accessories", "Videos", "Downloads & Media", "Benefits")


def clean_para(s, title=""):
    """Bo breadcrumb / nhan tab dinh o dau doan van."""
    s = (s or "").strip()
    if not s:
        return ""
    # bo cac nhan tab lap o dau
    changed = True
    while changed:
        changed = False
        for w in TAB_WORDS:
            if s.startswith(w):
                s = s[len(w):].lstrip(" :|-")
                changed = True
    # noi dung that thuong bat dau bang "<ten SP> <dong tu thuong>"
    if title:
        # ten SP co the lap vai lan trong breadcrumb ("... perma NOVA 360° perma
        # NOVA   perma NOVA can be used...") -> lay lan XUAT HIEN CUOI
        pos, last = 0, -1
        while True:
            m = re.search(re.escape(title) + r"\s+[a-z]", s[pos:])
            if not m:
                break
            last = pos + m.start()
            pos = last + len(title)
        if last >= 0:
            return s[last:].strip()
    # Khong co ten SP thi KHONG cat theo chu hoa: nhieu cau mo dau bang ten
    # thuong ("perma NOVA was ...") -> cat se mat chu "perma".
    return s


def ntrend_cats():
    """Danh muc SP tren web Ntrend, sheet '05 Perma' (review 22/10)."""
    try:
        from openpyxl import load_workbook
        p = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
        ws = load_workbook(p, read_only=True, data_only=True)["05 Perma"]
    except Exception:
        return {}
    out = {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or len(r) < 9 or str(r[7] or "").strip() != "x":
            continue
        nm = (r[1] or "").strip()
        cat = (r[8] or "").strip()
        if nm and cat:
            out[nm.lower()] = cat
    return out


NTREND_CAT = {}
_IPN = None


def name_vi_of(name):
    """Ten dong SP: la ma/ten rieng (perma NOVA, STAR) thi giu, cum tu thi dich."""
    global _IPN
    if _IPN is None:
        src = io.open(os.path.join(HERE, "c_sick.py"), encoding="utf-8").read()
        ns = {"re": re}
        exec(src[src.index("_COMMON = set"):src.index("IMG_EXT")], ns)
        _IPN = ns["is_proper_name"]
    if _IPN(name) or name.lower().startswith("perma"):
        return name
    return vi(name)


def srcset_largest(tag):
    ss = re.search(r'srcset="([^"]+)"', tag)
    if ss:
        best, bw = None, -1
        for part in ss.group(1).split(","):
            bits = part.strip().split()
            if not bits:
                continue
            w = int(re.sub(r"\D", "", bits[1]) or 0) if len(bits) > 1 else 0
            if w > bw:
                best, bw = bits[0], w
        if best:
            return best
    m = re.search(r'src="([^"]+)"', tag)
    return m.group(1) if m else None


def extract(t):
    """overview paragraphs, feature bullets, spec rows, hero image url."""
    body = re.sub(r"<script.*?</script>|<style.*?</style>", "", t, flags=re.S)
    # chi quet trong <main> de khoi dinh nav/footer (menu, terms, ...)
    mo = re.search(r"<main[\s>]", body)
    if mo:
        end = body.find("</main>", mo.start())
        region = body[mo.start():end if end > 0 else len(body)]
    else:
        region = body
    h1 = re.search(r"<h1[^>]*>(.*?)</h1>", region, re.S)
    main = region[h1.end():] if h1 else region
    title = txt(h1.group(1)) if h1 else ""

    paras = []
    for p in re.findall(r"<p[^>]*>(.*?)</p>", region, re.S):
        s = clean_para(txt(p), title)
        if len(s) < 50 or any(b.lower() in s.lower() for b in BAD_PARA):
            continue
        if s not in paras:
            paras.append(s)
        if len(paras) >= 5:
            break

    feats = []
    for x in re.findall(r"<li[^>]*>(.*?)</li>", main, re.S):
        if "<a " in x:  # li chua link = dieu huong, khong phai tinh nang
            continue
        s = txt(x)
        if not (15 < len(s) < 180):
            continue
        if any(b in s.lower() for b in BAD_LI):
            continue
        if s not in feats:
            feats.append(s)
        if len(feats) >= 10:
            break

    specs = []
    tb = re.search(r"<table[^>]*>(.*?)</table>", main, re.S)
    if tb:
        for row in re.findall(r"<tr[^>]*>(.*?)</tr>", tb.group(1), re.S):
            cells = [txt(c) for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, re.S)]
            if len(cells) >= 2 and cells[0] and cells[1]:
                specs.append([cells[0], cells[1]])
            if len(specs) >= 25:
                break

    img = None
    st = re.search(r'data-media-index="0"[^>]*>\s*(<img[^>]+>)', body)
    if st:
        img = srcset_largest(st.group(1))
    if not img:
        for m in re.finditer(r"<img[^>]+>", main):
            u = srcset_largest(m.group(0))
            if u and "Persistent" in u and "logo" not in u.lower():
                img = u
                break
    return title, paras, feats, specs, img


# tieu de item: trang Angles dung "text-lg text-black", trang INSTALLATION KIT
# dung "text-lg font-medium" -> bat chung "text-lg"; voi moi Art. No. lay tieu
# de GAN NHAT phia truoc de khoi vo nham div menu.
TITLE_RX = re.compile(r'<div class="text-lg[^"]*"[^>]*>(.*?)</div>', re.S)
ART_RX = re.compile(r"Art\.\s*No\.?:?\s*(\d{4,7})")


def items_fallback(t):
    """Trang phu kien: tra ve list dict {name, material, art} cho bang ma SP."""
    body = re.sub(r"<script.*?</script>|<style.*?</style>", "", t, flags=re.S)
    titles = [(m.end(), txt(m.group(1))) for m in TITLE_RX.finditer(body)]
    out, seen = [], set()
    for am in ART_RX.finditer(body):
        art = am.group(1)
        if art in seen:
            continue
        prev = [tt for tt in titles if tt[0] <= am.start()
                and am.start() - tt[0] < 1200]
        if not prev:
            continue
        end, name = prev[-1]
        if not name or len(name) > 120:
            continue
        seen.add(art)
        mat = re.search(r"Material:\s*([^<;,]{2,40})", txt(body[end:am.start()]))
        out.append({"name": name,
                    "material": mat.group(1).strip() if mat else "",
                    "art": art})
        if len(out) >= 30:
            break
    return out


def download(url, dst):
    if os.path.exists(dst) and os.path.getsize(dst) > 500:
        return True
    p = subprocess.run(["curl", "-sS", "-L", "--max-time", "60", "-o", dst, url],
                       capture_output=True, timeout=90)
    return p.returncode == 0 and os.path.exists(dst) and os.path.getsize(dst) > 500


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    with io.open(CAT_JSON, encoding="utf-8") as f:
        d = json.load(f)
    nodes = d["nodes"]
    leaves = [n for n in nodes
              if not any(len(m["path"]) == len(n["path"]) + 1 and
                         m["path"][:len(n["path"])] == n["path"] for m in nodes)]
    globals()["NTREND_CAT"] = ntrend_cats()
    print("bai can lam: %d | danh muc Ntrend: %d" % (len(leaves), len(NTREND_CAT)))

    arts, miss_img, miss_txt = [], [], []
    for i, n in enumerate(leaves, 1):
        name = n["path"][-1]
        parent = n["path"][-2] if len(n["path"]) > 1 else "Accessories"
        slug = slugify(name)
        t = try_fetch(n["url"], retries=3)
        if not t:
            print("  [FAIL]", n["url"])
            continue
        title, paras, feats, specs, img = extract(t)
        items = items_fallback(t)
        if not paras and not items:
            miss_txt.append(name)

        # anh
        img_file = None
        if img:
            if img.startswith("/"):
                img = "https://www.perma-tec.com" + img
            ext = os.path.splitext(img.split("?")[0])[1] or ".jpg"
            src = os.path.join(SRC_IMG, slug + ext)
            if download(img, src):
                try:
                    img_file = brandimg.compose(src, os.path.join(BRAND_IMG, slug + ".png"))
                except Exception as e:
                    print("  [img fail]", name, repr(e))
        if not img_file:
            miss_img.append(name)

        type_en = TYPE_EN.get(parent, TYPE_EN.get(n["path"][1] if len(n["path"]) > 1 else "", "Product"))
        type_vi = TYPE_VI.get(parent, TYPE_VI.get(n["path"][1] if len(n["path"]) > 1 else "", "Sản phẩm"))

        for it in items:                       # ten + vat lieu trong bang ma SP
            it["name_vi"] = vi(it.get("name", ""))
            it["material_vi"] = vi(it.get("material", "").title())

        art = {
            "name": name, "slug": slug, "url_src": n["url"],
            "name_vi": name_vi_of(name),
            "cat_vi": NTREND_CAT.get(name.strip().lower(), ""),
            "path": n["path"], "brand": "perma",
            "type_en": type_en, "type_vi": type_vi,
            "overview_en": paras, "overview_vi": [vi(p) for p in paras],
            "features_en": feats, "features_vi": [vi(x) for x in feats],
            "items": items,
            "specs_en": specs,
            "specs_vi": [[vi_spec_key(k), vi(v)] for k, v in specs],
            "path_vi": [TYPE_VI.get(x) or vi(x) for x in n["path"]],
            "img_src_url": img, "img_file": img_file,
        }
        arts.append(art)
        print("  %2d/%d %-55s paras=%d feats=%d specs=%d img=%s"
              % (i, len(leaves), name[:55], len(paras), len(feats), len(specs),
                 "OK" if img_file else "--"))

    _flush(True)
    with io.open(os.path.join(OUT_DIR, "content.json"), "w", encoding="utf-8") as f:
        json.dump({"brand": "05 Perma", "articles": arts}, f, ensure_ascii=False, indent=1)
    print("\nxong:", len(arts), "bai | thieu anh:", len(miss_img), "| thieu mo ta:", len(miss_txt))
    if miss_img:
        print("  thieu anh:", miss_img[:10])
    if miss_txt:
        print("  thieu mo ta:", miss_txt[:10])


if __name__ == "__main__":
    main()
