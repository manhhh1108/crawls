# -*- coding: utf-8 -*-
"""Lap ban do 'dong SP trong danh muc Ntrend' -> URL tren web hang.

    python cat_map.py espe balluff otennlux festo
    python cat_map.py                       # ca 4 hang

Khac voi wenglor_map/leuze_map/ifm_map/mitsu_map (phai tu bo web hang de tim
URL): 4 hang nay da co san cay danh muc KEM URL trong catalog/out/<hang>.json
do nhom catalog/b_*.py dung truoc day -> chi can doi chieu duong dan danh muc.

Ket qua: content/out/<hang>/urlmap.json, dung chung dinh dang voi cac hang cu
(stt / name / path / cat_vi / url) cong them "code" va "kind".
"""
import collections
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
NTREND_XLSX = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
CATALOG = os.path.join(REPO, "catalog", "out")

# ten viet tat -> (sheet trong file danh muc Ntrend = ten file catalog, ten hang)
BRANDS = {
    "festo":    ("10 Festo", "Festo"),
    "espe":     ("15 ESPE Technology", "ESPE"),
    "otennlux": ("16 OTENNLUX Lighting", "OTENNLUX"),
    "balluff":  ("17 Balluff", "Balluff"),
}


def norm(s):
    return re.sub(r"[^0-9a-z]+", "", (s or "").lower())


def wanted(sheet, brand_vi):
    """Doc cac dong SP can crawl tu mot sheet cua file danh muc Ntrend.

    Cac sheet KHONG cung so cot: Festo co Cap 1-6, ESPE chi Cap 1-2, va cot
    danh dau co sheet ghi 'La cap ngay tren ma SP', sheet khac ghi 'SP can lay
    noi dung' -> do dong tieu de (dong 5) chu khong ghim chi so cot.
    """
    from openpyxl import load_workbook
    ws = load_workbook(NTREND_XLSX, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(min_row=5, values_only=True))
    hdr = [str(x or "") for x in rows[0]]
    lv = [i for i, h in enumerate(hdr) if h.startswith("Cấp ")]
    mk = next(i for i, h in enumerate(hdr)
              if h.startswith("Là cấp") or h.startswith("SP cần"))
    nt = next(i for i, h in enumerate(hdr) if h.startswith("Danh mục SP"))

    out = []
    for r in rows[1:]:
        if not r or not r[0] or len(r) <= mk:
            continue
        if str(r[mk] or "").strip().lower() != "x":
            continue
        path = [str(r[i]).strip() for i in lv
                if i < len(r) and r[i] and str(r[i]).strip()]
        if not path:
            continue
        cat = str(r[nt] or "").strip() if nt < len(r) else ""
        # Cot danh muc chi ghi moi ten hang = chua phan loai tren web Ntrend.
        if cat == brand_vi or norm(cat) == norm(brand_vi):
            cat = ("Không có danh mục sản phẩm,  chỉ nằm trong danh mục "
                   "Thương Hiệu " + brand_vi)
        out.append({"stt": str(r[0]).strip(), "name": path[-1], "path": path,
                    "cat_vi": cat})
    return out


def run(key):
    sheet, brand_vi = BRANDS[key]
    doc = json.load(io.open(os.path.join(CATALOG, sheet + ".json"),
                            encoding="utf-8"))
    by_path, by_name = {}, collections.defaultdict(list)
    for n in doc["nodes"]:
        if not n.get("url") or not n.get("path"):
            continue
        by_path["/".join(norm(x) for x in n["path"])] = n
        by_name[norm(n["path"][-1])].append(n)

    lines = wanted(sheet, brand_vi)
    hit = 0
    for l in lines:
        n = by_path.get("/".join(norm(x) for x in l["path"]))
        if n is None:
            # du phong: khop theo ten la cuoi, uu tien ban trung nhieu cap cha
            cand = by_name.get(norm(l["name"])) or []
            if len(cand) > 1:
                par = {norm(x) for x in l["path"][:-1]}
                cand.sort(key=lambda c: -len(par & {norm(x) for x in c["path"]}))
            n = cand[0] if cand else None
        if n is None:
            continue
        l["url"] = n["url"]
        l["code"] = n.get("code") or ""
        l["kind"] = n.get("type") or ""
        hit += 1

    out_dir = os.path.join(HERE, "out", sheet)
    os.makedirs(out_dir, exist_ok=True)
    with io.open(os.path.join(out_dir, "urlmap.json"), "w",
                 encoding="utf-8") as f:
        json.dump({"lines": lines}, f, ensure_ascii=False)
    uniq = len({l["url"] for l in lines if l.get("url")})
    print("%-9s %-22s dong=%-5d khop=%-5d trang khac nhau=%-5d chua khop=%d"
          % (key, sheet, len(lines), hit, uniq, len(lines) - hit))
    for l in lines:
        if not l.get("url"):
            print("   [?]", " > ".join(l["path"]))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    for k in [a for a in sys.argv[1:] if a in BRANDS] or list(BRANDS):
        run(k)


if __name__ == "__main__":
    main()
