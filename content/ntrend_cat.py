# -*- coding: utf-8 -*-
"""Danh muc Ntrend cho tung dong SP SICK - NGUON DUY NHAT.

File goc: data/Danh muc SP web NTREND.xlsx, sheet "01 SICK"
          cot B = ten dong SP (EN), cot D = danh mục Ntrend (VI)

Ntrend tu sua truc tiep trong file xlsx do; chay lai build_sick_excel.py la
noi dung cap nhat theo, KHONG phai crawl lai.

(Ban cu lay o cot "Danh muc SP Tieng Viet" cua data/SICK - dong SP can crawl.xlsx
 - gio chi con dung lam du phong khi ten dong SP khong co trong sheet 01 SICK.)
"""
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
XLSX = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
SHEET = "01 SICK"

NO_CAT = "Không có danh mục sản phẩm, chỉ nằm trong danh mục Thương Hiệu SICK Sensor"

# Ban tieng Anh cua danh muc Ntrend (chi dung o cot C - soat chinh ta).
CAT_EN = {
    "Cảm biến công nghiệp": "Industrial Sensors",
    "Cảm biến công nghiệp, Cảm biến dạng fork": "Industrial Sensors, Fork Sensors",
    "Cảm biến công nghiệp, Cảm biến lưu lượng": "Industrial Sensors, Flow Sensors",
    "Cảm biến công nghiệp, Cảm biến màu và tương phản":
        "Industrial Sensors, Color and Contrast Sensors",
    "Cảm biến công nghiệp, Cảm biến mức": "Industrial Sensors, Level Sensors",
    "Cảm biến công nghiệp, Cảm biến quang dạng vùng":
        "Industrial Sensors, Area Sensors (Light Grids)",
    "Cảm biến công nghiệp, Cảm biến quang điện":
        "Industrial Sensors, Photoelectric Sensors",
    "Cảm biến công nghiệp, Cảm biến quét laser":
        "Industrial Sensors, Laser Scanner Sensors",
    "Cảm biến công nghiệp, Cảm biến radar": "Industrial Sensors, Radar Sensors",
    "Cảm biến công nghiệp, Cảm biến sợi quang":
        "Industrial Sensors, Fiber Optic Sensors",
    "Cảm biến công nghiệp, Cảm biến tiệm cận": "Industrial Sensors, Proximity Sensors",
    "Cảm biến công nghiệp, Cảm biến từ xi lanh":
        "Industrial Sensors, Cylinder Magnetic Sensors",
    "Cảm biến công nghiệp, Cảm biến áp suất": "Industrial Sensors, Pressure Sensors",
    "Cảm biến công nghiệp, Cảm biến đo khoảng cách":
        "Industrial Sensors, Distance Measurement Sensors",
    "Cảm biến màu và tương phản": "Industrial Sensors, Color and Contrast Sensors",
    "Encoder bộ mã hóa vòng quay": "Rotary Encoders",
    "Encoder bộ mã hóa vòng quay, Encoder bánh xe":
        "Rotary Encoders, Measuring Wheel Encoders",
    "Encoder bộ mã hóa vòng quay, Encoder dây kéo":
        "Rotary Encoders, Draw-Wire Encoders",
    "Encoder bộ mã hóa vòng quay, Encoder tuyến tính": "Rotary Encoders, Linear Encoders",
    "Encoder bộ mã hóa vòng quay, Encoder tuyệt đối":
        "Rotary Encoders, Absolute Encoders",
    "Encoder bộ mã hóa vòng quay, Encoder tương đối":
        "Rotary Encoders, Incremental Encoders",
    "Thiết bị an toàn máy": "Machine Safety Devices",
    "Thiết bị an toàn máy, Bộ điều khiển an toàn":
        "Machine Safety Devices, Safety Controllers",
    "Thiết bị an toàn máy, Camera 3D an toàn": "Machine Safety Devices, 3D Safety Cameras",
    "Thiết bị an toàn máy, Công tắc an toàn": "Machine Safety Devices, Safety Switches",
    "Thiết bị an toàn máy, Lưới quang an toàn":
        "Machine Safety Devices, Safety Light Curtains",
    "Thiết bị an toàn máy, Quét laser an toàn":
        "Machine Safety Devices, Safety Laser Scanners",
    "Thiết bị an toàn máy, Radar an toàn": "Machine Safety Devices, Safety Radar Sensors",
    "Thiết bị an toàn máy, Rơ le an toàn": "Machine Safety Devices, Safety Relays",
    "Thiết bị phân tích khí và bụi": "Gas and Dust Analyzers",
    "Vision và Nhận dạng, Camera công nghiệp 2D/3D":
        "Vision & Identification, 2D/3D Industrial Cameras",
    "Vision và Nhận dạng, Máy đọc mã vạch": "Vision & Identification, Barcode Scanners",
    "Vision và Nhận dạng, Đầu đọc RFID": "Vision & Identification, RFID Readers",
    "Bộ kết nối mạng công nghiệp, Thiết bị kết nối mạng công nghiệp":
        "Industrial Network Connectivity, Industrial Network Connectivity Devices",
    "Bộ kết nối mạng công nghiệp, Đầu nối và cáp":
        "Industrial Network Connectivity, Connectors & Cables",
    "Đo và Giám sát nhiệt độ, Cảm biến nhiệt độ":
        "Temperature Measurement & Monitoring, Temperature Sensors",
    NO_CAT: NO_CAT,
}

_CACHE = None


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").replace("\xa0", " ")).strip()


def _load():
    global _CACHE
    if _CACHE is not None:
        return _CACHE
    from openpyxl import load_workbook
    out = {}
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    try:
        ws = wb[SHEET]
        for r in ws.iter_rows(values_only=True):
            if not r or len(r) < 4 or r[0] is None or not r[1]:
                continue
            if not re.match(r"^[\d.]+$", str(r[0]).strip()):   # bo dong tieu de
                continue
            out[_norm(r[1])] = _norm(r[3])
    finally:
        wb.close()
    _CACHE = out
    return out


def cat_vi(name, fallback=""):
    """Danh muc Ntrend (tieng Viet) cua 1 dong SP."""
    v = _load().get(_norm(name))
    if v:
        return v
    return _norm(fallback)


def cat_en(vi, fallback=""):
    """Ban tieng Anh cua danh muc (chi de soat o cot C)."""
    return CAT_EN.get(_norm(vi), _norm(fallback))


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    d = _load()
    print("%s / sheet %s: %d dong" % (os.path.basename(XLSX), SHEET, len(d)))
    missing = sorted({v for v in d.values() if v and v not in CAT_EN})
    print("danh muc chua co ban EN:", missing or "(khong co)")
    blank = [k for k, v in d.items() if not v]
    print("dong chua phan loai:", blank or "(khong co)")
