# -*- coding: utf-8 -*-
"""Reer (hang 03) - crawl noi dung bai dang cho tung dong san pham.

    python c_reer.py            # chay tiep nhung dong chua co
    python c_reer.py 5          # chi 5 dong dau (xem thu)
    python c_reer.py --force    # lam lai tu dau

Khac SICK/P+F: reersafety.com la WooCommerce, KHONG chan curl -> tai bang HTTP
thuong, khong can Chrome. Phan trang /page/N/, 12 san pham moi trang.

Moi san pham tren trang danh muc co san bang thuoc tinh rieng:
    <table class="shop_attributes"><tr><th>Safety level</th><td>Type 4 - PL e</td>...
Ten cot khac nhau theo tung dong SP -> bang ma san pham co COT DONG (dung mau
trong Reer_template.docx).

Ket qua: out/03 Reer/content.json
"""
import io
import json
import os
import re
import sys
import html as H

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, REPO)

import out_dir                                         # noqa: E402

from shared.fetcher import try_fetch                 # noqa: E402

BRAND = "Reer"
SITE = "https://www.reersafety.com"
OUT = os.path.join(HERE, "out", "03 Reer")
CONTENT = os.path.join(OUT, "content.json")
SRC_IMG = os.path.join(OUT, "src_img")
XLSX_DIR = out_dir.brand_dir("Noi dung Reer")
IMG_DIR = os.path.join(XLSX_DIR, "images")
for d in (OUT, SRC_IMG, IMG_DIR):
    os.makedirs(d, exist_ok=True)

NTREND = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
FULL = os.path.join(REPO, "data", "Danh muc san pham 16 hang - FULL.xlsx")

MAX_PAGE = 40
MIN_HL, MAX_HL = 5, 8          # Mo ta ngan: 5-8 thong so


def txt(s):
    return H.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", s or ""))).strip()


# ---------------------------------------------------------------------------
# 1. Danh sach dong SP
# ---------------------------------------------------------------------------
def lines():
    """171 dong SP: ten + danh muc Ntrend (file Ntrend) + URL (file FULL)."""
    from openpyxl import load_workbook
    url_by = {}
    ws = load_workbook(FULL, read_only=True, data_only=True)["03 Reer"]
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or len(r) < 10 or str(r[8] or "").strip() != "x":
            continue
        url_by[str(r[0]).strip()] = (r[9] or "").strip()

    ws = load_workbook(NTREND, read_only=True, data_only=True)["03 Reer"]
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or len(r) < 10 or str(r[8] or "").strip() != "x":
            continue
        stt = str(r[0]).strip()
        name = (r[1] or "").strip()
        if not name:
            continue
        cat = (r[9] or "").strip()
        if cat == BRAND:
            cat = ("Không có danh mục sản phẩm,  chỉ nằm trong danh mục "
                   "Thương Hiệu Reer")
        out.append({"stt": stt, "name": name, "cat_vi": cat,
                    "url": url_by.get(stt, ""),
                    "path": [str(x).strip() for x in r[3:8] if x and str(x).strip()]})
    return [l for l in out if l["url"]]


# ---------------------------------------------------------------------------
# 2. Boc san pham tren trang danh muc
# ---------------------------------------------------------------------------
BLOCK = re.compile(r'<div class="product-small col has-hover(.*?)(?=<div class="product-small col has-hover|<footer)', re.S)
TITLE = re.compile(r'woocommerce-loop-product__title"><a[^>]*href="([^"]+)"[^>]*>(.*?)</a>', re.S)
SKU = re.compile(r'<div class="sku_wrapper">\s*(.*?)</div>', re.S)
EXC = re.compile(r'box-excerpt[^>]*>(.*?)</p>', re.S)
ATTR_TB = re.compile(r'<table class="shop_attributes">(.*?)</table>', re.S)
ATTR_ROW = re.compile(r"<tr><th>(.*?)</th><td>(.*?)</td></tr>", re.S)
IMG = re.compile(r'data-src="([^"]+)"')
# "Showing 1-12 of 75 results"  hoac  "Showing all 7 results" (chi 1 trang)
PAGE_H1 = re.compile(r'<h1[^>]*>(.*?)</h1>', re.S)
COUNT = re.compile(r"of\s+([\d,]+)\s+results")
COUNT_ALL = re.compile(r"Showing all\s+([\d,]+)\s+results")


def parse_items(html):
    out = []
    for blk in BLOCK.findall(html):
        t = TITLE.search(blk)
        if not t:
            continue
        s = SKU.search(blk)
        e = EXC.search(blk)
        tb = ATTR_TB.search(blk)
        attrs = []
        if tb:
            for k, v in ATTR_ROW.findall(tb.group(1)):
                k, v = txt(k), txt(v)
                if k and v:
                    attrs.append([k, v])
        im = IMG.search(blk)
        out.append({"name": txt(t.group(2)), "url": t.group(1),
                    "art": txt(s.group(1)) if s else "",
                    "excerpt": txt(e.group(1)) if e else "",
                    "attrs": attrs,
                    "img": im.group(1) if im else ""})
    return out


def page_title(html):
    """<h1> cua trang = TEN DAY DU cua dong SP (review 22/10).

    Danh muc chi luu ten rut gon ('30 mm beam spacing'), con web ghi day du
    ('Measuring and automation light curtains with 30 mm beam spacing').
    """
    for m in PAGE_H1.findall(html):
        t = txt(m)
        # Loai breadcrumb ("Home / Safety Light Curtains / EOS4 / ...") - dau '/'
        # CO KHOANG TRANG hai ben. Ten that van co the chua '/' lien ("Input/Output",
        # "Admiral/Vision") nen khong duoc loai het theo dau '/'.
        if not t or len(t) <= 3:
            continue
        if t.lower().startswith("home") or " / " in t:
            continue
        return t
    return ""


def fetch_line(url):
    """Tai het cac trang /page/N/ cua mot dong SP."""
    items, seen, total, h1 = [], set(), None, ""
    for page in range(1, MAX_PAGE + 1):
        u = url if page == 1 else url.rstrip("/") + "/page/%d/" % page
        html = try_fetch(u, retries=3, timeout=60)
        if not html:
            break
        if not h1:
            h1 = page_title(html)
        if total is None:
            t = txt(html)
            m = COUNT.search(t) or COUNT_ALL.search(t)
            total = int(m.group(1).replace(",", "")) if m else None
        got = parse_items(html)
        new = [x for x in got if (x["art"] or x["name"]) not in seen]
        for x in new:
            seen.add(x["art"] or x["name"])
        items += new
        if not new or (total and len(items) >= total):
            break
    return items, total, h1


# ---------------------------------------------------------------------------
# 3. Mo ta dai / ngan tu bang thuoc tinh
# ---------------------------------------------------------------------------
def groups(items):
    import collections
    g = collections.defaultdict(collections.Counter)
    for it in items:
        for k, v in it.get("attrs", []):
            g[k][v] += 1
    return sorted(g.items(), key=lambda kv: -sum(kv[1].values()))


def fmt(k, cnt, max_vals=6):
    vals = [v for v, _c in cnt.most_common(max_vals)]
    more = len(cnt) - len(vals)
    return "%s: %s%s" % (k, "; ".join(vals),
                         " (+%d tuỳ chọn khác)" % more if more > 0 else "")


def long_desc(items):
    out = [fmt(k, c) for k, c in groups(items)]
    if out:
        return out
    seen = []
    for it in items:                       # khong co bang thuoc tinh -> lay excerpt
        e = (it.get("excerpt") or "").strip()
        if e and e not in seen:
            seen.append(e)
    return seen[:15]


def attr_columns(items):
    """Ten cac cot thong so cho bang ma SP, theo thu tu pho bien nhat.

    Dong khong co <table class="shop_attributes"> (vd Test Rods) thi web van in
    mot cau mo ta cho tung ma ('Test rod, Diameter 14 mm') -> dung lam 1 cot.
    """
    cols = [k for k, _c in groups(items)]
    if not cols and any((it.get("excerpt") or "").strip() for it in items):
        return ["Description"]
    return cols


# ---------------------------------------------------------------------------
# 4. Ten tieng Viet + anh
# ---------------------------------------------------------------------------
_IPN = None


def is_code_name(s):
    """Ten dong SP la ma/ten rieng cua Reer (EOS4, Janus, Admiral) hay cum tu
    thuong (Safety Light Curtains, Brackets)? Dung lai ham cua c_sick.py."""
    global _IPN
    if _IPN is None:
        src = io.open(os.path.join(HERE, "c_sick.py"), encoding="utf-8").read()
        ns = {"re": re}
        exec(src[src.index("_COMMON = set"):src.index("IMG_EXT")], ns)
        _IPN = ns["is_proper_name"]
    return _IPN(s)


def name_vi_of(name, translate_many):
    if is_code_name(name):
        return name
    from vi_engine import lookup
    return lookup(name) or translate_many([name], "key")[0]


def grab_image(items, line_en, slugify):
    import urllib.request
    url = next((it.get("img") for it in items if it.get("img")), "")
    if not url:
        return ""
    base = slugify(line_en)
    ext = os.path.splitext(url.split("?")[0])[1][:5] or ".png"
    src = os.path.join(SRC_IMG, base + ext)
    try:
        if not (os.path.exists(src) and os.path.getsize(src) > 500):
            req = urllib.request.Request(url, headers={"User-Agent":
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/131.0.0.0"})
            data = urllib.request.urlopen(req, timeout=60).read()
            if len(data) < 500:
                return ""
            with open(src, "wb") as f:
                f.write(data)
        import brandimg
        return os.path.basename(
            brandimg.compose(src, os.path.join(IMG_DIR, base + ".jpg")))
    except Exception as e:
        print("     [anh loi]", repr(e)[:60])
        return ""


def save(ls, done):
    arts, seen = [], set()
    for l in ls:
        u = l["url"]
        if u in done and u not in seen:
            seen.add(u)
            arts.append(done[u])
    with io.open(CONTENT, "w", encoding="utf-8") as f:
        json.dump({"brand": "03 Reer", "articles": arts}, f, ensure_ascii=False)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    force = "--force" in sys.argv
    limit = next((int(a) for a in sys.argv[1:] if a.isdigit()), 0)

    from c_pepperl_try import slugify
    from vi_engine import translate_many
    from translate import _flush

    ls = lines()
    if limit:
        ls = ls[:limit]
    # bo trung theo URL: cung trang = cung noi dung
    uniq, seen = [], set()
    for l in ls:
        if l["url"] in seen:
            continue
        seen.add(l["url"])
        uniq.append(l)
    done = {}
    if os.path.exists(CONTENT) and not force:
        done = {a["url_src"]: a for a in
                json.load(io.open(CONTENT, encoding="utf-8"))["articles"]
                if a.get("url_src")}
    todo = [l for l in uniq if l["url"] not in done]
    print("dong SP: %d | bo trung cung trang: %d | bai se co: %d"
          % (len(ls), len(ls) - len(uniq), len(uniq)))
    print("da xong: %d | se lam: %d" % (len(done), len(todo)))

    for i, ln in enumerate(todo, 1):
        items, total, h1 = fetch_line(ln["url"])
        if not items:
            print("  %3d/%d %-34s KHONG CO MA SP" % (i, len(todo), ln["name"][:34]))
        cols = attr_columns(items)
        long_en = long_desc(items)
        short_en = long_en[:MAX_HL]
        img = grab_image(items, ln["name"], slugify)

        d_en = [it.get("excerpt", "") for it in items]
        vals = sorted({v for it in items for _k, v in it.get("attrs", [])})
        tr_long = translate_many(long_en, "value") if long_en else []
        tr_cols = translate_many(cols, "key") if cols else []
        tr_vals = dict(zip(vals, translate_many(vals, "value"))) if vals else {}
        tr_exc = dict(zip(d_en, translate_many(d_en, "prose"))) if any(d_en) else {}
        for it in items:
            it["excerpt_vi"] = tr_exc.get(it.get("excerpt", ""), "")
            it["attrs_vi"] = [[k, tr_vals.get(v, v)] for k, v in it.get("attrs", [])]
            if cols == ["Description"]:
                it["attrs"] = [["Description", it.get("excerpt", "")]]
                it["attrs_vi"] = [["Description", it.get("excerpt_vi", "")]]

        done[ln["url"]] = dict(ln, url_src=ln["url"], items=items, img_file=img,
                               cols_en=cols, cols_vi=tr_cols,
                               long_en=long_en, long_vi=tr_long,
                               short_en=short_en, short_vi=tr_long[:len(short_en)],
                               title_en=h1 or ln["name"],
                               name_vi=name_vi_of(h1 or ln["name"], translate_many))
        print("  %3d/%d %-34s ma=%-4d cot=%-2d ts=%-2d %s"
              % (i, len(todo), ln["name"][:34], len(items), len(cols),
                 len(long_en), "img" if img else "---"))
        if i % 5 == 0:
            _flush(True)
            save(uniq, done)
    _flush(True)
    save(uniq, done)
    print("\nxong %d bai -> %s" % (len(done), CONTENT))


if __name__ == "__main__":
    main()
