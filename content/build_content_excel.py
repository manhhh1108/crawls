# -*- coding: utf-8 -*-
"""Dung file noi dung theo review Ntrend: MOI BAI 1 FILE Excel trong 1 folder.
Layout: cot A = truong, cot B = TIENG VIET (ban dang web), cot C = tieng Anh
(chi de soat chinh ta ban dich). Mau van lay tu 'Crawl noi dung Ntrend.docx'."""
import io, json, os, re, shutil, sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
FIELD_FILL = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(bold=True, color="FFFFFF")
FIELD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

BRAND_FULL_VI = "perma-tec (Đức)"
BRAND_FULL_EN = "perma-tec (Germany)"

# ---------------- mau van theo review (docx) ----------------
LONG_TPL_VI = (
    "{type_c} {name_brand} là sản phẩm trong công nghiệp được sử dụng trong các "
    "hệ thống tự động hóa và điều khiển công nghiệp, đáp ứng các yêu cầu về độ "
    "chính xác, độ ổn định và độ tin cậy trong quá trình vận hành. Sản phẩm được "
    "thiết kế phù hợp với nhiều ứng dụng trong nhà máy và dây chuyền sản xuất, "
    "hỗ trợ đo lường, phát hiện, giám sát, điều khiển và tối ưu hóa quy trình.\n\n"
    "{type_c} {name} được sản xuất theo các tiêu chuẩn kỹ thuật của {brand}, với "
    "các thông số kỹ thuật được công bố rõ ràng, giúp kỹ sư và bộ phận kỹ thuật "
    "dễ dàng lựa chọn, tích hợp và triển khai vào hệ thống tự động hóa.\n\n"
    "Tùy thuộc vào từng model, sản phẩm có thể có các tùy chọn khác nhau đáp ứng "
    "đa dạng yêu cầu của các ứng dụng công nghiệp.")
LONG_TPL_EN = (
    "The {name_brand} {type_l} is an industrial product used in industrial "
    "automation and control systems, meeting requirements for precision, "
    "stability and reliability in operation. It is designed for a wide range of "
    "applications in factories and production lines, supporting measurement, "
    "detection, monitoring, control and process optimisation.\n\n"
    "The {name} {type_l} is manufactured to {brand}'s technical standards, with "
    "clearly published specifications that make it easy for engineers and "
    "technical teams to select, integrate and deploy in automation systems.\n\n"
    "Depending on the model, different options are available to meet the varied "
    "requirements of industrial applications.")

SHORT_TPL_VI = "{type_c} {name} là sản phẩm trong công nghiệp thuộc hãng {brand}"
SHORT_TPL_EN = "The {name} {type_l} is an industrial product from {brand}"

SPECDESC_TPL_VI = (
    "{type_c} {name_brand} là thiết bị công nghiệp được sử dụng trong các hệ "
    "thống tự động hóa và điều khiển công nghiệp, đáp ứng các yêu cầu về độ "
    "chính xác, độ ổn định và độ tin cậy trong quá trình vận hành. Sản phẩm được "
    "thiết kế phù hợp với nhiều ứng dụng trong nhà máy và dây chuyền sản xuất, "
    "hỗ trợ đo lường, phát hiện, giám sát, điều khiển và tối ưu hóa quy trình.\n\n"
    "{type_c} {name} được sản xuất theo các tiêu chuẩn kỹ thuật của {brand}, với "
    "các thông số kỹ thuật được công bố rõ ràng, giúp kỹ sư và bộ phận kỹ thuật "
    "dễ dàng lựa chọn, tích hợp và triển khai vào hệ thống tự động hóa.\n\n"
    "Tùy thuộc vào từng model, sản phẩm có thể có các tùy chọn khác nhau về kích "
    "thước, phạm vi hoạt động, nguồn cấp, tín hiệu đầu ra, giao diện kết nối, "
    "cấp bảo vệ và điều kiện môi trường, đáp ứng đa dạng yêu cầu của các ứng "
    "dụng công nghiệp.\n\n"
    "Dưới đây là bảng thông số kỹ thuật chi tiết của sản phẩm:")
SPECDESC_TPL_EN = (
    "The {name_brand} {type_l} is an industrial device used in industrial "
    "automation and control systems, meeting requirements for precision, "
    "stability and reliability in operation. It is designed for a wide range of "
    "applications in factories and production lines.\n\n"
    "The {name} {type_l} is manufactured to {brand}'s technical standards with "
    "clearly published specifications.\n\n"
    "Depending on the model, options vary in size, operating range, power "
    "supply, output signal, connection interface, protection class and ambient "
    "conditions.\n\nDetailed specifications are listed below:")

NOSPEC_DESC_VI = ("{type_c} {name_brand} là sản phẩm công nghiệp chính hãng của "
                  + BRAND_FULL_VI + "; bảng mã sản phẩm chi tiết như bảng dưới:")
NOSPEC_DESC_EN = ("The {name_brand} {type_l} is a genuine industrial product from "
                  + BRAND_FULL_EN + "; the detailed part-number table is below:")

META_TPL_VI = ("Mua {type_l} {name_brand} chính hãng – thông số kỹ thuật đầy đủ, "
               "giá tốt, hỗ trợ tư vấn và lựa chọn model phù hợp. "
               "Liên hệ Ntrend ngay!")
META_TPL_EN = ("Buy genuine {name_brand} {type_l} – full specifications, good "
               "price, support in selecting the right model. "
               "Contact Ntrend now for a consultation!")

ANGLE_RX = re.compile(r"^Angle\s+([\d.,]+°)\s+(.+?)\s+male\s*x\s*(.+?)\s+female$", re.I)


def with_brand(a):
    if a["brand"].lower() in a["name"].lower():
        return a["name"]
    return "%s %s" % (a["name"], a["brand"])


def vn_slug(s):
    """Keyword chinh (tieng Viet) -> URL: bo dau, thuong hoa, noi bang '-'."""
    import unicodedata
    s = (s or "").replace("Đ", "D").replace("đ", "d")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).lower().strip("-")
    return re.sub(r"-{2,}", "-", s) or "sp"


def ctx(a):
    return {"name": a["name"], "name_brand": with_brand(a), "brand": a["brand"],
            "type_c": a["type_vi"][:1].upper() + a["type_vi"][1:],
            "type_l": a["type_en"].lower()}


def with_brand_vi(a):
    nv = a.get("name_vi") or a["name"]
    if a["brand"].lower() in nv.lower():
        return nv
    return "%s %s" % (nv, a["brand"])


def ctx_vi(a):
    """Ngu canh cot tieng Viet - uu tien ten da dich.

    Review 23/10: truoc day chi dich ten o cac dong Tag/Keyword/URL/Meta title,
    con {name}/{name_brand} trong Mo ta dai / Meta description / Mo ta thong so
    van la ten TIENG ANH -> cau tieng Viet lai chen nguyen ten Anh vao giua.
    """
    c = ctx(a)
    c["type_l"] = a["type_vi"].lower()
    nv = a.get("name_vi") or a["name"]
    if nv != a["name"]:
        # ten da dich thi BAN THAN no da la cum danh tu tieng Viet, ghep them
        # {type_c} nua se lap ("Dau noi ong Dau noi ong cho ong...").
        c["type_c"] = nv[:1].upper() + nv[1:]
        first = nv.split(" ")[0]
        c["type_l"] = (nv if (len(first) > 1 and first.isupper())
                       else nv[:1].lower() + nv[1:])
        c["name_brand"] = a["brand"]
        c["name"] = ""
    else:
        c["name_brand"] = with_brand_vi(a)
    return c


def fmt_vi(tpl, cv):
    """format() roi don khoang trang thua do {name} bi bo rong."""
    return re.sub(r"[ 	]{2,}", " ", tpl.format(**cv)).replace(" .", ".")


def safe_name(s, n=60):
    return re.sub(r'[\\/:*?"<>|]+', "-", s).strip()[:n].rstrip(" .")


def add_row(ws, label, vn, en):
    ws.append([label, vn, en])
    r = ws.max_row
    ws.cell(row=r, column=1).fill = FIELD_FILL
    ws.cell(row=r, column=1).font = FIELD_FONT
    for c in (1, 2, 3):
        ws.cell(row=r, column=c).alignment = WRAP
    return r


def build_article(a, i, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "NOI DUNG"
    ws["A1"] = "NỘI DUNG SẢN PHẨM NTREND - %s" % a["name"]
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["Trường", "Tiếng Việt (bản đăng web)", "Tiếng Anh (soát chính tả)"])
    hr = ws.max_row
    for c in range(1, 4):
        cell = ws.cell(row=hr, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT

    has_specs = bool(a["specs_en"])
    web_overview = a.get("overview_en") and True or False
    web_feats = bool(a.get("features_en"))
    cv, ce = ctx_vi(a), ctx(a)
    slug = a["slug"] if a["brand"].lower() in a["slug"].lower() \
        else "%s-%s" % (a["slug"], a["brand"])

    # ----- noi dung dai / ngan / thong so theo review -----
    if web_overview:
        long_vi = "\n\n".join(a["overview_vi"])
        long_en = "\n\n".join(a["overview_en"])
    else:
        long_vi = fmt_vi(LONG_TPL_VI, cv)
        long_en = LONG_TPL_EN.format(**ce)
    if web_feats:
        short_vi = "\n".join("- " + x for x in a["features_vi"])
        short_en = "\n".join("- " + x for x in a["features_en"])
    else:
        short_vi = fmt_vi(SHORT_TPL_VI, cv)
        short_en = SHORT_TPL_EN.format(**ce)
    if has_specs:
        sd_vi, sd_en = fmt_vi(SPECDESC_TPL_VI, cv), SPECDESC_TPL_EN.format(**ce)
    else:
        sd_vi, sd_en = fmt_vi(NOSPEC_DESC_VI, cv), NOSPEC_DESC_EN.format(**ce)

    add_row(ws, "Ảnh sản phẩm",
            os.path.basename(a["img_file"]) if a.get("img_file") else "(chưa có ảnh)", "")
    add_row(ws, "Danh mục sản phẩm", a.get("cat_vi") or "(chưa phân loại)",
            " > ".join(a["path"][:-1]) + "  (danh mục gốc perma, tham khảo)")
    nv = a.get("name_vi") or a["name"]
    add_row(ws, "Thẻ tag", "%s, %s" % (nv, a["brand"]),
            "%s, %s" % (a["name"], a["brand"]))
    add_row(ws, "Thương hiệu", a["brand"], a["brand"])
    kw_vi = "%s %s" % (nv, a["brand"]) if a["brand"].lower() not in nv.lower() else nv
    add_row(ws, "URL", vn_slug(kw_vi), slug)
    add_row(ws, "Keyword chính", kw_vi, with_brand(a))
    tc = a["type_vi"][:1].upper() + a["type_vi"][1:]
    add_row(ws, "Meta title",
            kw_vi if tc.lower() in kw_vi.lower() else "%s %s" % (tc, kw_vi),
            "%s %s" % (with_brand(a), a["type_en"].lower()))
    add_row(ws, "Meta description (155 - 165 characters)",
            fmt_vi(META_TPL_VI, cv), META_TPL_EN.format(**ce))
    add_row(ws, "Mô tả dài của sản phẩm", long_vi, long_en)
    add_row(ws, "Mô tả ngắn của sản phẩm", short_vi, short_en)
    add_row(ws, "Mô tả thông số", sd_vi, sd_en)

    # ----- bang: thong so ky thuat HOAC bang ma san pham -----
    if has_specs:
        ws.append(["Thông số chi tiết", "Tên thông số", "Giá trị",
                   "Parameter (EN)", "Value (EN)"])
        r = ws.max_row
        ws.cell(row=r, column=1).fill = FIELD_FILL
        ws.cell(row=r, column=1).font = FIELD_FONT
        for c in range(2, 6):
            ws.cell(row=r, column=c).font = Font(bold=True, italic=True)
        for (ke, ve), (kv, vv) in zip(a["specs_en"], a["specs_vi"]):
            ws.append(["", kv, vv, ke, ve])
            for c in range(2, 6):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP
    else:
        items = a.get("items") or []
        angle_rows = [ANGLE_RX.match(it["name"]) for it in items]
        if items and all(angle_rows):
            ws.append(["Bảng mã sản phẩm", "Góc", "Ren ngoài", "Ren trong",
                       "Vật liệu", "Mã đặt hàng (Part No.)"])
            r = ws.max_row
            ws.cell(row=r, column=1).fill = FIELD_FILL
            ws.cell(row=r, column=1).font = FIELD_FONT
            for c in range(2, 7):
                ws.cell(row=r, column=c).font = Font(bold=True, italic=True)
            for it, m in zip(items, angle_rows):
                ws.append(["", m.group(1), m.group(2), m.group(3),
                           it["material"].title(), it["art"]])
        elif items:
            # review 22/10: bang ma phai co ca ban tieng Viet lan tieng Anh
            # review 23/10: xep HET cot tieng Viet truoc roi moi den cot tieng
            # Anh, khong xen ke - tool dang web chi can khoi VI ben trai.
            ws.append(["Bảng mã sản phẩm", "Sản phẩm (VI)", "Vật liệu",
                       "Mã đặt hàng (Part No.)", "Product (EN)", "Material (EN)"])
            r = ws.max_row
            ws.cell(row=r, column=1).fill = FIELD_FILL
            ws.cell(row=r, column=1).font = FIELD_FONT
            for c in range(2, 7):
                ws.cell(row=r, column=c).font = Font(bold=True, italic=True)
            for it in items:
                ws.append(["", it.get("name_vi") or it["name"],
                           it.get("material_vi") or it["material"].title(),
                           it["art"], it["name"], it["material"].title()])
                for c in (2, 3, 5, 6):
                    ws.cell(row=ws.max_row, column=c).alignment = WRAP
        else:
            ws.append(["Bảng mã sản phẩm", "(trang hãng không có bảng mã)", "", ""])
            ws.cell(row=ws.max_row, column=1).fill = FIELD_FILL
            ws.cell(row=ws.max_row, column=1).font = FIELD_FONT

    ws.append([])
    ws.append(["Nguồn", a["url_src"], ""])
    for i2, w in enumerate([34, 56, 38, 26, 56, 30], 1):
        ws.column_dimensions[get_column_letter(i2)].width = w

    if a.get("img_file") and os.path.exists(a["img_file"]):
        # anh de rieng 1 folder cho de dang web (review 22/10), khong nhung xlsx
        img_dir = os.path.join(out_dir, "images")
        os.makedirs(img_dir, exist_ok=True)
        try:
            shutil.copy2(a["img_file"],
                         os.path.join(img_dir, os.path.basename(a["img_file"])))
        except Exception as e:
            print("  [anh loi]", a["name"][:30], repr(e)[:60])

    fname = "%02d - %s.xlsx" % (i, safe_name(a["name"]))
    wb.save(os.path.join(out_dir, fname))
    return fname


def build(brand_dir, out_dir, brand_label):
    with io.open(os.path.join(brand_dir, "content.json"), encoding="utf-8") as f:
        data = json.load(f)
    arts = data["articles"]
    os.makedirs(out_dir, exist_ok=True)

    idx = Workbook()
    iw = idx.active
    iw.title = "MUC LUC"
    iw["A1"] = "MỤC LỤC BÀI ĐĂNG - %s" % brand_label
    iw["A1"].font = Font(bold=True, size=13, color="1F4E79")
    iw.append([])
    iw.append(["STT", "Dòng sản phẩm", "File", "Ảnh", "Có TSKT",
               "Danh mục gốc (perma)", "Nguồn"])
    for c in range(1, 8):
        cell = iw.cell(row=3, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
    for i, w in enumerate([6, 42, 52, 6, 9, 50, 60], 1):
        iw.column_dimensions[get_column_letter(i)].width = w

    for i, a in enumerate(arts, 1):
        fname = build_article(a, i, out_dir)
        iw.append([i, a["name"], fname, "x" if a.get("img_file") else "",
                   "x" if a["specs_en"] else "",
                   " > ".join(a["path"][:-1]), a["url_src"]])
    idx.save(os.path.join(out_dir, "00 MUC LUC.xlsx"))
    print("Saved %d bai -> %s" % (len(arts), out_dir))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    build(os.path.join(HERE, "out", "05 Perma"),
          os.path.join(REPO, "data", "Noi dung Perma"),
          "05 Perma")
