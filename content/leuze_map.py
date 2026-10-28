# -*- coding: utf-8 -*-
"""Leuze - lap ban do 'nhanh danh muc' -> URL trang danh sach SP.

    python leuze_map.py

leuze.com chan moi request khong phai trinh duyet (403 ke ca robots.txt), nen
sitemap cung phai tai qua Chrome that (Playwright). Cay danh muc trong
'Danh mục SP web NTREND.xlsx' duoc dung tu chinh sitemap en-int nen duong dan
trong URL khop 1-1 voi cac cot Cap 1..5.

Ket qua: out/08 Leuze/urlmap.json
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

SITE = "https://www.leuze.com"
HOME = SITE + "/en-int"
ROBOTS = SITE + "/robots.txt"
OUT = os.path.join(HERE, "out", "08 Leuze")
NTREND_XLSX = os.path.join(REPO, "data", "Danh mục SP web NTREND.xlsx")
SHEET = "08 Leuze"

JS = """async (u) => { const r = await fetch(u, {credentials:'include'});
                      return {s: r.status, t: await r.text()}; }"""


def key(parts):
    """['Accessories','Mounting systems'] -> 'accessories/mountingsystems'."""
    out = []
    for p in parts:
        p = re.sub(r"[^0-9a-zA-Z]+", "", str(p or "")).lower()
        if p:
            out.append(p)
    return "/".join(out)


def url_key(u):
    path = u.split("/en-int/products/", 1)[-1].rstrip("/")
    return key(path.split("/"))


def fetch_sitemap():
    """robots.txt -> sitemap index en-int -> danh sach URL (tat ca qua Chrome)."""
    from playwright.sync_api import sync_playwright
    urls = set()
    with sync_playwright() as pw:
        b = pw.chromium.launch(headless=False, channel="chrome")
        pg = b.new_page()
        pg.goto(HOME, wait_until="domcontentloaded", timeout=120000)
        pg.wait_for_timeout(3000)

        def get(u):
            # file sitemap that nam o files.leuze.com -> fetch() bi CORS chan,
            # phai dieu hinh trinh duyet sang thang URL do roi doc body.
            if "files.leuze.com" in u or not u.startswith(SITE):
                pg.goto(u, wait_until="domcontentloaded", timeout=120000)
                return pg.evaluate("() => document.documentElement.textContent")
            r = pg.evaluate(JS, u)
            return r["t"] if r["s"] == 200 else ""

        rb = get(ROBOTS)
        sm = [l.split(":", 1)[1].strip() for l in rb.splitlines()
              if l.lower().startswith("sitemap:") and "/en-int/" in l]
        for sm_url in sm:
            idx = get(sm_url)
            locs = [x.strip() for x in re.findall(r"<loc>(.*?)</loc>", idx, re.S)]
            for loc in locs:
                for u in re.findall(r"<loc>(.*?)</loc>", get(loc), re.S):
                    u = u.strip()
                    if "/en-int/products/" in u:
                        urls.add(u)
        b.close()
    return {url_key(u): u for u in sorted(urls)}


BRAND_VI = "Leuze"


def cat_of(v):
    """Cot danh muc Ntrend chi ghi moi ten hang = chua phan loai."""
    c = (v or "").strip()
    if c == BRAND_VI:
        return ("Không có danh mục sản phẩm,  chỉ nằm trong danh mục "
                "Thương Hiệu " + BRAND_VI)
    return c


def wanted():
    """165 dong SP (cot 'La cap ngay tren ma SP' = x)."""
    from openpyxl import load_workbook
    ws = load_workbook(NTREND_XLSX, read_only=True, data_only=True)[SHEET]
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or len(r) < 10 or str(r[8] or "").strip() != "x":
            continue
        path = [str(x).strip() for x in r[3:8] if x and str(x).strip()]
        if not path:
            continue
        out.append({"stt": str(r[0]).strip(), "name": path[-1], "path": path,
                    "cat_vi": cat_of(r[9])})
    return out


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    os.makedirs(OUT, exist_ok=True)
    cache = os.path.join(OUT, "sitemap.json")
    if os.path.exists(cache) and "--refresh" not in sys.argv:
        m = json.load(io.open(cache, encoding="utf-8"))
    else:
        m = fetch_sitemap()
        json.dump(m, io.open(cache, "w", encoding="utf-8"), ensure_ascii=False)
    print("sitemap: %d trang /products/" % len(m))

    lines = wanted()
    hit, miss = [], []
    for l in lines:
        k = key(l["path"])
        u = m.get(k)
        if not u:
            cand = [v for kk, v in m.items() if kk.startswith(k)]
            if len(cand) == 1:
                u = cand[0]
        if u:
            l["url"] = u
            hit.append(l)
        else:
            miss.append(l)
    json.dump({"lines": lines}, io.open(os.path.join(OUT, "urlmap.json"), "w",
              encoding="utf-8"), ensure_ascii=False)
    print("dong SP: %d | khop: %d | chua khop: %d"
          % (len(lines), len(hit), len(miss)))
    for l in miss[:25]:
        print("   [?]", " > ".join(l["path"]))


if __name__ == "__main__":
    main()
