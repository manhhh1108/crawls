# -*- coding: utf-8 -*-
"""SICK - PHA 3: content.json -> moi bai 1 file Excel trong data/Noi dung SICK/.

Giu dung layout da duyet o Perma: cot A = truong, cot B = TIENG VIET (ban dang
web), cot C = tieng Anh (de soat chinh ta ban dich). Mau van theo
'Crawl noi dung Ntrend.docx' + 'ntrend template.xlsx'.

Them so voi Perma:
  - "Danh muc san pham" dien san tu cot "Danh muc SP Tieng Viet" cua file thong ke
  - co CA bang thong so ky thuat CA bang ma san pham
  - sinh them "00 CAN SOAT DICH.xlsx" = cac cau may dich con dang ngo
"""
import io
import json
import os
import re
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from openpyxl import Workbook                              # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill   # noqa: E402  # noqa
from openpyxl.utils import get_column_letter               # noqa: E402

from sick_names_vi import name_vi                          # noqa: E402
from build_content_excel import (LONG_TPL_VI, LONG_TPL_EN, SHORT_TPL_VI,       # noqa: E402
                                 SHORT_TPL_EN, SPECDESC_TPL_VI, SPECDESC_TPL_EN,
                                 HDR_FILL, FIELD_FILL, HDR_FONT, FIELD_FONT, WRAP)

IMG_SUBDIR = "images"       # anh de rieng 1 folder, khong nhung vao xlsx

BRAND_FULL_VI = "SICK (Đức)"
BRAND_FULL_EN = "SICK (Germany)"
MAX_ITEM_ROWS = None        # None = in het, khong cat bang ma san pham

NOSPEC_DESC_VI = ("{type_c} {name_brand} là sản phẩm công nghiệp chính hãng của "
                  + BRAND_FULL_VI + "; bảng mã sản phẩm chi tiết như bảng dưới:")
NOSPEC_DESC_EN = ("The {name_brand} {type_l} is a genuine industrial product from "
                  + BRAND_FULL_EN + "; the detailed part-number table is below:")
# mau chot trong review 'Crawl noi dung Ntrend.docx' - ngan hon ban cu 15 ky tu
META_TPL_VI = ("Mua {type_l} {name_brand} chính hãng – thông số kỹ thuật đầy đủ, "
               "giá tốt, hỗ trợ tư vấn và lựa chọn model phù hợp. "
               "Liên hệ Ntrend ngay!")
META_TPL_EN = ("Buy genuine {name_brand} {type_l} – full specifications, good "
               "price, support in selecting the right model. "
               "Contact Ntrend now for a consultation!")


def with_brand(a):
    return a["name"] if a["brand"].lower() in a["name"].lower() \
        else "%s %s" % (a["name"], a["brand"])


def with_brand_vi(a):
    """Nhu with_brand nhung dung ten da dich (xem sick_names_vi.py)."""
    n = name_vi(a["name"])
    return n if a["brand"].lower() in n.lower() else "%s %s" % (n, a["brand"])


def title_vi(cv, a):
    """'<loai> <ten> SICK', bo <loai> neu ten da chua san (vd Anten + Anten RFID)."""
    nb = with_brand_vi(a)
    t = cv["type_c"]
    return nb if not t or t.lower() in nb.lower() else "%s %s" % (t, nb)


MIN_HL, MAX_HL = 8, 10      # so gach dau dong cua "Mo ta dai"
MAX_HL_LEN = 220            # 1 gach dau dong dai toi da bao nhieu ky tu

# Truong KHONG bao gio dua vao "Mo ta dai" - la metadata catalog, dang web vo
# nghia: "San pham moi: Khong", "San pham: Bo loc bao ve" (lap lai chinh ten),
# "Phan khuc san pham: ..." (lap lai danh muc).
SPEC_SKIP = re.compile(
    r"^\s*(new\s*products?|products?|product\s*(sub)?segments?|product\s*family|"
    r"product\s*categor(y|ies)|part\s*(no|number)s?|art(icle)?\.?\s*(no|number)?s?|"
    r"order(ing)?\s*(no|number|code)s?|type\s*codes?|eclass|etim|unspsc|gtin|ean|"
    r"customs\s*tariff.*|country\s*of\s*origin|brand|manufacturer)\.?\s*$", re.I)

# thu tu uu tien khi bu them thong so - theo dung danh sach trong file review
SPEC_PRIORITY = [
    r"measuring range|detection range|sensing range|scanning range|working range",
    r"resolution",
    r"accurac|repeatab|linearity",
    r"response time|switching frequency|speed|scanning rate",
    r"supply voltage|power consumption|current consumption",
    r"output|signal",
    r"interface|communication|protocol|fieldbus",
    r"enclosure rating|protection class|degree of protection|ip\b",
    r"ambient temperature|operating temperature",
    r"housing material|connection type|light source|dimension|weight",
]


def _trim(s, n=MAX_HL_LEN):
    """SICK gop nhieu bien the vao 1 o, ngan bang ';'. Lay ve dau cho tron cau
    thay vi cat cung giua chung roi bo dau '…'."""
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    if len(s) > n and ";" in s:
        head = s.split(";")[0].strip()
        if 25 <= len(head) <= n:          # ve dau du nghia thi dung luon
            return head
    return s if len(s) <= n else s[:n].rsplit(" ", 1)[0] + "…"


def highlights(a):
    """'Mo ta dai' = 8-10 gach dau dong: Features/Benefits truoc, thieu thi bu
    bang thong so ky thuat quan trong (theo review Ntrend)."""
    vi = [x for x in (a.get("features_vi") or []) if x]
    en = [x for x in (a.get("features_en") or []) if x]
    n = min(len(vi), len(en))
    vi, en = vi[:n], en[:n]

    sp_en, sp_vi = a.get("specs_en") or [], a.get("specs_vi") or []
    used = set()

    def add_specs(patterns):
        for rx in patterns:
            for i, kvp in enumerate(sp_en):
                if len(vi) >= MAX_HL:
                    return
                if i in used or i >= len(sp_vi):
                    continue
                if SPEC_SKIP.match(str(kvp[0])):      # metadata catalog -> bo
                    continue
                if not re.search(rx, str(kvp[0]), re.I):
                    continue
                kv, vv = sp_vi[i][0], sp_vi[i][1]
                if not vv:
                    continue
                used.add(i)
                vi.append("%s: %s" % (kv, _trim(vv)))
                en.append("%s: %s" % (kvp[0], _trim(kvp[1])))

    if len(vi) < MAX_HL:
        add_specs(SPEC_PRIORITY)
    if len(vi) < MIN_HL:                       # van thieu -> lay thong so con lai
        add_specs([r"."])
    if not vi:
        return "", ""
    return ("\n".join("- " + x for x in vi[:MAX_HL]),
            "\n".join("- " + x for x in en[:MAX_HL]))


def vn_slug(s):
    """Keyword chinh (tieng Viet) -> URL: bo dau, thuong hoa, noi bang '-'.

    Chuan SEO Ntrend (chot 19/10): keyword chinh phai nam trong URL.
        'Anten GPS SICK' -> 'anten-gps-sick'
    """
    import unicodedata
    s = (s or "").replace("Đ", "D").replace("đ", "d")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).lower().strip("-")
    return re.sub(r"-{2,}", "-", s) or "sp"


def fmt_vi(tpl, cv):
    """format() cho cot tieng Viet, don khoang trang thua khi type_c/type_l rong."""
    return re.sub(r"\s{2,}", " ", tpl.format(**cv)).strip()


def singular(s):
    """'safety controllers' -> 'safety controller' (nhom ho SP cua SICK deu so nhieu)."""
    w = (s or "").split()
    if w and len(w[-1]) > 3 and w[-1].endswith("s") and not w[-1].endswith(("ss", "us", "is")):
        w[-1] = w[-1][:-1] if not w[-1].endswith("ies") else w[-1][:-3] + "y"
    return " ".join(w)


def ctx(a):
    t = singular(a["type_en"] or "sensor")
    return {"name": a["name"], "name_brand": with_brand(a), "brand": a["brand"],
            "type_c": (a["type_vi"] or "Sản phẩm")[:1].upper() + (a["type_vi"] or "Sản phẩm")[1:],
            "type_l": t.lower()}


def ctx_vi(a):
    c = ctx(a)
    c["type_l"] = (a["type_vi"] or "sản phẩm").lower()
    c["name"] = name_vi(a["name"])
    c["name_brand"] = with_brand_vi(a)
    # ten da dich sang tieng Viet thi ban than no da noi ro loai SP roi -> dung
    # luon ten lam "<loai>", khong viet "Bo cap nguon va cap cap nguon Bo nguon SICK"
    # cac mau cau ghep theo dang "{type_c} {name_brand}" / "{type_c} {name}",
    # nen dat ten da dich vao {type_c}/{type_l} va bo trong {name} - fmt_vi se
    # don khoang trang thua.
    if c["name"] != a["name"]:
        c["type_c"] = c["name"]
        # viet thuong chu dau cho hop giua cau, NHUNG giu nguyen neu tu dau la
        # chu viet tat (CSS, RFID, GPS...) - neu khong se ra "Mua cSS do phan giai cao"
        first = c["name"].split(" ")[0].split("/")[0]
        c["type_l"] = (c["name"] if (len(first) > 1 and first.isupper())
                       else c["name"][:1].lower() + c["name"][1:])
        c["name_brand"] = a["brand"]
        c["name"] = ""
    return c


def safe_name(s, n=58):
    return re.sub(r'[\\/:*?"<>|]+', "-", s).strip()[:n].rstrip(" .")


def add_row(ws, label, vn, en):
    ws.append([label, vn, en])
    r = ws.max_row
    ws.cell(row=r, column=1).fill = FIELD_FILL
    ws.cell(row=r, column=1).font = FIELD_FONT
    for c in (1, 2, 3):
        ws.cell(row=r, column=c).alignment = WRAP
    return r


def table_header(ws, label, cols):
    ws.append([label] + cols)
    r = ws.max_row
    ws.cell(row=r, column=1).fill = FIELD_FILL
    ws.cell(row=r, column=1).font = FIELD_FONT
    for c in range(2, len(cols) + 2):
        ws.cell(row=r, column=c).font = Font(bold=True, italic=True)
    return r


def build_article(a, i, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "NOI DUNG"
    ws["A1"] = "NỘI DUNG SẢN PHẨM NTREND - %s" % a["name"]
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.append([])
    ws.append(["Trường", "Tiếng Việt (bản đăng web)", "Tiếng Anh (soát chính tả)"])
    hr = ws.max_row
    for c in range(1, 4):
        ws.cell(row=hr, column=c).fill = HDR_FILL
        ws.cell(row=hr, column=c).font = HDR_FONT

    cv, ce = ctx_vi(a), ctx(a)
    has_specs = bool(a.get("specs_en"))
    # URL = keyword chinh tieng Viet, bo dau, noi bang '-' (chot 19/10)
    slug_en = a["slug"] if a["brand"].lower() in a["slug"].lower() \
        else "%s-%s" % (a["slug"], a["brand"].lower())
    slug = vn_slug(with_brand_vi(a))

    # Theo review 'Crawl noi dung Ntrend.docx':
    #   Mo ta DAI   = 8-10 gach dau dong dac diem/thong so quan trong
    #   Mo ta NGAN  = doan overview cua web hang
    if a.get("overview_vi"):
        short_vi = "\n\n".join(a["overview_vi"])
        short_en = "\n\n".join(a["overview_en"])
    else:
        short_vi, short_en = fmt_vi(SHORT_TPL_VI, cv), SHORT_TPL_EN.format(**ce)

    # Review 13/10: 83 dong SP hang khong cong bo Features/Benefits thi "Mo ta
    # dai" lay y nguyen "Mo ta ngan". Truoc day ghep tu metadata catalog -> ra
    # list "San pham moi: Khong / Phan khuc san pham: ..." vo nghia.
    long_vi, long_en = highlights(a) if a.get("features_vi") else ("", "")
    if not long_vi:
        long_vi, long_en = short_vi, short_en
    if has_specs:
        sd_vi, sd_en = fmt_vi(SPECDESC_TPL_VI, cv), SPECDESC_TPL_EN.format(**ce)
    elif a.get("items"):
        sd_vi, sd_en = fmt_vi(NOSPEC_DESC_VI, cv), NOSPEC_DESC_EN.format(**ce)
    else:
        sd_vi = fmt_vi("{type_c} {name_brand} là sản phẩm công nghiệp chính hãng của "
                       + BRAND_FULL_VI + ".", cv)
        sd_en = ("The {name_brand} {type_l} is a genuine industrial product from "
                 + BRAND_FULL_EN + ".").format(**ce)

    add_row(ws, "Ảnh sản phẩm",
            os.path.basename(a["img_file"]) if a.get("img_file") else "(chưa có ảnh)", "")
    add_row(ws, "Danh mục sản phẩm", a.get("cat_vi") or "(chưa phân loại)",
            a.get("cat_en") or " > ".join(a.get("path", [])[1:]))
    add_row(ws, "Thẻ tag", "%s, %s" % (name_vi(a["name"]), a["brand"]),
            "%s, %s" % (a["name"], a["brand"]))
    add_row(ws, "Thương hiệu", a.get("brand_field") or a["brand"],
            a.get("brand_field") or a["brand"])
    add_row(ws, "URL", slug, slug_en)
    add_row(ws, "Keyword chính", with_brand_vi(a), with_brand(a))
    add_row(ws, "Meta title", title_vi(cv, a),
            "%s %s" % (with_brand(a), ce["type_l"]))
    add_row(ws, "Meta description (155 - 165 characters)",
            fmt_vi(META_TPL_VI, cv), META_TPL_EN.format(**ce))
    add_row(ws, "Mô tả dài của sản phẩm", long_vi, long_en)
    add_row(ws, "Mô tả ngắn của sản phẩm", short_vi, short_en)
    add_row(ws, "Mô tả thông số", sd_vi, sd_en)

    if has_specs:
        table_header(ws, "Thông số chi tiết",
                     ["Tên thông số", "Giá trị", "Parameter (EN)", "Value (EN)"])
        for (ke, ve), (kv, vv) in zip(a["specs_en"], a["specs_vi"]):
            ws.append(["", kv, vv, ke, ve])
            for c in range(2, 6):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP

    items = a.get("items") or []
    if items:
        shown = items[:MAX_ITEM_ROWS]
        # Review 13/10: dong SP khong co bang TECHNICAL DATA OVERVIEW thi bang
        # ma san pham chinh LA bang thong so -> lay nguyen bang o tab Variants,
        # ke ca cot Description.
        has_desc = any(it.get("desc") for it in shown)
        label = "Bảng mã sản phẩm" if has_specs \
            else "Thông số chi tiết (bảng mã ở tab Variants)"
        if len(items) > len(shown):
            label += " (%d/%d model)" % (len(shown), len(items))
        cols = ["Model", "Mã đặt hàng (Part No.)"]
        if has_desc:
            cols += ["Mô tả", "Description (EN)"]
        table_header(ws, label, cols)
        for it in shown:
            row = ["", it["name"], it["art"]]
            if has_desc:
                row += [it.get("desc_vi") or "", it.get("desc") or ""]
            ws.append(row)
            for c in range(2, len(row) + 1):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP
        if len(items) > len(shown):
            ws.append(["", "... còn %d model khác, xem đầy đủ tại trang nguồn"
                       % (len(items) - len(shown)), ""])

    if not has_specs and not items:
        table_header(ws, "Bảng mã sản phẩm", ["(trang hãng không có dữ liệu)"])

    ws.append([])
    if a.get("not_on_web"):
        ws.append(["CẢNH BÁO",
                   "Dòng này không còn trên sick.com/sg/en (đã ngừng hoặc đổi tên) – "
                   "nội dung đang dùng mẫu văn, cần kiểm tra lại trước khi đăng", ""])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor="FFC7CE")
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="9C0006")
        ws.cell(row=ws.max_row, column=2).alignment = WRAP
    ws.append(["Nguồn", a["url_src"] or "(không còn trang trên web SICK)", ""])
    ws.append(["Mã danh mục SICK", a.get("code", ""), ""])
    for i2, w in enumerate([34, 62, 62, 34, 62], 1):
        ws.column_dimensions[get_column_letter(i2)].width = w

    if a.get("img_file") and os.path.exists(a["img_file"]):
        img_dir = os.path.join(out_dir, IMG_SUBDIR)
        os.makedirs(img_dir, exist_ok=True)
        try:
            shutil.copy2(a["img_file"],
                         os.path.join(img_dir, os.path.basename(a["img_file"])))
        except Exception:
            pass

    fname = "%03d - %s.xlsx" % (i, safe_name(a["name"]))
    wb.save(os.path.join(out_dir, fname))
    return fname


def build_review(out_dir, src_dir):
    p = os.path.join(src_dir, "review_dich.json")
    if not os.path.exists(p):
        return 0
    with io.open(p, encoding="utf-8") as f:
        rows = json.load(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "CAN SOAT"
    ws["A1"] = "CÁC CÂU DỊCH CẦN SOÁT TAY - SICK"
    ws["A1"].font = Font(bold=True, size=13, color="C00000")
    ws.append([])
    ws.append(["Dòng sản phẩm", "Trường", "Lý do nghi ngờ", "Tiếng Anh", "Bản dịch"])
    for c in range(1, 6):
        ws.cell(row=3, column=c).fill = HDR_FILL
        ws.cell(row=3, column=c).font = HDR_FONT
    for r in rows:
        ws.append(r)
        for c in range(1, 6):
            ws.cell(row=ws.max_row, column=c).alignment = WRAP
    for i, w in enumerate([26, 20, 26, 70, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    wb.save(os.path.join(out_dir, "00 CAN SOAT DICH.xlsx"))
    return len(rows)


def build(src_dir, out_dir):
    with io.open(os.path.join(src_dir, "content.json"), encoding="utf-8") as f:
        arts = json.load(f)["articles"]
    os.makedirs(out_dir, exist_ok=True)

    idx = Workbook()
    iw = idx.active
    iw.title = "MUC LUC"
    iw["A1"] = "MỤC LỤC BÀI ĐĂNG - 01 SICK"
    iw["A1"].font = Font(bold=True, size=13, color="1F4E79")
    iw.append([])
    iw.append(["STT", "Dòng sản phẩm", "File", "Ảnh", "Nguồn thông số", "Số model",
               "Danh mục Ntrend (VI)", "Nguồn"])
    for c in range(1, 9):
        iw.cell(row=3, column=c).fill = HDR_FILL
        iw.cell(row=3, column=c).font = HDR_FONT
    for i, w in enumerate([6, 34, 50, 6, 22, 10, 52, 62], 1):
        iw.column_dimensions[get_column_letter(i)].width = w
    iw.freeze_panes = "A4"

    written = set()
    for i, a in enumerate(arts, 1):
        fname = build_article(a, i, out_dir)
        written.add(fname)
        iw.append([i, a["name"], fname, "x" if a.get("img_file") else "",
                   "Technical data overview" if a.get("specs_en")
                   else ("Bảng mã ở Variants" if a.get("items") else "(không có)"),
                   len(a.get("items") or []),
                   a.get("cat_vi", ""), a["url_src"]])
    idx.save(os.path.join(out_dir, "00 MUC LUC.xlsx"))
    n = build_review(out_dir, src_dir)

    # Doi thu tu / doi ten dong SP la file cu khong bi ghi de -> thu muc lan
    # lon ca 2 phien ban (lan 18/10 con sot 618 file cua lan 10/10). Xoa het
    # file bai KHONG nam trong lan dung nay. Chi dong toi file dat ten theo
    # dung mau "NNN - <ten>.xlsx", khong dong vao "00 ..." hay thu muc images/.
    stale = [f for f in os.listdir(out_dir)
             if re.match(r"^\d{3} - .+\.xlsx$", f) and f not in written]
    for f in stale:
        try:
            os.remove(os.path.join(out_dir, f))
        except OSError:
            pass

    print("Saved %d bai -> %s" % (len(arts), out_dir))
    if stale:
        print("  da xoa %d file bai cu khong con dung" % len(stale))
    if n:
        print("  + 00 CAN SOAT DICH.xlsx (%d dong can soat tay)" % n)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    build(os.path.join(HERE, "out", "01 SICK"),
          os.path.join(REPO, "data", "Noi dung SICK"))
